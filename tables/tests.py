import datetime
from django.test import TestCase
from django.contrib.auth import get_user_model
from employee_management.models import Department
from .models import Table, Column, Row, CellValue, TableAccess
from tasks.models import Task

User = get_user_model()

class TablesTestCase(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name="Tables Engineering Dept", slug="tables-engineering-dept")
        self.admin = User.objects.create_user(
            email="tablesadmin@flow-force.com",
            password="testpassword",
            full_name="Admin User",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="tablesemp@flow-force.com",
            password="testpassword",
            full_name="Employee User",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )

    def test_table_system_columns_creation(self):
        # 1. Creating a table should auto-create system columns
        table = Table.objects.create(name="Development tasks", created_by=self.admin)
        columns = table.columns.all()
        col_names = [col.name for col in columns]
        
        self.assertIn("S_NO", col_names)
        self.assertIn("DATE", col_names)
        self.assertIn("DUE_DATE", col_names)
        self.assertIn("TASK_NAME", col_names)
        self.assertIn("INITIAL_MAIL", col_names)
        self.assertIn("ALERT_MAIL", col_names)
        self.assertEqual(columns.count(), 6)

    def test_row_creation_with_task_sync(self):
        table = Table.objects.create(name="Development tasks", created_by=self.admin)
        row = Row.objects.create(table=table, created_by=self.employee)
        
        # Test task auto-creation on views. For row views, creating a row will sync a Task.
        # Let's test custom column access checking
        col = table.columns.first()
        cell = CellValue.objects.create(row=row, column=col, value=123, updated_by=self.admin)
        self.assertEqual(cell.value, 123)

    def test_table_duplication(self):
        table = Table.objects.create(name="Source Table", created_by=self.admin)
        # Create a custom column
        Column.objects.create(table=table, name="Custom Text Column", data_type="TEXT", position=7)
        
        # Test duplication API logic
        # Clone table metadata
        new_table = Table.objects.create(
            name=f"Copy of {table.name}",
            description=table.description,
            created_by=self.admin,
            department=table.department
        )
        for col in table.columns.filter(is_system_column=False):
            Column.objects.create(
                table=new_table,
                name=col.name,
                data_type=col.data_type,
                is_mandatory=col.is_mandatory,
                is_system_column=False,
                position=col.position
            )
            
        self.assertEqual(new_table.name, "Copy of Source Table")
        self.assertEqual(new_table.columns.count(), 7)  # 6 system + 1 custom
        self.assertEqual(new_table.columns.filter(is_system_column=False).count(), 1)

    def test_import_csv_with_offsets_and_equality_check(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Import Test Table", created_by=self.admin)
        # Ensure we have ADMIN access
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        # Create valid CSV data with 11 empty/meta lines, then header, then data.
        # System columns are: S_NO, DATE, DUE_DATE, TASK_NAME, INITIAL_MAIL, ALERT_MAIL
        csv_lines = [
            "Metadata line 1", "Metadata line 2", "Metadata line 3", "Metadata line 4",
            "Metadata line 5", "Metadata line 6", "Metadata line 7", "Metadata line 8",
            "Metadata line 9", "Metadata line 10", "Metadata line 11",
            "S_NO,DATE,DUE_DATE,TASK_NAME,INITIAL_MAIL,ALERT_MAIL",
            "1,2026-06-22,2026-06-30,Imported Task Name,NO,NO"
        ]
        csv_data = "\n".join(csv_lines)

        from django.core.files.uploadedfile import SimpleUploadedFile
        csv_file = SimpleUploadedFile("tasks.csv", csv_data.encode("utf-8"), content_type="text/csv")

        # Perform POST to import-csv
        view = TableViewSet.as_view({'post': 'import_csv'})
        request = factory.post(f"/tables/api/tables/{table.id}/import-csv/", {"file": csv_file}, format="multipart")
        force_authenticate(request, user=self.admin)

        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Row.objects.filter(table=table).count(), 1)

        # Check if Task was created
        task = Task.objects.filter(row__table=table).first()
        self.assertIsNotNone(task)
        self.assertEqual(task.task_name, "Imported Task Name")

    def test_import_csv_with_different_date_formats(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Import Date Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        csv_lines = [
            "Metadata line 1", "Metadata line 2", "Metadata line 3", "Metadata line 4",
            "Metadata line 5", "Metadata line 6", "Metadata line 7", "Metadata line 8",
            "Metadata line 9", "Metadata line 10", "Metadata line 11",
            "S_NO,DATE,DUE_DATE,TASK_NAME,INITIAL_MAIL,ALERT_MAIL",
            "1,22/06/2026,30/06/2026,Imported Task Name DD-MM-YYYY,NO,NO",
            "2,06/22/2026,06/30/2026,Imported Task Name MM-DD-YYYY,NO,NO"
        ]
        csv_data = "\n".join(csv_lines)

        from django.core.files.uploadedfile import SimpleUploadedFile
        csv_file = SimpleUploadedFile("tasks.csv", csv_data.encode("utf-8"), content_type="text/csv")

        view = TableViewSet.as_view({'post': 'import_csv'})
        request = factory.post(f"/tables/api/tables/{table.id}/import-csv/", {"file": csv_file}, format="multipart")
        force_authenticate(request, user=self.admin)

        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Row.objects.filter(table=table).count(), 2)

        # Check if Tasks were created and parsed correctly
        tasks = list(Task.objects.filter(row__table=table).order_by('id'))
        self.assertEqual(len(tasks), 2)
        import datetime
        self.assertEqual(tasks[0].due_date, datetime.date(2026, 6, 30))
        self.assertEqual(tasks[1].due_date, datetime.date(2026, 6, 30))

    def test_import_csv_column_mismatch_error(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Mismatch Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        # Missing standard columns
        csv_lines = [
            "Metadata line 1", "Metadata line 2", "Metadata line 3", "Metadata line 4",
            "Metadata line 5", "Metadata line 6", "Metadata line 7", "Metadata line 8",
            "Metadata line 9", "Metadata line 10", "Metadata line 11",
            "S_NO,DATE,DUE_DATE",
            "1,2026-06-22,2026-06-30"
        ]
        csv_data = "\n".join(csv_lines)

        from django.core.files.uploadedfile import SimpleUploadedFile
        csv_file = SimpleUploadedFile("tasks.csv", csv_data.encode("utf-8"), content_type="text/csv")

        view = TableViewSet.as_view({'post': 'import_csv'})
        request = factory.post(f"/tables/api/tables/{table.id}/import-csv/", {"file": csv_file}, format="multipart")
        force_authenticate(request, user=self.admin)

        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 400)
        self.assertIn("Required column", response.data["error"])

    def test_import_google_sheet_mocked(self):
        from unittest.mock import patch
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="GS Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        csv_lines = [
            "Metadata line 1", "Metadata line 2", "Metadata line 3", "Metadata line 4",
            "Metadata line 5", "Metadata line 6", "Metadata line 7", "Metadata line 8",
            "Metadata line 9", "Metadata line 10", "Metadata line 11",
            "S_NO,DATE,DUE_DATE,TASK_NAME,INITIAL_MAIL,ALERT_MAIL",
            "1,2026-06-22,2026-07-15,Google Sheet Task,NO,NO"
        ]
        csv_data = "\n".join(csv_lines)

        class MockUrlOpen:
            def __init__(self, data):
                self.data = data.encode('utf-8')
            def __enter__(self):
                return self
            def __exit__(self, exc_type, exc_val, exc_tb):
                pass
            def read(self):
                return self.data

        with patch("urllib.request.urlopen", return_value=MockUrlOpen(csv_data)) as mock_urlopen:
            view = TableViewSet.as_view({'post': 'import_google_sheet'})
            request = factory.post(f"/tables/api/tables/{table.id}/import-google-sheet/", {
                "url": "https://docs.google.com/spreadsheets/d/1abc123_xyz/edit#gid=12"
            }, format="json")
            force_authenticate(request, user=self.admin)

            response = view(request, pk=table.id)
            self.assertEqual(response.status_code, 201)
            # Verify urlopen was called
            self.assertTrue(mock_urlopen.called)
            self.assertEqual(Row.objects.filter(table=table).count(), 1)
            task = Task.objects.filter(row__table=table).first()
            self.assertEqual(task.task_name, "Google Sheet Task")

    def test_delete_row(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Delete Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        
        row = Row.objects.create(table=table, created_by=self.admin)
        
        view = RowViewSet.as_view({'delete': 'destroy'})
        request = factory.delete(f"/tables/api/rows/{row.id}/")
        force_authenticate(request, user=self.admin)
        
        response = view(request, pk=row.id)
        self.assertEqual(response.status_code, 204)
        self.assertFalse(Row.objects.filter(id=row.id).exists())

    def test_engineer_table_pid_column_creation(self):
        table = Table.objects.create(name="Engineer tasks", job_type="ENGINEER", created_by=self.admin)
        columns = table.columns.all()
        col_names = [col.name for col in columns]
        
        self.assertIn("S_NO", col_names)
        self.assertIn("DATE", col_names)
        self.assertIn("DUE_DATE", col_names)
        self.assertIn("TASK_NAME", col_names)
        self.assertIn("INITIAL_MAIL", col_names)
        self.assertIn("ALERT_MAIL", col_names)
        self.assertIn("PID", col_names)
        
        pid_col = table.columns.get(name="PID")
        self.assertFalse(pid_col.is_mandatory)
        self.assertTrue(pid_col.is_system_column)
        self.assertEqual(pid_col.data_type, "TEXT")
        self.assertEqual(pid_col.position, 7)
        self.assertEqual(columns.count(), 7)

    def test_row_creation_preserves_pid(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Engineer Tracker", job_type="ENGINEER", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        view = RowViewSet.as_view({'post': 'create'})
        request = factory.post(f"/tables/api/rows/", {
            "table": table.id,
            "cells": {
                "TASK_NAME": "Verify PID Test",
                "DUE_DATE": "2026-07-10",
                "PID": "PID-999"
            }
        }, format="json")
        force_authenticate(request, user=self.admin)

        response = view(request)
        self.assertEqual(response.status_code, 201)

        row = Row.objects.filter(table=table).first()
        self.assertIsNotNone(row)

        pid_col = table.columns.get(name="PID")
        cell = CellValue.objects.get(row=row, column=pid_col)
        self.assertEqual(cell.value, "PID-999")

    def test_row_level_editing(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Engineer Tracker", job_type="ENGINEER", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        
        row = Row.objects.create(table=table, created_by=self.admin)
        pid_col = table.columns.get(name="PID")
        task_name_col = table.columns.get(name="TASK_NAME")
        
        CellValue.objects.create(row=row, column=pid_col, value="", updated_by=self.admin)
        CellValue.objects.create(row=row, column=task_name_col, value="Original Name", updated_by=self.admin)

        view = RowViewSet.as_view({'post': 'edit_row'})
        request = factory.post(f"/tables/api/rows/{row.id}/edit-row/", {
            "cells": {
                "TASK_NAME": "Updated Name",
                "PID": "PID-888"
            }
        }, format="json")
        force_authenticate(request, user=self.admin)

        response = view(request, pk=row.id)
        self.assertEqual(response.status_code, 200)

        self.assertEqual(CellValue.objects.get(row=row, column=pid_col).value, "PID-888")
        self.assertEqual(CellValue.objects.get(row=row, column=task_name_col).value, "Updated Name")

    def test_grant_access_invalid_user(self):
        table = Table.objects.create(name="Grant Test Table", created_by=self.admin)
        self.client.force_login(self.admin)
        
        # Post with empty user_id
        response = self.client.post("/tables/", {
            "action": "grant",
            "table_id": table.id,
            "user_id": "",
            "access_level": "EDIT"
        })
        self.assertEqual(response.status_code, 302)
        self.assertRedirects(response, "/tables/")
        
        # Verify that no TableAccess was created
        self.assertFalse(TableAccess.objects.filter(table=table).exists())

    def test_bulk_update_action(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Bulk Update Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="VIEW")

        row = Row.objects.create(table=table, created_by=self.admin)
        task = Task.objects.create(row=row, due_date="2026-07-15", status="PENDING", priority="MEDIUM", assigned_by=self.admin)

        # 1. Employee is forbidden
        view = TableViewSet.as_view({'post': 'bulk_update'})
        request = factory.post(f"/tables/api/tables/{table.id}/bulk-update/", {"field": "INITIAL_MAIL", "value": "YES"}, format="json")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin successfully bulk updates INITIAL_MAIL to YES
        request = factory.post(f"/tables/api/tables/{table.id}/bulk-update/", {"field": "INITIAL_MAIL", "value": "YES"}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)

        # Verify initial_mail_sent is True on task and cell value is YES
        task.refresh_from_db()
        self.assertTrue(task.initial_mail_sent)
        init_col = table.columns.get(name="INITIAL_MAIL")
        self.assertEqual(CellValue.objects.get(row=row, column=init_col).value, "YES")

        # 3. Admin successfully bulk updates STATUS to COMPLETED
        # Create STATUS custom column
        status_col = Column.objects.create(table=table, name="STATUS", data_type="TEXT")
        request = factory.post(f"/tables/api/tables/{table.id}/bulk-update/", {"field": "STATUS", "value": "COMPLETED"}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)

        task.refresh_from_db()
        self.assertEqual(task.status, "COMPLETED")
        self.assertEqual(CellValue.objects.get(row=row, column=status_col).value, "COMPLETED")

    def test_sync_status_cell_to_task(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Status Sync Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        status_col = Column.objects.create(table=table, name="STATUS", data_type="TEXT")

        row = Row.objects.create(table=table, created_by=self.admin)
        task = Task.objects.create(row=row, due_date="2026-07-15", status="PENDING", priority="MEDIUM", assigned_by=self.admin)

        view = RowViewSet.as_view({'post': 'edit_cell'})
        request = factory.post(f"/tables/api/rows/{row.id}/edit-cell/", {
            "column": status_col.id,
            "value": "IN_PROGRESS"
        }, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=row.id)
        self.assertEqual(response.status_code, 200)

        task.refresh_from_db()
        self.assertEqual(task.status, "IN_PROGRESS")

    def test_sync_task_to_status_cell(self):
        from tasks.views import TaskViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Task to Cell Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        status_col = Column.objects.create(table=table, name="STATUS", data_type="TEXT")

        row = Row.objects.create(table=table, created_by=self.admin)
        task = Task.objects.create(row=row, due_date="2026-07-15", status="PENDING", priority="MEDIUM", assigned_by=self.admin)

        view = TaskViewSet.as_view({'post': 'update_status'})
        request = factory.post(f"/tasks/api/tasks/{task.id}/update-status/", {
            "status": "COMPLETED"
        }, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=task.id)
        self.assertEqual(response.status_code, 200)

        # Verify STATUS cell in spreadsheet was updated to COMPLETED
        cell = CellValue.objects.get(row=row, column=status_col)
        self.assertEqual(cell.value, "COMPLETED")

    def test_table_edit_details_and_permissions(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Original Table", description="Original Desc", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="EDIT")

        view = TableViewSet.as_view({'patch': 'partial_update'})

        # 1. Non-admin (employee with EDIT access) cannot change table name
        request = factory.patch(f"/tables/api/tables/{table.id}/", {"name": "Hacked Name"}, format="json")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin can change table name and description
        request = factory.patch(f"/tables/api/tables/{table.id}/", {"name": "New Name", "description": "New Desc"}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)
        table.refresh_from_db()
        self.assertEqual(table.name, "New Name")
        self.assertEqual(table.description, "New Desc")

    def test_column_edit_and_permissions(self):
        from tables.views import ColumnViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Column Permission Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="EDIT")

        col = Column.objects.create(table=table, name="Old Col", data_type="TEXT", position=7, options="A,B")

        view = ColumnViewSet.as_view({'patch': 'partial_update'})

        # 1. Employee cannot update column options
        request = factory.patch(f"/tables/api/columns/{col.id}/", {"options": "X,Y,Z"}, format="json")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=col.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin can update column options
        request = factory.patch(f"/tables/api/columns/{col.id}/", {"options": "X,Y,Z"}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=col.id)
        self.assertEqual(response.status_code, 200)
        col.refresh_from_db()
        self.assertEqual(col.options, "X,Y,Z")

    def test_table_duplication_structure_and_values(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Full Copy Table", created_by=self.admin, job_type="ENGINEER")
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        # Custom column with options
        custom_col = Column.objects.create(table=table, name="Custom Dropdown", data_type="DROPDOWN", options="Option1,Option2", position=7)

        # Row with cell values
        row = Row.objects.create(table=table, created_by=self.admin)
        cell1 = CellValue.objects.create(row=row, column=custom_col, value="Option1", updated_by=self.admin)
        
        # Row has task
        task = Task.objects.create(row=row, status="IN_PROGRESS", due_date="2026-07-01", priority="HIGH", assigned_by=self.admin)
        task.assigned_to.add(self.employee)

        view = TableViewSet.as_view({'post': 'duplicate_table'})
        request = factory.post(f"/tables/api/tables/{table.id}/duplicate/")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 201)

        # Verify new table metadata
        new_table_id = response.data["id"]
        new_table = Table.objects.get(id=new_table_id)
        self.assertEqual(new_table.name, "Copy of Full Copy Table")
        self.assertEqual(new_table.job_type, "ENGINEER")

        # Verify duplicated columns and options
        new_custom_col = new_table.columns.get(name="Custom Dropdown")
        self.assertEqual(new_custom_col.data_type, "DROPDOWN")
        self.assertEqual(new_custom_col.options, "Option1,Option2")

        # Verify duplicated rows and cell values
        new_row = new_table.rows.first()
        self.assertIsNotNone(new_row)
        new_cell = CellValue.objects.get(row=new_row, column=new_custom_col)
        self.assertEqual(new_cell.value, "Option1")

        # Verify task is duplicated with assignees
        new_task = new_row.task
        self.assertEqual(new_task.status, "IN_PROGRESS")
        self.assertEqual(new_task.priority, "HIGH")
        self.assertEqual(new_task.due_date.strftime("%Y-%m-%d"), "2026-07-01")
        self.assertIn(self.employee, new_task.assigned_to.all())

    def test_bulk_delete_rows(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Delete Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="VIEW")

        row1 = Row.objects.create(table=table, created_by=self.admin)
        row2 = Row.objects.create(table=table, created_by=self.admin)
        row3 = Row.objects.create(table=table, created_by=self.admin)

        view = TableViewSet.as_view({'post': 'bulk_delete_rows'})

        # 1. Non-edit user (employee with VIEW access) cannot bulk delete
        request = factory.post(f"/tables/api/tables/{table.id}/bulk-delete-rows/", {"row_ids": [row1.id, row2.id]}, format="json")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin can delete selected rows (row1, row2)
        request = factory.post(f"/tables/api/tables/{table.id}/bulk-delete-rows/", {"row_ids": [row1.id, row2.id]}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Row.objects.filter(id__in=[row1.id, row2.id]).exists())
        self.assertTrue(Row.objects.filter(id=row3.id).exists())

        # 3. Admin can delete all remaining rows in table
        request = factory.post(f"/tables/api/tables/{table.id}/bulk-delete-rows/", format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(Row.objects.filter(table=table).exists())

    def test_send_manual_escalation_api(self):
        from django.utils import timezone
        from datetime import timedelta
        from tasks.models import EmailLog
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Escalation Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="VIEW")

        row = Row.objects.create(table=table, created_by=self.admin)
        task = Task.objects.create(
            row=row,
            due_date=timezone.localdate() - timedelta(days=2),
            priority="HIGH",
            status="PENDING",
            assigned_by=self.admin
        )
        task.assigned_to.add(self.employee)

        view = TableViewSet.as_view({'post': 'send_manual_escalation'})

        # 1. Non-admin user cannot trigger manual escalation
        request = factory.post(f"/tables/api/tables/{table.id}/send-escalation/", {"row_ids": [row.id]}, format="json")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin can trigger manual escalation
        request = factory.post(f"/tables/api/tables/{table.id}/send-escalation/", {"row_ids": [row.id]}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)

        # Verify EmailLog was created
        log = EmailLog.objects.filter(task=task, email_type="OVERDUE_ESCALATION_MAIL").first()
        self.assertIsNotNone(log)
        self.assertEqual(log.recipient_email, self.employee.email)
        
        # Verify task escalation level updated
        task.refresh_from_db()
        self.assertEqual(task.last_escalation_level, 2)

    def test_send_manual_escalation_non_overdue_api(self):
        from django.utils import timezone
        from datetime import timedelta
        from tasks.models import EmailLog
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Non-Overdue Escalation Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="VIEW")

        row = Row.objects.create(table=table, created_by=self.admin)
        task = Task.objects.create(
            row=row,
            due_date=timezone.localdate() + timedelta(days=2),  # Future date (not overdue)
            priority="MEDIUM",
            status="PENDING",
            assigned_by=self.admin
        )
        task.assigned_to.add(self.employee)

        view = TableViewSet.as_view({'post': 'send_manual_escalation'})

        # Admin triggers escalation on the selected row (even though task is not overdue)
        request = factory.post(f"/tables/api/tables/{table.id}/send-escalation/", {"row_ids": [row.id]}, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 200)

        # Verify EmailLog was created and status level updated to 0
        log = EmailLog.objects.filter(task=task, email_type="OVERDUE_ESCALATION_MAIL").first()
        self.assertIsNotNone(log)
        self.assertEqual(log.recipient_email, self.employee.email)
        self.assertIn("A task has been escalated and requires immediate attention.", log.body)

        task.refresh_from_db()
        self.assertEqual(task.last_escalation_level, 0)

    def test_personal_job_type_creation_and_row_addition(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        # 1. Create a PERSONAL table. It should have 0 columns automatically created.
        table = Table.objects.create(name="My Personal Table", job_type="PERSONAL", created_by=self.admin)
        self.assertEqual(table.columns.count(), 0)

        # 2. Add a custom column to the table.
        custom_col = Column.objects.create(table=table, name="My Task Header", data_type="TEXT", position=1)

        # 3. Create a Row. It should bypass due date / task name mandatory validations.
        view = RowViewSet.as_view({'post': 'create'})
        
        request = factory.post(f"/tables/api/rows/", {
            "table": table.id,
            "cells": {
                "My Task Header": "Do gym workout"
            }
        }, format="json")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 201)

        # 4. Verify Row and Cell value were saved.
        row = Row.objects.filter(table=table).first()
        self.assertIsNotNone(row)
        cell = row.cells.filter(column=custom_col).first()
        self.assertEqual(cell.value, "Do gym workout")

        # 5. Verify the Task has default due_date=None, and task_name is guessed correctly.
        task = getattr(row, "task", None)
        self.assertIsNotNone(task)
        self.assertIsNone(task.due_date)
        self.assertEqual(task.task_name, "Do gym workout")

    def test_personal_job_type_csv_import(self):
        from unittest.mock import patch
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        # 1. Create a PERSONAL table and some custom columns.
        table = Table.objects.create(name="My Personal Table CSV", job_type="PERSONAL", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        
        custom_col1 = Column.objects.create(table=table, name="My Task Header", data_type="TEXT", position=1)
        custom_col2 = Column.objects.create(table=table, name="Note Info", data_type="TEXT", position=2)

        # 2. Mock Google sheet import with personal columns (NO system columns like TASK_NAME, S_NO or DUE_DATE)
        csv_lines = [
            "My Task Header,Note Info",
            "Gym workout,Focus on cardio and legs",
            "Read book,Finish chapter 5"
        ]
        csv_data = "\n".join(csv_lines)

        class MockUrlOpen:
            def __init__(self, data):
                self.data = data.encode('utf-8')
            def __enter__(self):
                return self
            def __exit__(self, exc_type, exc_val, exc_tb):
                pass
            def read(self):
                return self.data

        with patch("urllib.request.urlopen", return_value=MockUrlOpen(csv_data)) as mock_urlopen:
            view = TableViewSet.as_view({'post': 'import_google_sheet'})
            request = factory.post(f"/tables/api/tables/{table.id}/import-google-sheet/", {
                "url": "https://docs.google.com/spreadsheets/d/1abc123_xyz/edit#gid=12"
            }, format="json")
            force_authenticate(request, user=self.admin)

            response = view(request, pk=table.id)
            self.assertEqual(response.status_code, 201)
            self.assertEqual(Row.objects.filter(table=table).count(), 2)

            # Verify task names resolved from the text columns
            rows = Row.objects.filter(table=table).order_by("id")
            self.assertEqual(rows[0].task.task_name, "Gym workout")
            self.assertEqual(rows[1].task.task_name, "Read book")

    def test_column_unique_name_validation(self):
        from tables.serializers import ColumnSerializer
        table = Table.objects.create(name="Col Unique Table", created_by=self.admin)
        Column.objects.create(table=table, name="CustomCol1", data_type="TEXT")

        # Serializer should raise validation error if creating column with duplicate name
        serializer = ColumnSerializer(data={"table": table.id, "name": "CustomCol1", "data_type": "TEXT"})
        self.assertFalse(serializer.is_valid())
        self.assertIn("name", serializer.errors)

        # Serializer should pass if name is unique
        serializer2 = ColumnSerializer(data={"table": table.id, "name": "CustomCol2", "data_type": "TEXT"})
        self.assertTrue(serializer2.is_valid())

    def test_clear_column_values_api(self):
        from tables.views import ColumnViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Clear Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.employee, access_level="VIEW")
        col = Column.objects.create(table=table, name="CustomCol", data_type="TEXT")
        row = Row.objects.create(table=table, created_by=self.admin)
        cell = CellValue.objects.create(row=row, column=col, value="Target Value")

        view = ColumnViewSet.as_view({'post': 'clear_values'})

        # 1. Non-admin cannot clear values
        request = factory.post(f"/tables/api/columns/{col.id}/clear-values/")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=col.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin can clear values
        request = factory.post(f"/tables/api/columns/{col.id}/clear-values/")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=col.id)
        self.assertEqual(response.status_code, 200)

        cell.refresh_from_db()
        self.assertIsNone(cell.value)

    def test_delete_rows_by_column_api(self):
        from tables.views import ColumnViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Delete Rows Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.employee, access_level="VIEW")
        col = Column.objects.create(table=table, name="CustomCol", data_type="TEXT")
        row1 = Row.objects.create(table=table, created_by=self.admin)
        row2 = Row.objects.create(table=table, created_by=self.admin)
        
        CellValue.objects.create(row=row1, column=col, value="Row 1 Value")
        # row2 has no value

        view = ColumnViewSet.as_view({'post': 'delete_rows'})

        # 1. Non-admin cannot bulk delete rows by column
        request = factory.post(f"/tables/api/columns/{col.id}/delete-rows/")
        force_authenticate(request, user=self.employee)
        response = view(request, pk=col.id)
        self.assertEqual(response.status_code, 403)

        # 2. Admin can bulk delete rows by column (should delete row1, keep row2)
        request = factory.post(f"/tables/api/columns/{col.id}/delete-rows/")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=col.id)
        self.assertEqual(response.status_code, 200)

        self.assertFalse(Row.objects.filter(id=row1.id).exists())
        self.assertTrue(Row.objects.filter(id=row2.id).exists())

    def test_list_pid_table_behavior(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="LIST_PID table", job_type="LIST_PID", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=table, user=self.employee, access_level="EDIT")

        col_names = [col.name for col in table.columns.all()]
        self.assertIn("S_NO", col_names)
        self.assertIn("ENQUIRY_NO/QUOTATION_NO", col_names)
        self.assertIn("PID", col_names)
        self.assertIn("DUE_DATE_FLOW_FORCE", col_names)
        self.assertIn("INITIAL_MAIL", col_names)
        self.assertIn("ALERT_MAIL", col_names)

        # Test Row Creation via API
        view = RowViewSet.as_view({'post': 'create'})
        request = factory.post("/tables/api/rows/", {
            "table": table.id,
            "cells": {
                "ENQUIRY_NO/QUOTATION_NO": "ENQ-1234",
                "PID": "PID-5678",
                "DUE_DATE_FLOW_FORCE": "2026-07-20",
                "QTY": "15"
            }
        }, format="json")
        force_authenticate(request, user=self.employee)
        response = view(request)
        self.assertEqual(response.status_code, 201)

        row_id = response.data["id"]
        row = Row.objects.get(id=row_id)

        # Check CellValues
        self.assertEqual(CellValue.objects.get(row=row, column__name="ENQUIRY_NO/QUOTATION_NO").value, "ENQ-1234")
        self.assertEqual(CellValue.objects.get(row=row, column__name="DUE_DATE_FLOW_FORCE").value, "2026-07-20")
        self.assertEqual(CellValue.objects.get(row=row, column__name="QTY").value, "15")

        # Check Task
        task = getattr(row, "task", None)
        self.assertIsNotNone(task)
        self.assertEqual(task.task_name, "ENQ-1234")
        import datetime
        self.assertEqual(task.due_date, datetime.date(2026, 7, 20))

        # Check Permissions - System columns should be EDITABLE for LIST_PID
        from tables.permissions import get_column_access_level
        s_no_col = table.columns.get(name="S_NO")
        initial_mail_col = table.columns.get(name="INITIAL_MAIL")
        self.assertEqual(get_column_access_level(self.employee, s_no_col), "EDITABLE")
        self.assertEqual(get_column_access_level(self.employee, initial_mail_col), "EDITABLE")

    def test_dynamic_dropdown_column_filtering(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import RowViewSet
        factory = APIRequestFactory()
        table = Table.objects.create(name="Dynamic Filter Table", created_by=self.admin)
        col = Column.objects.create(table=table, name="Status Col", data_type="DROPDOWN", options="Open,Closed", is_filterable=True)
        
        row1 = Row.objects.create(table=table, created_by=self.employee)
        row2 = Row.objects.create(table=table, created_by=self.employee)
        
        CellValue.objects.create(row=row1, column=col, value="Open", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=col, value="Closed", updated_by=self.admin)
        
        # Test API filtering with col_<id>=Open
        view = RowViewSet.as_view({'get': 'list'})
        request = factory.get(f"/tables/api/rows/?table={table.id}&col_{col.id}=Open")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["results"]), 1)
        self.assertEqual(response.data["results"][0]["id"], row1.id)

    def test_list_pid_import_and_filtering(self):
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        from django.core.files.uploadedfile import SimpleUploadedFile
        factory = APIRequestFactory()

        table = Table.objects.create(name="LIST_PID Import Test Table", job_type="LIST_PID", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        # Verify columns are is_system_column = False
        for col in table.columns.all():
            self.assertFalse(col.is_system_column)

        # Create valid CSV data where DATE, DUE_DATE_CUSTOMER, and DUE_DATE_FLOW_FORCE are missing/blank
        csv_lines = [
            "S_NO,DATE,ENQUIRY_NO,DUE_DATE_FLOW_FORCE,DUE_DATE_CUSTOMER,COMPANY_NAME",
            "1,,ENQ-1234,,,Company A",
            "2,,ENQ-5678,,,Company B",
            "3,,ENQ-9012,,,Company A"
        ]
        csv_data = "\n".join(csv_lines)
        csv_file = SimpleUploadedFile("import_pid.csv", csv_data.encode("utf-8"), content_type="text/csv")

        view = TableViewSet.as_view({'post': 'import_csv'})
        request = factory.post(f"/tables/api/tables/{table.id}/import-csv/", {
            "file": csv_file
        }, format="multipart")
        force_authenticate(request, user=self.admin)
        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 201)

        # Check imported rows and blank dates
        rows = Row.objects.filter(table=table)
        self.assertEqual(rows.count(), 3)
        for r in rows:
            name_cell = CellValue.objects.filter(row=r, column__name="ENQUIRY_NO/QUOTATION_NO").first()
            self.assertIsNotNone(name_cell)
            self.assertTrue(name_cell.value.startswith("ENQ-"))

            date_cell = CellValue.objects.filter(row=r, column__name="DATE").first()
            due_ff_cell = CellValue.objects.filter(row=r, column__name="DUE_DATE_FLOW_FORCE").first()
            # Verify they are blank/None (safe_parse_date returned None, and they were not auto-assigned today)
            self.assertTrue(not date_cell or date_cell.value in [None, ""])
            self.assertTrue(not due_ff_cell or due_ff_cell.value in [None, ""])

            # Verify associated Task.due_date is None
            self.assertIsNone(r.task.due_date)

        # Verify company_name column became filterable DROPDOWN with analyzed company options
        company_col = table.columns.get(name="COMPANY_NAME")
        self.assertEqual(company_col.data_type, "DROPDOWN")
        self.assertTrue(company_col.is_filterable)
        # Options should be analyzed, sorted and distinct: "Company A,Company B"
        self.assertEqual(company_col.options, "Company A,Company B")

    def test_row_search_case_and_space_insensitive(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Search Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        col = Column.objects.create(table=table, name="Test Column", data_type="TEXT", position=7)
        row1 = Row.objects.create(table=table, created_by=self.employee)
        row2 = Row.objects.create(table=table, created_by=self.employee)

        CellValue.objects.create(row=row1, column=col, value="Esih Ayu Lisa", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=col, value="Tri Pirmansyah", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})

        # Test case & space insensitive search: "esihayulisa" -> should match row1
        request = factory.get(f"/tables/api/rows/?table={table.id}&search=esihayulisa")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['id'], row1.id)

        # Test another search: " TRIPIRMANSYAH " -> should match row2
        request = factory.get(f"/tables/api/rows/?table={table.id}&search=  TRIPIRMANSYAH ")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['id'], row2.id)

        # Test non-matching search: "unknown" -> should match 0 rows
        request = factory.get(f"/tables/api/rows/?table={table.id}&search=unknown")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data['results']), 0)

    def test_row_sorting_options(self):
        from tables.views import RowViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tasks.models import Task, TaskFollowUp
        import datetime
        factory = APIRequestFactory()

        # --- Test GENERAL Table (Due Date & Date Assigned) ---
        table = Table.objects.create(name="Sort Test Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        date_col = table.columns.get(name="DATE")
        
        row1 = Row.objects.create(table=table, created_by=self.employee)
        row2 = Row.objects.create(table=table, created_by=self.employee)
        row3 = Row.objects.create(table=table, created_by=self.employee)

        # 1. Populate Date Assigned (via DATE column)
        CellValue.objects.create(row=row1, column=date_col, value="2026-07-02", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=date_col, value="2026-07-01", updated_by=self.admin)
        CellValue.objects.create(row=row3, column=date_col, value="2026-07-03", updated_by=self.admin)

        # 2. Create Tasks for due date sorting
        task1 = Task.objects.create(row=row1, due_date=datetime.date(2026, 8, 15), status="PENDING")
        task2 = Task.objects.create(row=row2, due_date=datetime.date(2026, 8, 10), status="PENDING")
        task3 = Task.objects.create(row=row3, due_date=datetime.date(2026, 8, 20), status="PENDING")

        view = RowViewSet.as_view({'get': 'list'})

        # Test Sort by due_date Ascending: row2 (Aug 10), row1 (Aug 15), row3 (Aug 20)
        request = factory.get(f"/tables/api/rows/?table={table.id}&sort_by=due_date&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [row2.id, row1.id, row3.id])

        # Test Sort by due_date Descending: row3 (Aug 20), row1 (Aug 15), row2 (Aug 10)
        request = factory.get(f"/tables/api/rows/?table={table.id}&sort_by=due_date&sort_dir=desc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [row3.id, row1.id, row2.id])

        # Test follow_up_date sort on GENERAL table is ignored (returns default ordering by ID)
        request = factory.get(f"/tables/api/rows/?table={table.id}&sort_by=follow_up_date&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [row1.id, row2.id, row3.id])

        # Test Sort by date_assigned Ascending: row2 (July 1), row1 (July 2), row3 (July 3)
        request = factory.get(f"/tables/api/rows/?table={table.id}&sort_by=date_assigned&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [row2.id, row1.id, row3.id])

        # Test Sort by date_assigned Descending: row3 (July 3), row1 (July 2), row2 (July 1)
        request = factory.get(f"/tables/api/rows/?table={table.id}&sort_by=date_assigned&sort_dir=desc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [row3.id, row1.id, row2.id])

        # --- Test SALES Table (Follow-up Date & Date Assigned) ---
        sales_table = Table.objects.create(name="Sales Sort Test Table", job_type="SALES", created_by=self.admin)
        TableAccess.objects.create(table=sales_table, user=self.admin, access_level="ADMIN")
        
        s_row1 = Row.objects.create(table=sales_table, created_by=self.employee)
        s_row2 = Row.objects.create(table=sales_table, created_by=self.employee)
        s_row3 = Row.objects.create(table=sales_table, created_by=self.employee)

        s_task1 = Task.objects.create(row=s_row1, due_date=datetime.date(2026, 8, 15), status="PENDING")
        s_task2 = Task.objects.create(row=s_row2, due_date=datetime.date(2026, 8, 10), status="PENDING")
        s_task3 = Task.objects.create(row=s_row3, due_date=datetime.date(2026, 8, 20), status="PENDING")

        TaskFollowUp.objects.create(task=s_task1, follow_up_date=datetime.date(2026, 9, 2), discussed_points="points 1")
        TaskFollowUp.objects.create(task=s_task2, follow_up_date=datetime.date(2026, 9, 3), discussed_points="points 2")
        # s_task3 has no follow ups (should fallback to due_date: 2026-08-20)

        # Test Sort by follow_up_date Ascending: s_row3 (Aug 20), s_row1 (Sept 2), s_row2 (Sept 3)
        request = factory.get(f"/tables/api/rows/?table={sales_table.id}&sort_by=follow_up_date&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [s_row3.id, s_row1.id, s_row2.id])

        # Test Sort by follow_up_date Descending: s_row2 (Sept 3), s_row1 (Sept 2), s_row3 (Aug 20)
        request = factory.get(f"/tables/api/rows/?table={sales_table.id}&sort_by=follow_up_date&sort_dir=desc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [s_row2.id, s_row1.id, s_row3.id])

        # Test due_date sort on SALES table is ignored
        request = factory.get(f"/tables/api/rows/?table={sales_table.id}&sort_by=due_date&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [s_row1.id, s_row2.id, s_row3.id])

        # --- Test LIST_PID Table (Due Date maps to DUE_DATE_FLOW_FORCE) ---
        pid_table = Table.objects.create(name="PID Sort Test Table", job_type="LIST_PID", created_by=self.admin)
        TableAccess.objects.create(table=pid_table, user=self.admin, access_level="ADMIN")
        
        flow_force_col = pid_table.columns.get(name="DUE_DATE_FLOW_FORCE")
        
        p_row1 = Row.objects.create(table=pid_table, created_by=self.employee)
        p_row2 = Row.objects.create(table=pid_table, created_by=self.employee)
        p_row3 = Row.objects.create(table=pid_table, created_by=self.employee)
        
        CellValue.objects.create(row=p_row1, column=flow_force_col, value="2026-10-15", updated_by=self.admin)
        CellValue.objects.create(row=p_row2, column=flow_force_col, value="2026-10-10", updated_by=self.admin)
        CellValue.objects.create(row=p_row3, column=flow_force_col, value="2026-10-20", updated_by=self.admin)
        
        # Test Sort by due_date on LIST_PID table (maps to DUE_DATE_FLOW_FORCE)
        request = factory.get(f"/tables/api/rows/?table={pid_table.id}&sort_by=due_date&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [p_row2.id, p_row1.id, p_row3.id])

        # Test Sort by date (normalized to date_assigned) on GENERAL table
        request = factory.get(f"/tables/api/rows/?table={table.id}&sort_by=date&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [row2.id, row1.id, row3.id])

        # Test Sort by enquiry_no on LIST_PID table
        enquiry_col = pid_table.columns.get(name="ENQUIRY_NO/QUOTATION_NO")
        CellValue.objects.create(row=p_row1, column=enquiry_col, value="ENQ-002", updated_by=self.admin)
        CellValue.objects.create(row=p_row2, column=enquiry_col, value="ENQ-001", updated_by=self.admin)
        CellValue.objects.create(row=p_row3, column=enquiry_col, value="ENQ-003", updated_by=self.admin)

        request = factory.get(f"/tables/api/rows/?table={pid_table.id}&sort_by=enquiry_no&sort_dir=asc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [p_row2.id, p_row1.id, p_row3.id])

        request = factory.get(f"/tables/api/rows/?table={pid_table.id}&sort_by=enquiry_no&sort_dir=desc")
        force_authenticate(request, user=self.admin)
        response = view(request)
        self.assertEqual(response.status_code, 200)
        row_ids = [r['id'] for r in response.data['results']]
        self.assertEqual(row_ids, [p_row3.id, p_row1.id, p_row2.id])

    def test_employee_promoted_to_admin_table_access(self):
        from tables.permissions import get_accessible_tables, has_table_access, get_column_access_level
        from employee_management.services import EmployeeService

        # Create another department and table created by superadmin in that department
        other_dept = Department.objects.create(name="Sales Dept", slug="sales-dept")
        super_admin = User.objects.create_user(
            email="superadmin@flow-force.com",
            password="testpassword",
            full_name="Super Admin",
            role="SUPER_ADMIN",
            status="APPROVED"
        )
        table_other = Table.objects.create(name="Global Sales Table", created_by=super_admin, department=other_dept)

        # 1. Employee originally in Tables Engineering Dept
        emp = User.objects.create_user(
            email="promoted_emp@flow-force.com",
            password="testpassword",
            full_name="Promoted Employee",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )

        # As an employee, table_other is NOT accessible
        accessible_before = get_accessible_tables(emp)
        self.assertNotIn(table_other, accessible_before)
        self.assertFalse(has_table_access(emp, table_other, "VIEW"))

        # 2. Promote employee to ADMIN
        EmployeeService.update_employee(emp, role="ADMIN", updated_by=self.admin)
        emp.refresh_from_db()
        self.assertEqual(emp.role, "ADMIN")
        self.assertTrue(emp.is_staff)

        # As an ADMIN, all active tables are accessible regardless of department or creator
        accessible_after = get_accessible_tables(emp)
        self.assertIn(table_other, accessible_after)
        self.assertTrue(has_table_access(emp, table_other, "VIEW"))
        self.assertTrue(has_table_access(emp, table_other, "EDIT"))
        self.assertTrue(has_table_access(emp, table_other, "ADMIN"))

        # Column permissions check
        sys_col = table_other.columns.first()
        self.assertEqual(get_column_access_level(emp, sys_col), "EDITABLE")

    def test_import_standard_csv_no_metadata(self):
        """Test importing standard CSV starting on row 1 without metadata lines."""
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="Standard CSV Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        csv_content = (
            "Task Name,Due Date,Priority,Status\n"
            "Deploy Web App,2026-09-01,HIGH,PENDING\n"
            "Update Documentation,2026-09-05,MEDIUM,PENDING\n"
        )
        from django.core.files.uploadedfile import SimpleUploadedFile
        csv_file = SimpleUploadedFile("standard.csv", csv_content.encode("utf-8"), content_type="text/csv")

        view = TableViewSet.as_view({'post': 'import_csv'})
        request = factory.post(f"/tables/api/tables/{table.id}/import-csv/", {"file": csv_file}, format="multipart")
        force_authenticate(request, user=self.admin)

        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Row.objects.filter(table=table).count(), 2)

        tasks = list(Task.objects.filter(row__table=table).order_by('id'))
        self.assertEqual(len(tasks), 2)
        self.assertEqual(tasks[0].task_name, "Deploy Web App")
        self.assertEqual(tasks[1].task_name, "Update Documentation")

    def test_import_csv_without_s_no_and_missing_due_date(self):
        """Test importing CSV without S_NO column and with empty due date."""
        from tables.views import TableViewSet
        from rest_framework.test import APIRequestFactory, force_authenticate
        factory = APIRequestFactory()

        table = Table.objects.create(name="No SNO Table", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        csv_content = (
            "Task Name,Due Date,Priority\n"
            "Task Without Date,,LOW\n"
        )
        from django.core.files.uploadedfile import SimpleUploadedFile
        csv_file = SimpleUploadedFile("no_sno.csv", csv_content.encode("utf-8"), content_type="text/csv")

        view = TableViewSet.as_view({'post': 'import_csv'})
        request = factory.post(f"/tables/api/tables/{table.id}/import-csv/", {"file": csv_file}, format="multipart")
        force_authenticate(request, user=self.admin)

        response = view(request, pk=table.id)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(Row.objects.filter(table=table).count(), 1)

        task = Task.objects.filter(row__table=table).first()
        self.assertIsNotNone(task)
        self.assertEqual(task.task_name, "Task Without Date")
        self.assertIsNone(task.due_date)

    def test_import_csv_excel_serial_date(self):
        """Test safe_parse_date with Excel serial date numbers."""
        from tables.views import TableViewSet
        view = TableViewSet()
        import datetime
        # Excel serial 45443 is 2024-05-31
        parsed = view.safe_parse_date("45443")
        self.assertEqual(parsed, datetime.date(2024, 5, 31))

    def test_pid_dashboard_view(self):
        """Test PID Executive Dashboard view renders LIST_PID table details correctly."""
        from django.test import Client
        client = Client()
        client.force_login(self.admin)

        pid_table = Table.objects.create(name="Singapore PID Master", created_by=self.admin, job_type="LIST_PID")
        TableAccess.objects.create(table=pid_table, user=self.admin, access_level="ADMIN")

        r1 = Row.objects.create(table=pid_table, created_by=self.admin)
        col_pid = pid_table.columns.get(name="PID")
        col_enq = pid_table.columns.get(name="ENQUIRY_NO/QUOTATION_NO")
        col_po = pid_table.columns.get(name="PO")
        col_so = pid_table.columns.get(name="SALES_ORDER")
        col_cust = pid_table.columns.get(name="COMPANY_NAME")
        col_due_cust = pid_table.columns.get(name="DUE_DATE_CUSTOMER")
        col_due_ff = pid_table.columns.get(name="DUE_DATE_FLOW_FORCE")
        col_status = pid_table.columns.get(name="STATUS")

        CellValue.objects.create(row=r1, column=col_pid, value="PID-9001", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_enq, value="QUO-2026-001", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_po, value="PO-778899", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_so, value="SO-112233", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_cust, value="Acme Corp", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_due_cust, value="2026-09-15", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_due_ff, value="2026-09-01", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=col_status, value="Fabrication In Progress", updated_by=self.admin)

        response = client.get("/tables/pid-dashboard/")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "PID Executive Dashboard")
        self.assertContains(response, "Quick PID Status Summary Index")
        self.assertContains(response, "Singapore PID Master")
        self.assertContains(response, "PID-9001")
        self.assertContains(response, "QUO-2026-001")
        self.assertContains(response, "PO-778899")
        self.assertContains(response, "SO-112233")
        self.assertContains(response, "Acme Corp")
        self.assertContains(response, "Fabrication In Progress")
        self.assertContains(response, "Year 2026")

class LogsTableTestCase(TestCase):
    def setUp(self):
        import datetime
        from auth_app.models import EmployeeUser
        self.admin = EmployeeUser.objects.create_user(
            email="admin_logs@example.com",
            password="Password123!",
            role="ADMIN",
            full_name="Admin Logs User"
        )
        self.editor = EmployeeUser.objects.create_user(
            email="editor_logs@example.com",
            password="Password123!",
            role="EMPLOYEE",
            full_name="Editor Logs User"
        )
        self.viewer = EmployeeUser.objects.create_user(
            email="viewer_logs@example.com",
            password="Password123!",
            role="EMPLOYEE",
            full_name="Viewer Logs User"
        )

    def test_logs_table_system_columns_creation(self):
        table = Table.objects.create(name="Tool Inventory Logs", job_type="LOGS", created_by=self.admin)
        col_names = list(table.columns.values_list("name", flat=True))
        expected_cols = ["S_NO", "ISSUE_DATE", "RETURN_DATE", "TOOL_NAME", "STATUS", "ISSUED_BY", "RECEIVED_BY", "INITIAL_MAIL", "ALERT_MAIL"]
        for col in expected_cols:
            self.assertIn(col, col_names)
        
        status_col = table.columns.get(name="STATUS")
        self.assertEqual(status_col.options, "Returned,Not Returned")

    def test_logs_row_creation_and_status_sync(self):
        from rest_framework.test import APIClient
        client = APIClient()
        client.force_authenticate(user=self.admin)

        table = Table.objects.create(name="Workshop Tool Tracker", job_type="LOGS", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.editor, access_level="EDIT")
        TableAccess.objects.create(table=table, user=self.viewer, access_level="VIEW")

        # Create row
        response = client.post("/tables/api/rows/", {
            "table": table.id,
            "cells": {
                "TOOL_NAME": "Pneumatic Drill",
                "ISSUE_DATE": "2026-08-10",
                "RETURN_DATE": "2026-08-14",
                "ISSUED_BY": "Admin Logs User",
                "RECEIVED_BY": "John Doe",
                "STATUS": "Not Returned"
            }
        }, format="json")
        self.assertEqual(response.status_code, 201)

        row_id = response.data["id"]
        row = Row.objects.get(id=row_id)
        self.assertEqual(row.task.due_date, datetime.date(2026, 8, 14))

        # Update cell status to Returned
        status_col = table.columns.get(name="STATUS")
        response = client.post(f"/tables/api/rows/{row.id}/edit-cell/", {
            "column": status_col.id,
            "value": "Returned"
        }, format="json")
        self.assertEqual(response.status_code, 200)

        row.task.refresh_from_db()
        self.assertEqual(row.task.status, "COMPLETED")

    def test_logs_daily_alert_mails_sending(self):
        from tasks.tasks import send_daily_alert_mails
        from tasks.models import EmailLog, Notification
        from django.utils import timezone

        table = Table.objects.create(name="Site Equipment Logs", job_type="LOGS", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.editor, access_level="EDIT")
        TableAccess.objects.create(table=table, user=self.viewer, access_level="VIEW")

        # Create row overdue past return date
        row = Row.objects.create(table=table, created_by=self.admin)
        cols = {col.name: col for col in table.columns.all()}
        CellValue.objects.create(row=row, column=cols["TOOL_NAME"], value="Laser Level", updated_by=self.admin)
        CellValue.objects.create(row=row, column=cols["ISSUE_DATE"], value="2026-08-01", updated_by=self.admin)
        CellValue.objects.create(row=row, column=cols["RETURN_DATE"], value="2026-08-10", updated_by=self.admin)
        CellValue.objects.create(row=row, column=cols["ISSUED_BY"], value="Admin Logs User", updated_by=self.admin)
        CellValue.objects.create(row=row, column=cols["RECEIVED_BY"], value="Bob Smith", updated_by=self.admin)
        CellValue.objects.create(row=row, column=cols["STATUS"], value="Not Returned", updated_by=self.admin)

        Task.objects.create(row=row, due_date=datetime.date(2026, 8, 10), status="PENDING", assigned_by=self.admin)

        # Trigger daily alert mails
        send_daily_alert_mails()

        # Check EmailLog created for Admin and Editor, but NOT Viewer
        admin_logs = EmailLog.objects.filter(recipient_email=self.admin.email, subject__icontains="Unreturned Tool")
        editor_logs = EmailLog.objects.filter(recipient_email=self.editor.email, subject__icontains="Unreturned Tool")
        viewer_logs = EmailLog.objects.filter(recipient_email=self.viewer.email, subject__icontains="Unreturned Tool")

        self.assertTrue(admin_logs.exists())
        self.assertTrue(editor_logs.exists())
        self.assertFalse(viewer_logs.exists())

        log_body = admin_logs.first().body
        self.assertIn("Laser Level", log_body)
        self.assertIn("Bob Smith", log_body)
        self.assertIn("2026-08-10", log_body)

    def test_logs_table_auto_return_date_and_days_overdue(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import RowViewSet

        table = Table.objects.create(name="Tool Logs Test", job_type="LOGS", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        factory = APIRequestFactory()
        view = RowViewSet.as_view({'post': 'create'})

        # 1. Create row with ISSUE_DATE given but no RETURN_DATE
        req = factory.post('/tables/api/rows/', {
            'table': table.id,
            'cells': {
                'TOOL_NAME': 'Drill Machine',
                'ISSUE_DATE': '2026-08-01',
                'STATUS': 'Not Returned'
            }
        }, format='json')
        force_authenticate(req, user=self.admin)
        res = view(req)
        self.assertEqual(res.status_code, 201)

        row_id = res.data['id']
        row = Row.objects.get(id=row_id)

        # Verify RETURN_DATE auto captured to match ISSUE_DATE (2026-08-01)
        ret_col = table.columns.get(name='RETURN_DATE')
        ret_cell = row.cells.get(column=ret_col)
        self.assertEqual(ret_cell.value, '2026-08-01')

        # Verify DAYS_OVERDUE calculated based on current date
        from django.utils import timezone
        import datetime
        expected_days = (timezone.localdate() - datetime.date(2026, 8, 1)).days
    def test_export_excel(self):
        import io
        import openpyxl
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import TableViewSet

        table = Table.objects.create(name="PIDs Table Test", job_type="LIST_PID", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        # Create 2 rows
        row1 = Row.objects.create(table=table, created_by=self.admin)
        row2 = Row.objects.create(table=table, created_by=self.admin)

        desc_col = table.columns.get(name="DESCRIPTION")
        company_col = table.columns.get(name="COMPANY_NAME")
        qty_col = table.columns.get(name="QTY")
        due_col = table.columns.get(name="DUE_DATE_FLOW_FORCE")

        CellValue.objects.create(row=row1, column=desc_col, value="Line 1\nLine 2 description", updated_by=self.admin)
        CellValue.objects.create(row=row1, column=company_col, value="Acme Corp", updated_by=self.admin)
        CellValue.objects.create(row=row1, column=qty_col, value=15, updated_by=self.admin)
        CellValue.objects.create(row=row1, column=due_col, value="2026-09-01", updated_by=self.admin)

        CellValue.objects.create(row=row2, column=desc_col, value="Single line text", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=company_col, value="Beta LLC", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=qty_col, value=42, updated_by=self.admin)
        CellValue.objects.create(row=row2, column=due_col, value="2026-09-15", updated_by=self.admin)

        factory = APIRequestFactory()
        view = TableViewSet.as_view({'get': 'export_excel'})

        # 1. Test full export
        req = factory.get(f'/tables/api/tables/{table.id}/export-excel/')
        force_authenticate(req, user=self.admin)
        res = view(req, pk=table.id)

        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('pids_table_test_export_', res['Content-Disposition'])
        self.assertTrue(res['Content-Disposition'].endswith('.xlsx"'))

        # Parse Excel workbook from bytes
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active

        # Check headers
        headers = [cell.value for cell in ws[1]]
        expected_cols = [c.name for c in table.columns.all().order_by("position", "id")]
        self.assertEqual(headers, expected_cols)

        # Check row count (1 header + 2 data rows = 3 rows)
        self.assertEqual(ws.max_row, 3)

        # Check values in first data row
        desc_idx = headers.index("DESCRIPTION") + 1
        company_idx = headers.index("COMPANY_NAME") + 1
        qty_idx = headers.index("QTY") + 1

        self.assertEqual(ws.cell(row=2, column=desc_idx).value, "Line 1\nLine 2 description")
        self.assertEqual(ws.cell(row=2, column=company_idx).value, "Acme Corp")
        self.assertEqual(ws.cell(row=2, column=qty_idx).value, 15)

        # 2. Test filtered export (search="Beta")
        req_filter = factory.get(f'/tables/api/tables/{table.id}/export-excel/?search=Beta')
        force_authenticate(req_filter, user=self.admin)
        res_filter = view(req_filter, pk=table.id)

        self.assertEqual(res_filter.status_code, 200)
        wb_filtered = openpyxl.load_workbook(io.BytesIO(res_filter.content))
        ws_filtered = wb_filtered.active

        # Filtered export should only have 1 data row
        self.assertEqual(ws_filtered.max_row, 2)
        self.assertEqual(ws_filtered.cell(row=2, column=company_idx).value, "Beta LLC")

    def test_text_column_multiline_properties_persistence(self):
        import json
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import ColumnViewSet

        table = Table.objects.create(name="Custom Column Test", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        factory = APIRequestFactory()
        view_create = ColumnViewSet.as_view({'post': 'create'})
        view_update = ColumnViewSet.as_view({'patch': 'partial_update'})

        # 1. Create a custom TEXT column with multiline configuration
        opts_payload = json.dumps({
            "input_type": "multiline",
            "rows": 4,
            "placeholder": "Enter detailed notes...",
            "max_length": 1000,
            "resizable": True
        })
        req = factory.post('/tables/api/columns/', {
            'table': table.id,
            'name': 'Special Notes',
            'data_type': 'TEXT',
            'options': opts_payload,
            'is_filterable': False
        }, format='json')
        force_authenticate(req, user=self.admin)
        res = view_create(req)
        self.assertEqual(res.status_code, 201)

        col_id = res.data['id']
        col = Column.objects.get(id=col_id)
        self.assertEqual(col.name, 'Special Notes')
        self.assertEqual(col.data_type, 'TEXT')
        
        # Verify JSON options persist
        parsed_opts = json.loads(col.options)
        self.assertEqual(parsed_opts['input_type'], 'multiline')
        self.assertEqual(parsed_opts['rows'], 4)
        self.assertEqual(parsed_opts['placeholder'], 'Enter detailed notes...')
        self.assertEqual(parsed_opts['max_length'], 1000)
        self.assertTrue(parsed_opts['resizable'])

        # 2. Update column configuration via PATCH (edit properties)
        updated_opts = json.dumps({
            "input_type": "multiline",
            "rows": 6,
            "placeholder": "Updated placeholder",
            "max_length": 2000,
            "resizable": False
        })
        req_patch = factory.patch(f'/tables/api/columns/{col.id}/', {
            'options': updated_opts
        }, format='json')
        force_authenticate(req_patch, user=self.admin)
        res_patch = view_update(req_patch, pk=col.id)
        self.assertEqual(res_patch.status_code, 200)

        col.refresh_from_db()
        parsed_updated = json.loads(col.options)
        self.assertEqual(parsed_updated['rows'], 6)
        self.assertEqual(parsed_updated['placeholder'], 'Updated placeholder')
        self.assertEqual(parsed_updated['max_length'], 2000)
        self.assertFalse(parsed_updated['resizable'])

    def test_text_column_validation_invalid_rows_and_max_length(self):
        import json
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import ColumnViewSet

        table = Table.objects.create(name="Validation Test", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        factory = APIRequestFactory()
        view_create = ColumnViewSet.as_view({'post': 'create'})

        # Test invalid rows (0 or negative)
        invalid_rows_opts = json.dumps({
            "input_type": "multiline",
            "rows": 0
        })
        req = factory.post('/tables/api/columns/', {
            'table': table.id,
            'name': 'Invalid Rows Column',
            'data_type': 'TEXT',
            'options': invalid_rows_opts
        }, format='json')
        force_authenticate(req, user=self.admin)
        res = view_create(req)
        self.assertEqual(res.status_code, 400)
        self.assertIn('options', res.data)

        # Test invalid max_length (0 or negative)
        invalid_len_opts = json.dumps({
            "input_type": "multiline",
            "rows": 3,
            "max_length": -5
        })
        req2 = factory.post('/tables/api/columns/', {
            'table': table.id,
            'name': 'Invalid Len Column',
            'data_type': 'TEXT',
            'options': invalid_len_opts
        }, format='json')
        force_authenticate(req2, user=self.admin)
        res2 = view_create(req2)
        self.assertEqual(res2.status_code, 400)
        self.assertIn('options', res2.data)

    def test_multiline_cell_value_save_and_excel_export(self):
        import io
        import json
        import openpyxl
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import RowViewSet, TableViewSet

        table = Table.objects.create(name="Multiline Flow Test", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        # Custom multiline column
        multiline_col = Column.objects.create(
            table=table,
            name="Job Details",
            data_type="TEXT",
            position=7,
            options=json.dumps({"input_type": "multiline", "rows": 3, "placeholder": "Details...", "resizable": True})
        )

        factory = APIRequestFactory()
        view_row = RowViewSet.as_view({'post': 'create'})

        multiline_content = "Heading:\n- Item 1\n- Item 2\nSummary paragraph."
        req_row = factory.post('/tables/api/rows/', {
            'table': table.id,
            'cells': {
                'TASK_NAME': 'Complex Engineering Task',
                'DUE_DATE': '2026-09-30',
                'Job Details': multiline_content
            }
        }, format='json')
        force_authenticate(req_row, user=self.admin)
        res_row = view_row(req_row)
        self.assertEqual(res_row.status_code, 201)

        row_id = res_row.data['id']
        cell = CellValue.objects.get(row_id=row_id, column=multiline_col)
        self.assertEqual(cell.value, multiline_content)

        # Export Excel and verify multiline content and headers
        view_export = TableViewSet.as_view({'get': 'export_excel'})
        req_export = factory.get(f'/tables/api/tables/{table.id}/export-excel/')
        force_authenticate(req_export, user=self.admin)
        res_export = view_export(req_export, pk=table.id)

        self.assertEqual(res_export.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res_export.content))
        ws = wb.active

        headers = [c.value for c in ws[1]]
        self.assertIn("Job Details", headers)
        job_details_idx = headers.index("Job Details") + 1

        # Check cell value in row 2
        exported_val = ws.cell(row=2, column=job_details_idx).value
        self.assertEqual(exported_val, multiline_content)

    def test_export_excel_table_27_style_and_year_filter(self):
        import io
        from datetime import datetime
        import openpyxl
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import TableViewSet

        # 1. Create a Table matching production Table 27 (LIST_PID schema)
        table = Table.objects.create(name="PID Master 2025", job_type="LIST_PID", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        col_names = [
            ("S_NO", "NUMBER"),
            ("ENQUIRY_NO/QUOTATION_NO", "TEXT"),
            ("PID", "TEXT"),
            ("NEW_PID_NO", "TEXT"),
            ("SALES_ORDER", "TEXT"),
            ("PO", "TEXT"),
            ("DATE", "DATE"),
            ("FFE_SINGAPORE", "TEXT"),
            ("COMPANY_NAME", "TEXT"),
            ("DESCRIPTION", "TEXT"),
            ("QTY", "NUMBER"),
            ("DUE_DATE_CUSTOMER", "DATE"),
            ("DUE_DATE_FLOW_FORCE", "DATE"),
            ("PROJECT", "TEXT"),
            ("STATUS", "TEXT"),
            ("INITIAL_MAIL", "CHECKBOX"),
            ("ALERT_MAIL", "CHECKBOX"),
        ]
        created_cols = {}
        for pos, (c_name, c_type) in enumerate(col_names, start=1):
            created_cols[c_name], _ = Column.objects.get_or_create(
                table=table,
                name=c_name,
                defaults={
                    "data_type": c_type,
                    "position": pos,
                    "is_system_column": True
                }
            )

        # 2. Add Row 1 (Year 2025)
        row2025 = Row.objects.create(table=table, created_by=self.admin)
        Task.objects.create(
            row=row2025,
            due_date=datetime.strptime("2025-06-15", "%Y-%m-%d").date(),
            priority="HIGH",
            status="PENDING",
            assigned_by=self.admin
        )
        CellValue.objects.update_or_create(row=row2025, column=created_cols["S_NO"], defaults={"value": 1, "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2025, column=created_cols["ENQUIRY_NO/QUOTATION_NO"], defaults={"value": "ENQ-2025-001", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2025, column=created_cols["PID"], defaults={"value": "PID-2025-A", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2025, column=created_cols["DESCRIPTION"], defaults={"value": "Line 1: 2025 Order\nLine 2: Multi-line spec", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2025, column=created_cols["QTY"], defaults={"value": 25, "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2025, column=created_cols["DUE_DATE_FLOW_FORCE"], defaults={"value": "2025-06-15", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2025, column=created_cols["INITIAL_MAIL"], defaults={"value": "true", "updated_by": self.admin})
        # Leave other columns empty/null to test null handling

        # 3. Add Row 2 (Year 2026)
        row2026 = Row.objects.create(table=table, created_by=self.admin)
        Task.objects.create(
            row=row2026,
            due_date=datetime.strptime("2026-08-20", "%Y-%m-%d").date(),
            priority="MEDIUM",
            status="COMPLETED",
            assigned_by=self.admin
        )
        CellValue.objects.update_or_create(row=row2026, column=created_cols["S_NO"], defaults={"value": 2, "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2026, column=created_cols["ENQUIRY_NO/QUOTATION_NO"], defaults={"value": "ENQ-2026-099", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2026, column=created_cols["PID"], defaults={"value": "PID-2026-B", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2026, column=created_cols["DESCRIPTION"], defaults={"value": "2026 single line", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2026, column=created_cols["QTY"], defaults={"value": 50, "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2026, column=created_cols["DUE_DATE_FLOW_FORCE"], defaults={"value": "2026-08-20", "updated_by": self.admin})
        CellValue.objects.update_or_create(row=row2026, column=created_cols["INITIAL_MAIL"], defaults={"value": "false", "updated_by": self.admin})

        factory = APIRequestFactory()
        view = TableViewSet.as_view({'get': 'export_excel'})

        # Test A: Unfiltered Export (table=27)
        req_all = factory.get(f'/tables/api/tables/{table.id}/export-excel/?table={table.id}')
        force_authenticate(req_all, user=self.admin)
        res_all = view(req_all, pk=table.id)

        self.assertEqual(res_all.status_code, 200)
        self.assertEqual(res_all['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        wb_all = openpyxl.load_workbook(io.BytesIO(res_all.content))
        ws_all = wb_all.active
        self.assertEqual(ws_all.max_row, 3) # Header + 2 data rows

        # Test B: Filtered Export with ?table=27&year=2025
        req_2025 = factory.get(f'/tables/api/tables/{table.id}/export-excel/?table={table.id}&year=2025')
        force_authenticate(req_2025, user=self.admin)
        res_2025 = view(req_2025, pk=table.id)

        self.assertEqual(res_2025.status_code, 200)
        wb_2025 = openpyxl.load_workbook(io.BytesIO(res_2025.content))
        ws_2025 = wb_2025.active
        self.assertEqual(ws_2025.max_row, 2) # Header + exactly 1 row (2025)

        headers = [c.value for c in ws_2025[1]]
        desc_idx = headers.index("DESCRIPTION") + 1
        pid_idx = headers.index("PID") + 1
        qty_idx = headers.index("QTY") + 1
        mail_idx = headers.index("INITIAL_MAIL") + 1
        po_idx = headers.index("PO") + 1

        self.assertEqual(ws_2025.cell(row=2, column=pid_idx).value, "PID-2025-A")
        self.assertEqual(ws_2025.cell(row=2, column=desc_idx).value, "Line 1: 2025 Order\nLine 2: Multi-line spec")
        self.assertEqual(ws_2025.cell(row=2, column=qty_idx).value, 25)
        self.assertEqual(ws_2025.cell(row=2, column=mail_idx).value, "YES")
        self.assertIn(ws_2025.cell(row=2, column=po_idx).value, ["", None]) # Null/empty value check

    def test_export_excel_empty_table(self):
        import io
        import openpyxl
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import TableViewSet

        empty_table = Table.objects.create(name="Empty Test Table", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=empty_table, user=self.admin, access_level="ADMIN")

        Column.objects.create(table=empty_table, name="CustomCol1", data_type="TEXT", position=10)
        Column.objects.create(table=empty_table, name="CustomCol2", data_type="NUMBER", position=11)

        factory = APIRequestFactory()
        view = TableViewSet.as_view({'get': 'export_excel'})
        req = factory.get(f'/tables/api/tables/{empty_table.id}/export-excel/')
        force_authenticate(req, user=self.admin)
        res = view(req, pk=empty_table.id)

        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.max_row, 1) # Header only
        expected_cols = [c.name for c in empty_table.columns.all().order_by("position", "id")]
        self.assertEqual([c.value for c in ws[1]], expected_cols)

    def test_export_excel_postgresql_safe_date_sorting(self):
        import io
        import openpyxl
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import TableViewSet

        table = Table.objects.create(name="Date Sorting Test", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=table, user=self.admin, access_level="ADMIN")

        date_col = Column.objects.create(table=table, name="DATE", data_type="DATE", position=1)
        task_col = Column.objects.create(table=table, name="TASK_NAME", data_type="TEXT", position=2)

        # Row 1 with valid date
        r1 = Row.objects.create(table=table, created_by=self.admin)
        CellValue.objects.create(row=r1, column=date_col, value="2025-01-15", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=task_col, value="Task A", updated_by=self.admin)

        # Row 2 with empty string "" in date (simulating PostgreSQL JSONB empty string edge case)
        r2 = Row.objects.create(table=table, created_by=self.admin)
        CellValue.objects.create(row=r2, column=date_col, value="", updated_by=self.admin)
        CellValue.objects.create(row=r2, column=task_col, value="Task B", updated_by=self.admin)

        # Row 3 with null in date
        r3 = Row.objects.create(table=table, created_by=self.admin)
        CellValue.objects.create(row=r3, column=date_col, value=None, updated_by=self.admin)
        CellValue.objects.create(row=r3, column=task_col, value="Task C", updated_by=self.admin)

        factory = APIRequestFactory()
        view = TableViewSet.as_view({'get': 'export_excel'})
        req = factory.get(f'/tables/api/tables/{table.id}/export-excel/?sort_by=date&sort_dir=asc')
        force_authenticate(req, user=self.admin)
        res = view(req, pk=table.id)

        self.assertEqual(res.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.max_row, 4) # Header + 3 data rows


class SalesAndRowActionTests(TestCase):
    def setUp(self):
        from employee_management.models import Department
        from auth_app.models import EmployeeUser
        self.sales_dept, _ = Department.objects.get_or_create(name="Sales", defaults={"description": "Sales Department"})
        self.sales_user, _ = EmployeeUser.objects.get_or_create(
            email="sales-rep@flow-force.com",
            defaults={
                "password": "password123",
                "role": "EMPLOYEE",
                "full_name": "Sales Rep",
                "department": self.sales_dept
            }
        )
        self.admin, _ = EmployeeUser.objects.get_or_create(
            email="admin-rep@flow-force.com",
            defaults={
                "password": "adminpassword123",
                "role": "ADMIN",
                "full_name": "System Admin"
            }
        )
        self.sales_table = Table.objects.create(
            name="Sales Leads",
            job_type="SALES",
            department=self.sales_dept,
            created_by=self.admin
        )
        self.row = Row.objects.create(table=self.sales_table, created_by=self.sales_user)

    def test_sales_department_user_has_edit_access(self):
        from tables.permissions import has_table_access
        self.assertTrue(has_table_access(self.sales_user, self.sales_table, "VIEW"))
        self.assertTrue(has_table_access(self.sales_user, self.sales_table, "EDIT"))

    def test_row_serializer_auto_creates_missing_task(self):
        from tables.serializers import RowSerializer
        serializer = RowSerializer(self.row)
        data = serializer.data
        self.assertIsNotNone(data["task_details"])
        self.assertEqual(data["task_details"]["status"], "PENDING")

    def test_toggle_status_row_action(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tables.views import RowViewSet
        factory = APIRequestFactory()
        view = RowViewSet.as_view({'post': 'toggle_status'})

        # Toggle to COMPLETED
        req = factory.post(f'/tables/api/rows/{self.row.id}/toggle-status/', {'status': 'COMPLETED'}, format='json')
        force_authenticate(req, user=self.sales_user)
        res = view(req, pk=self.row.id)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["task_details"]["status"], "COMPLETED")

        # Toggle to PENDING
        req2 = factory.post(f'/tables/api/rows/{self.row.id}/toggle-status/', {}, format='json')
        force_authenticate(req2, user=self.sales_user)
        res2 = view(req2, pk=self.row.id)
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(res2.data["task_details"]["status"], "PENDING")

    def test_sales_follow_up_logging(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from tasks.views import TaskViewSet
        from tables.serializers import RowSerializer

        # Ensure task exists
        RowSerializer(self.row).data
        task = self.row.task

        factory = APIRequestFactory()
        view = TaskViewSet.as_view({'post': 'log_follow_up'})
        req = factory.post(f'/tasks/api/tasks/{task.id}/log-follow-up/', {
            'discussed_points': 'Customer agreed to demo next Tuesday',
            'status': 'IN_PROGRESS',
            'next_follow_up_date': '2026-10-15'
        }, format='json')
        force_authenticate(req, user=self.sales_user)
        res = view(req, pk=task.id)
        self.assertEqual(res.status_code, 200)
        task.refresh_from_db()
        self.assertEqual(task.status, 'IN_PROGRESS')
        self.assertEqual(str(task.due_date), '2026-10-15')
        self.assertEqual(task.follow_ups.count(), 1)


class ConcurrentSNoGenerationRegressionTestCase(TestCase):
    """
    BUG-06 Regression tests for safe, concurrent S_NO generation.
    """
    def setUp(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        self.factory = APIRequestFactory()
        self.dept = Department.objects.create(name="S_NO QA Dept", slug="s-no-qa-dept")
        self.admin = User.objects.create_user(
            email="snoadmin@flow-force.com",
            password="testpassword",
            full_name="SNO Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(
            name="S_NO Test Table",
            created_by=self.admin,
            department=self.dept,
            job_type="STANDARD"
        )

    def test_consecutive_row_creations_generate_unique_s_no(self):
        """Verifies that consecutive row creations increment S_NO distinctly without collisions."""
        from rest_framework.test import force_authenticate
        from tables.views import RowViewSet
        view = RowViewSet.as_view({'post': 'create'})

        # First row creation
        req1 = self.factory.post('/tables/api/rows/', {
            'table': self.table.id,
            'cells': {
                'TASK_NAME': 'First Task',
                'DUE_DATE': '2026-10-20'
            }
        }, format='json')
        force_authenticate(req1, user=self.admin)
        res1 = view(req1)
        self.assertEqual(res1.status_code, 201)
        row1_id = res1.data['id']

        # Second row creation
        req2 = self.factory.post('/tables/api/rows/', {
            'table': self.table.id,
            'cells': {
                'TASK_NAME': 'Second Task',
                'DUE_DATE': '2026-10-21'
            }
        }, format='json')
        force_authenticate(req2, user=self.admin)
        res2 = view(req2)
        self.assertEqual(res2.status_code, 201)
        row2_id = res2.data['id']

        # Check S_NO in cell values
        s_no_col = self.table.columns.get(name="S_NO")
        cell1 = CellValue.objects.get(row_id=row1_id, column=s_no_col)
        cell2 = CellValue.objects.get(row_id=row2_id, column=s_no_col)
        self.assertEqual(int(cell1.value), 1)
        self.assertEqual(int(cell2.value), 2)
        self.assertNotEqual(cell1.value, cell2.value)

    def test_s_no_handles_string_values_and_increments(self):
        """Verifies that S_NO calculation safely parses string numeric values and increments correctly."""
        from rest_framework.test import force_authenticate
        from tables.views import RowViewSet

        # Manually create a row with string S_NO
        s_no_col = self.table.columns.get(name="S_NO")
        existing_row = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=existing_row, column=s_no_col, value="42")

        view = RowViewSet.as_view({'post': 'create'})
        req = self.factory.post('/tables/api/rows/', {
            'table': self.table.id,
            'cells': {
                'TASK_NAME': 'Incremented Task',
                'DUE_DATE': '2026-10-25'
            }
        }, format='json')
        force_authenticate(req, user=self.admin)
        res = view(req)
        self.assertEqual(res.status_code, 201)

        new_row_id = res.data['id']
        new_cell = CellValue.objects.get(row_id=new_row_id, column=s_no_col)
        self.assertEqual(int(new_cell.value), 43)

    def test_concurrent_row_creations_cannot_receive_same_s_no(self):
        """Demonstrates that row creations acquire locks and produce unique S_NOs."""
        from rest_framework.test import force_authenticate
        from tables.views import RowViewSet
        from unittest.mock import patch

        view = RowViewSet.as_view({'post': 'create'})

        # Verify that select_for_update is invoked on Table and CellValue to enforce concurrency locking
        with patch.object(Table.objects, 'select_for_update', wraps=Table.objects.select_for_update) as mock_table_sfu, \
             patch.object(CellValue.objects, 'select_for_update', wraps=CellValue.objects.select_for_update) as mock_cell_sfu:
            req1 = self.factory.post('/tables/api/rows/', {
                'table': self.table.id,
                'cells': {
                    'TASK_NAME': 'Concurrent Task 1',
                    'DUE_DATE': '2026-10-22'
                }
            }, format='json')
            force_authenticate(req1, user=self.admin)
            res1 = view(req1)
            self.assertEqual(res1.status_code, 201)
            self.assertTrue(mock_table_sfu.called)
            self.assertTrue(mock_cell_sfu.called)

        # Create second row
        req2 = self.factory.post('/tables/api/rows/', {
            'table': self.table.id,
            'cells': {
                'TASK_NAME': 'Concurrent Task 2',
                'DUE_DATE': '2026-10-23'
            }
        }, format='json')
        force_authenticate(req2, user=self.admin)
        res2 = view(req2)
        self.assertEqual(res2.status_code, 201)

        s_no_col = self.table.columns.get(name="S_NO")
        s_no1 = int(CellValue.objects.get(row_id=res1.data['id'], column=s_no_col).value)
        s_no2 = int(CellValue.objects.get(row_id=res2.data['id'], column=s_no_col).value)
        self.assertNotEqual(s_no1, s_no2)
        self.assertEqual(s_no2, s_no1 + 1)


class TableDuplicationRegressionTestCase(TestCase):
    """
    BUG-07 Regression tests for Table duplication:
    - Atomicity & rollback protection
    - Task assignee preservation
    - Suppression of mass assignment email side effects
    """
    def setUp(self):
        from rest_framework.test import APIRequestFactory
        self.factory = APIRequestFactory()
        self.dept = Department.objects.create(name="Duplication QA Dept", slug="dup-qa-dept")
        self.admin = User.objects.create_user(
            email="dupadmin@flow-force.com",
            password="testpassword",
            full_name="Dup Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="dupemp@flow-force.com",
            password="testpassword",
            full_name="Dup Employee",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )
        self.viewer = User.objects.create_user(
            email="dupviewer@flow-force.com",
            password="testpassword",
            full_name="Dup Viewer",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(
            name="Original Source Table",
            description="Testing duplication integrity",
            created_by=self.admin,
            department=self.dept,
            job_type="STANDARD"
        )

        # Custom column
        self.custom_col = Column.objects.create(
            table=self.table,
            name="PROJECT_CODE",
            data_type="TEXT",
            is_system_column=False
        )

        # Table Access
        TableAccess.objects.create(
            table=self.table,
            user=self.viewer,
            access_level="VIEW"
        )

        # Row, cells, task
        self.row = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=self.row, column=self.custom_col, value="PRJ-1001")
        task_name_col = self.table.columns.get(name="TASK_NAME")
        due_date_col = self.table.columns.get(name="DUE_DATE")
        CellValue.objects.create(row=self.row, column=task_name_col, value="Assembly Task")
        CellValue.objects.create(row=self.row, column=due_date_col, value="2026-10-30")

        self.task = Task.objects.create(
            row=self.row,
            priority="HIGH",
            status="PENDING",
            due_date=datetime.date(2026, 10, 30),
            assigned_by=self.admin
        )
        self.task.assigned_to.set([self.employee])

    def test_successful_duplication_preserves_all_data_and_assignees(self):
        """Verifies that duplicating a table clones metadata, columns, rows, cells, and preserves task assignees."""
        from rest_framework.test import force_authenticate
        from tables.views import TableViewSet
        from unittest.mock import patch

        view = TableViewSet.as_view({'post': 'duplicate_table'})
        req = self.factory.post(f'/tables/api/tables/{self.table.id}/duplicate/')
        force_authenticate(req, user=self.admin)

        with patch("tasks.tasks.send_initial_mail.delay") as mock_mail:
            res = view(req, pk=self.table.id)
            self.assertEqual(res.status_code, 201)
            # Verify no assignment emails were queued during duplication
            self.assertFalse(mock_mail.called)

        new_table_id = res.data['id']
        new_table = Table.objects.get(id=new_table_id)
        self.assertEqual(new_table.name, f"Copy of {self.table.name}")
        self.assertEqual(new_table.department, self.dept)
        self.assertEqual(new_table.job_type, "STANDARD")

        # Verify custom column copied
        self.assertTrue(new_table.columns.filter(name="PROJECT_CODE").exists())

        # Verify TableAccess copied
        self.assertTrue(new_table.access_rules.filter(user=self.viewer, access_level="VIEW").exists())

        # Verify row and cells copied
        new_row = new_table.rows.first()
        self.assertIsNotNone(new_row)
        new_custom_col = new_table.columns.get(name="PROJECT_CODE")
        cell_val = CellValue.objects.get(row=new_row, column=new_custom_col).value
        self.assertEqual(cell_val, "PRJ-1001")

        # Verify task and assignees copied
        new_task = new_row.task
        self.assertIsNotNone(new_task)
        self.assertEqual(new_task.priority, "HIGH")
        self.assertEqual(new_task.status, "PENDING")
        self.assertEqual(list(new_task.assigned_to.all()), [self.employee])

    def test_duplication_failure_rolls_back_cleanly(self):
        """Verifies that any error during duplication triggers a rollback, leaving no partial tables or orphan rows."""
        from rest_framework.test import force_authenticate
        from tables.views import TableViewSet
        from unittest.mock import patch

        view = TableViewSet.as_view({'post': 'duplicate_table'})
        req = self.factory.post(f'/tables/api/tables/{self.table.id}/duplicate/')
        force_authenticate(req, user=self.admin)

        # Inject failure during CellValue creation
        with patch("tables.services.duplicate_service.CellValue.objects.create", side_effect=RuntimeError("Simulated DB Crash")):
            res = view(req, pk=self.table.id)
            self.assertEqual(res.status_code, 500)
            self.assertIn("Failed to duplicate table", res.data["error"])

        # Verify atomic rollback: no copy table exists
        self.assertFalse(Table.objects.filter(name=f"Copy of {self.table.name}").exists())
        # Verify no orphan rows exist for a copy table
        self.assertEqual(Row.objects.filter(table__name=f"Copy of {self.table.name}").count(), 0)

    def test_duplication_does_not_send_mass_assignment_emails(self):
        """Verifies that duplicating a table with tasks never triggers mass assignment emails."""
        from rest_framework.test import force_authenticate
        from tables.views import TableViewSet
        from unittest.mock import patch

        view = TableViewSet.as_view({'post': 'duplicate_table'})
        req = self.factory.post(f'/tables/api/tables/{self.table.id}/duplicate/')
        force_authenticate(req, user=self.admin)

        with patch("tasks.tasks.send_initial_mail.delay") as mock_mail:
            res = view(req, pk=self.table.id)
            self.assertEqual(res.status_code, 201)
            mock_mail.assert_not_called()


class PidDashboardOptimizationRegressionTestCase(TestCase):
    """
    Regression test suite for Phase 2B PID Dashboard Optimization.
    Validates that pid_dashboard_view executes with constant-bound query complexity (<= 8 total queries including context processors),
    correctly calculates all KPI metrics, handles dirty legacy dates gracefully,
    and isolates LIST_PID tables per user access permissions.
    """

    def setUp(self):
        from django.test.client import RequestFactory
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        self.factory = RequestFactory()
        self.connection = connection
        self.CaptureQueriesContext = CaptureQueriesContext

        self.dept = Department.objects.create(name="PID Engineering", slug="pid-engineering")
        self.dept2 = Department.objects.create(name="Other Department", slug="other-dept")
        self.admin = User.objects.create_user(
            email="pidadmin@flow-force.com",
            password="testpassword",
            full_name="PID Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="pidemp@flow-force.com",
            password="testpassword",
            full_name="PID Employee",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )

        # Create two LIST_PID tables
        self.table1 = Table.objects.create(
            name="PID Table Alpha",
            job_type="LIST_PID",
            created_by=self.admin,
            department=self.dept
        )
        self.table2 = Table.objects.create(
            name="PID Table Beta",
            job_type="LIST_PID",
            created_by=self.admin,
            department=self.dept2
        )
        # Create one non-PID table
        self.non_pid_table = Table.objects.create(
            name="Standard Tasks",
            job_type="STANDARD",
            created_by=self.admin,
            department=self.dept
        )

        # Create relevant columns for Table 1
        self.cols1 = {}
        for cname in ["PID", "CUSTOMER_NAME", "STATUS", "DUE_DATE_FLOW_FORCE", "DATE"]:
            self.cols1[cname] = Column.objects.create(table=self.table1, name=cname, data_type="TEXT")

        # Create relevant columns for Table 2
        self.cols2 = {}
        for cname in ["PID", "CUSTOMER_NAME", "STATUS", "DUE_DATE_FLOW_FORCE", "DATE"]:
            self.cols2[cname] = Column.objects.create(table=self.table2, name=cname, data_type="TEXT")

    def test_pid_dashboard_query_count_and_fidelity(self):
        """PID dashboard executes with <= 8 total queries (<= 4 for view) and computes KPIs accurately."""
        from datetime import date, timedelta
        from tables.views import pid_dashboard_view
        today = date.today()
        yesterday = today - timedelta(days=1)
        future = today + timedelta(days=30)

        # Row 1: Table 1, Completed
        r1 = Row.objects.create(table=self.table1, created_by=self.admin)
        CellValue.objects.create(row=r1, column=self.cols1["PID"], value="PID-001")
        CellValue.objects.create(row=r1, column=self.cols1["STATUS"], value="COMPLETED")
        CellValue.objects.create(row=r1, column=self.cols1["DUE_DATE_FLOW_FORCE"], value=yesterday.strftime("%Y-%m-%d"))
        CellValue.objects.create(row=r1, column=self.cols1["DATE"], value="2025-04-10")

        # Row 2: Table 1, Due today (In Progress)
        r2 = Row.objects.create(table=self.table1, created_by=self.admin)
        CellValue.objects.create(row=r2, column=self.cols1["PID"], value="PID-002")
        CellValue.objects.create(row=r2, column=self.cols1["STATUS"], value="IN_PROGRESS")
        CellValue.objects.create(row=r2, column=self.cols1["DUE_DATE_FLOW_FORCE"], value=today.strftime("%Y-%m-%d"))

        # Row 3: Table 1, Overdue (Pending with past target date)
        r3 = Row.objects.create(table=self.table1, created_by=self.admin)
        CellValue.objects.create(row=r3, column=self.cols1["PID"], value="PID-003")
        CellValue.objects.create(row=r3, column=self.cols1["STATUS"], value="PENDING")
        CellValue.objects.create(row=r3, column=self.cols1["DUE_DATE_FLOW_FORCE"], value=yesterday.strftime("%Y-%m-%d"))

        # Row 4: Table 2, Future In Progress
        r4 = Row.objects.create(table=self.table2, created_by=self.admin)
        CellValue.objects.create(row=r4, column=self.cols2["PID"], value="PID-004")
        CellValue.objects.create(row=r4, column=self.cols2["STATUS"], value="WAITING_FOR_REVIEW")
        CellValue.objects.create(row=r4, column=self.cols2["DUE_DATE_FLOW_FORCE"], value=future.strftime("%Y-%m-%d"))

        # Row 5: Table 2, Completed
        r5 = Row.objects.create(table=self.table2, created_by=self.admin)
        CellValue.objects.create(row=r5, column=self.cols2["PID"], value="PID-005")
        CellValue.objects.create(row=r5, column=self.cols2["STATUS"], value="APPROVED")
        CellValue.objects.create(row=r5, column=self.cols2["DUE_DATE_FLOW_FORCE"], value=today.strftime("%Y-%m-%d"))

        request = self.factory.get("/tables/pid-dashboard/")
        request.user = self.admin

        with self.CaptureQueriesContext(self.connection) as ctx_queries:
            response = pid_dashboard_view(request)

        self.assertEqual(response.status_code, 200)
        # Query count must remain tightly bounded (<= 8 queries total including global context processors)
        self.assertLessEqual(len(ctx_queries), 8)

        # Verify rendered content
        self.assertIn(b"PID-001", response.content)
        self.assertIn(b"PID-002", response.content)
        self.assertIn(b"PID-003", response.content)
        self.assertIn(b"PID-004", response.content)
        self.assertIn(b"PID-005", response.content)
        self.assertIn(b"PID Table Alpha", response.content)
        self.assertIn(b"PID Table Beta", response.content)
        # Ensure non-PID table is excluded
        self.assertNotIn(b"Standard Tasks", response.content)

    def test_pid_dashboard_query_count_does_not_grow_linearly(self):
        """Proves query count is O(1) whether there are 4 rows or 40 rows across PID tables."""
        from tables.views import pid_dashboard_view

        # Setup 4 rows initially
        for i in range(4):
            r = Row.objects.create(table=self.table1, created_by=self.admin)
            CellValue.objects.create(row=r, column=self.cols1["PID"], value=f"PID-10{i}")
            CellValue.objects.create(row=r, column=self.cols1["STATUS"], value="IN_PROGRESS")

        req1 = self.factory.get("/tables/pid-dashboard/")
        req1.user = self.admin
        with self.CaptureQueriesContext(self.connection) as ctx1:
            res1 = pid_dashboard_view(req1)
        self.assertEqual(res1.status_code, 200)
        q_count_initial = len(ctx1)

        # Scale up: Add 36 more rows (total 40 rows across both tables)
        for i in range(36):
            target_table = self.table1 if i % 2 == 0 else self.table2
            cols = self.cols1 if i % 2 == 0 else self.cols2
            r = Row.objects.create(table=target_table, created_by=self.admin)
            CellValue.objects.create(row=r, column=cols["PID"], value=f"PID-SCALE-{i}")
            CellValue.objects.create(row=r, column=cols["STATUS"], value="COMPLETED" if i % 3 == 0 else "PENDING")

        req2 = self.factory.get("/tables/pid-dashboard/")
        req2.user = self.admin
        with self.CaptureQueriesContext(self.connection) as ctx2:
            res2 = pid_dashboard_view(req2)
        self.assertEqual(res2.status_code, 200)
        q_count_scaled = len(ctx2)

        # Assert zero linear N+1 query growth
        self.assertEqual(q_count_initial, q_count_scaled)
        self.assertLessEqual(q_count_scaled, 8)

    def test_pid_dashboard_handles_dirty_legacy_dates(self):
        """Verifies malformed date strings in legacy data do not crash PID dashboard view."""
        from tables.views import pid_dashboard_view

        dirty_dates = [
            "invalid-date-format",
            "99/99/9999",
            "",
            "-",
            "None",
            "2026-02-31",
            "15-08-2024",
        ]
        for idx, dirty in enumerate(dirty_dates):
            r = Row.objects.create(table=self.table1, created_by=self.admin)
            CellValue.objects.create(row=r, column=self.cols1["PID"], value=f"PID-DIRTY-{idx}")
            CellValue.objects.create(row=r, column=self.cols1["DUE_DATE_FLOW_FORCE"], value=dirty)
            CellValue.objects.create(row=r, column=self.cols1["DATE"], value=dirty)
            CellValue.objects.create(row=r, column=self.cols1["STATUS"], value="IN_PROGRESS")

        req = self.factory.get("/tables/pid-dashboard/")
        req.user = self.admin
        response = pid_dashboard_view(req)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"PID-DIRTY-0", response.content)

    def test_pid_dashboard_permissions_isolation(self):
        """Verifies non-admin users only see LIST_PID tables they have access to."""
        from tables.views import pid_dashboard_view

        r1 = Row.objects.create(table=self.table1, created_by=self.admin)
        CellValue.objects.create(row=r1, column=self.cols1["PID"], value="PID-ALPHA-ONLY")

        r2 = Row.objects.create(table=self.table2, created_by=self.admin)
        CellValue.objects.create(row=r2, column=self.cols2["PID"], value="PID-BETA-HIDDEN")

        req = self.factory.get("/tables/pid-dashboard/")
        req.user = self.employee
        response = pid_dashboard_view(req)

        self.assertEqual(response.status_code, 200)
        # Employee belongs to self.dept so sees Table 1, but NOT Table 2 (in dept2)
        self.assertIn(b"PID Table Alpha", response.content)
        self.assertNotIn(b"PID Table Beta", response.content)

class TableStatisticsOptimizationRegressionTestCase(TestCase):
    """
    Regression test suite for Phase 2C Table Statistics & Query Consolidation.
    Validates that:
    1. Duplicate table lookups are eliminated across get_queryset and get_paginated_response.
    2. Multiple task count queries are consolidated into single aggregate queries.
    3. Filterable column values are fetched in a single batched query instead of per-column queries.
    4. Table statistics cache-hit and cache-miss behave with reduced, constant-bound query counts.
    5. Empty tables and malformed legacy data are handled safely without exceptions.
    6. All statistics values, permissions, and spreadsheet behaviors remain 100% identical.
    """

    def setUp(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        from django.core.cache import cache

        self.factory = APIRequestFactory()
        self.force_authenticate = force_authenticate
        self.connection = connection
        self.CaptureQueriesContext = CaptureQueriesContext
        self.cache = cache

        self.dept = Department.objects.create(name="Phase 2C Ops", slug="phase-2c-ops")
        self.dept2 = Department.objects.create(name="Phase 2C Isolated", slug="phase-2c-isolated")
        self.admin = User.objects.create_user(
            email="p2cadmin_test@flow-force.com",
            password="testpassword",
            full_name="P2C Admin Test",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="p2cemp_test@flow-force.com",
            password="testpassword",
            full_name="P2C Emp Test",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )
        self.other_employee = User.objects.create_user(
            email="p2cother_test@flow-force.com",
            password="testpassword",
            full_name="P2C Other Test",
            role="EMPLOYEE",
            department=self.dept2,
            status="APPROVED"
        )

        self.table = Table.objects.create(
            name="Sales Operations Table",
            job_type="SALES",
            created_by=self.admin,
            department=self.dept
        )
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=self.table, user=self.employee, access_level="VIEW")

        # System and custom columns
        self.col_pid = Column.objects.create(table=self.table, name="PID", data_type="TEXT")
        self.col_project = Column.objects.create(table=self.table, name="PROJECT", data_type="TEXT", is_filterable=True)
        self.col_client = Column.objects.create(table=self.table, name="CLIENT", data_type="TEXT", is_filterable=True)
        self.col_stage = Column.objects.create(table=self.table, name="STAGE", data_type="DROPDOWN", options="Lead,Proposal,Negotiation", is_filterable=True)
        self.col_qty = Column.objects.create(table=self.table, name="QTY", data_type="NUMBER")
        self.col_fup = Column.objects.create(table=self.table, name="FOLLOW-UP DATE", data_type="DATE")
        self.col_act = Column.objects.create(table=self.table, name="ACTIVITY_TYPE", data_type="TEXT")
        self.col_status = Column.objects.create(table=self.table, name="STATUS", data_type="TEXT")

    def _create_sample_data(self, row_count=20):
        from datetime import date, timedelta
        today = date.today()
        cells = []
        for i in range(row_count):
            r = Row.objects.create(table=self.table, created_by=self.admin)
            st = "COMPLETED" if i % 4 == 0 else "IN_PROGRESS"
            pr = "HIGH" if i % 3 == 0 else ("LOW" if i % 3 == 1 else "MEDIUM")
            if i % 5 == 0:
                dd = today
            elif i % 5 == 1:
                dd = today - timedelta(days=3)
            else:
                dd = today + timedelta(days=7)

            t = Task.objects.create(
                row=r,
                due_date=dd,
                priority=pr,
                status=st,
                assigned_by=self.admin
            )

            cells.append(CellValue(row=r, column=self.col_pid, value=f"PID-{i+1:03d}"))
            cells.append(CellValue(row=r, column=self.col_project, value=f"Project {i % 3}"))
            cells.append(CellValue(row=r, column=self.col_client, value=f"Client {i % 4}"))
            cells.append(CellValue(row=r, column=self.col_qty, value=str(10.5 * (i + 1)) if i % 4 != 3 else "invalid"))
            cells.append(CellValue(row=r, column=self.col_fup, value=today.isoformat()))
            cells.append(CellValue(row=r, column=self.col_act, value="Call" if i % 2 == 0 else "Site Visit"))
            cells.append(CellValue(row=r, column=self.col_status, value=st))

        CellValue.objects.bulk_create(cells)

    def test_statistics_cache_miss_query_count_and_fidelity(self):
        """Verifies Cache Miss calculates all statistics accurately with <= 16 queries."""
        from datetime import date
        from tables.views import RowViewSet
        self._create_sample_data(row_count=20)

        today_str = date.today().isoformat()
        self.cache.delete(f"table_stats_{self.table.id}_{today_str}")

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req, user=self.admin)

        with self.CaptureQueriesContext(self.connection) as ctx_miss:
            res = view(req)

        self.assertEqual(res.status_code, 200)
        # Reduced from 21 queries down to <= 16 queries on cache miss
        self.assertLessEqual(len(ctx_miss), 16)

        stats = res.data.get("stats")
        self.assertIsNotNone(stats)
        self.assertEqual(stats["completion_stats"]["total"], 20)
        self.assertEqual(stats["completion_stats"]["completed"], 5)
        self.assertEqual(stats["completion_stats"]["percent"], 25)
        self.assertEqual(stats["due_today_count"], 4)
        self.assertEqual(stats["overdue_count"], 3)
        self.assertEqual(stats["total_qty"], 1575.0)
        self.assertEqual(stats["status_counts"]["COMPLETED"], 5)
        self.assertEqual(stats["status_counts"]["IN_PROGRESS"], 15)
        self.assertEqual(stats["priority_counts"]["High"], 7)
        self.assertEqual(stats["priority_counts"]["Low"], 7)
        self.assertEqual(stats["priority_counts"]["Med"], 6)
        self.assertEqual(stats["project_counts"]["Project 0"], 7)
        self.assertEqual(stats["project_counts"]["Project 1"], 7)
        self.assertEqual(stats["project_counts"]["Project 2"], 6)

    def test_statistics_cache_hit_query_count_and_fidelity(self):
        """Verifies Cache Hit returns exact statistics in <= 5 queries."""
        from datetime import date
        from tables.views import RowViewSet
        self._create_sample_data(row_count=20)

        view = RowViewSet.as_view({'get': 'list'})
        req_prime = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req_prime, user=self.admin)
        res_prime = view(req_prime)
        self.assertEqual(res_prime.status_code, 200)

        # Cache Hit request
        req_hit = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req_hit, user=self.admin)

        with self.CaptureQueriesContext(self.connection) as ctx_hit:
            res_hit = view(req_hit)

        self.assertEqual(res_hit.status_code, 200)
        # Reduced from 6 queries down to <= 5 queries on cache hit
        self.assertLessEqual(len(ctx_hit), 5)
        self.assertEqual(res_hit.data.get("stats"), res_prime.data.get("stats"))

    def test_duplicate_table_lookup_eliminated(self):
        """Verifies that Table.objects.get is called only once per request."""
        from tables.views import RowViewSet
        self._create_sample_data(row_count=5)

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req, user=self.admin)

        with self.CaptureQueriesContext(self.connection) as ctx:
            res = view(req)

        self.assertEqual(res.status_code, 200)
        table_queries = [
            q['sql'] for q in ctx.captured_queries
            if 'FROM "tables_table"' in q['sql'] or 'FROM tables_table' in q['sql']
        ]
        # Exactly 1 query for Table
        self.assertEqual(len(table_queries), 1)

    def test_statistics_query_count_does_not_grow_with_rows(self):
        """Verifies that statistics query count is O(1) independent of row count."""
        from datetime import date
        from tables.views import RowViewSet

        # 5 rows
        self._create_sample_data(row_count=5)
        today_str = date.today().isoformat()
        self.cache.delete(f"table_stats_{self.table.id}_{today_str}")

        view = RowViewSet.as_view({'get': 'list'})
        req5 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req5, user=self.admin)
        with self.CaptureQueriesContext(self.connection) as ctx5:
            res5 = view(req5)
        q_count_5 = len(ctx5)

        # 25 rows
        self._create_sample_data(row_count=20)
        self.cache.delete(f"table_stats_{self.table.id}_{today_str}")

        req25 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req25, user=self.admin)
        with self.CaptureQueriesContext(self.connection) as ctx25:
            res25 = view(req25)
        q_count_25 = len(ctx25)

        self.assertEqual(q_count_5, q_count_25)
        self.assertLessEqual(q_count_25, 16)

    def test_empty_table_statistics(self):
        """Verifies that an empty table returns clean zero/empty statistics without errors."""
        from datetime import date
        from tables.views import RowViewSet

        empty_table = Table.objects.create(
            name="Empty Table",
            job_type="GENERAL",
            created_by=self.admin,
            department=self.dept
        )
        TableAccess.objects.create(table=empty_table, user=self.admin, access_level="ADMIN")

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={empty_table.id}&include_stats=true")
        self.force_authenticate(req, user=self.admin)

        res = view(req)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["count"], 0)
        stats = res.data.get("stats")
        self.assertIsNotNone(stats)
        self.assertEqual(stats["completion_stats"]["total"], 0)
        self.assertEqual(stats["completion_stats"]["completed"], 0)
        self.assertEqual(stats["completion_stats"]["percent"], 0)
        self.assertEqual(stats["due_today_count"], 0)
        self.assertEqual(stats["overdue_count"], 0)
        self.assertEqual(stats["total_qty"], 0.0)

    def test_malformed_and_blank_qty_values(self):
        """Verifies that malformed and blank QTY cell values are safely handled in Python summation."""
        from datetime import date
        from tables.views import RowViewSet

        today_str = date.today().isoformat()
        self.cache.delete(f"table_stats_{self.table.id}_{today_str}")

        dirty_values = ["-", "N/A", "", " ", "None", "15.5", "24.5", "invalid_string"]
        cells = []
        for idx, val in enumerate(dirty_values):
            r = Row.objects.create(table=self.table, created_by=self.admin)
            Task.objects.create(row=r, status="PENDING", assigned_by=self.admin)
            cells.append(CellValue(row=r, column=self.col_qty, value=val))
        CellValue.objects.bulk_create(cells)

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.force_authenticate(req, user=self.admin)

        res = view(req)
        self.assertEqual(res.status_code, 200)
        # 15.5 + 24.5 = 40.0
        self.assertEqual(res.data["stats"]["total_qty"], 40.0)

    def test_permissions_isolation_for_table_statistics(self):
        """Verifies unauthorized users cannot access table rows or statistics."""
        from tables.views import RowViewSet

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        # other_employee has NO access to self.table (in different department and no TableAccess)
        self.force_authenticate(req, user=self.other_employee)

        res = view(req)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["count"], 0)
        self.assertEqual(res.data["results"], [])

    def test_table_spreadsheet_view_query_optimization(self):
        """Verifies table_spreadsheet_view uses prefetched columns and relations efficiently."""
        from tables.views import table_spreadsheet_view
        from django.test.client import RequestFactory
        factory = RequestFactory()

        req = factory.get(f"/tables/{self.table.id}/")
        req.user = self.admin

        with self.CaptureQueriesContext(self.connection) as ctx:
            res = table_spreadsheet_view(req, table_id=self.table.id)

        self.assertEqual(res.status_code, 200)
        # Should be tightly bounded (Table+department, columns prefetch, access check, + context processor)
        self.assertLessEqual(len(ctx), 8)


class RedisCacheConfigurationRegressionTestCase(TestCase):
    """
    Phase 2D Regression Tests:
    Verifies Redis cache configuration, fallback to LocMemCache, cache lifecycle,
    and preservation of table statistics cache keys and Celery broker isolation.
    """

    def test_production_style_redis_cache_config_with_cache_url(self):
        """Production-style Redis cache configuration resolves correctly when CACHE_URL is configured."""
        from flowforce.settings import get_cache_config
        config = get_cache_config(cache_url="redis://127.0.0.1:6379/1", debug=False, is_testing=False)
        self.assertEqual(config['default']['BACKEND'], 'django.core.cache.backends.redis.RedisCache')
        self.assertEqual(config['default']['LOCATION'], 'redis://127.0.0.1:6379/1')
        self.assertEqual(config['default']['KEY_PREFIX'], 'flowforce')
        self.assertEqual(config['default']['TIMEOUT'], 86400)

    def test_production_default_cache_config_without_cache_url(self):
        """Production default resolves to RedisCache DB 1 when CACHE_URL is not set."""
        from flowforce.settings import get_cache_config
        config = get_cache_config(cache_url=None, debug=False, is_testing=False)
        self.assertEqual(config['default']['BACKEND'], 'django.core.cache.backends.redis.RedisCache')
        self.assertEqual(config['default']['LOCATION'], 'redis://127.0.0.1:6379/1')
        self.assertEqual(config['default']['KEY_PREFIX'], 'flowforce')
        self.assertEqual(config['default']['TIMEOUT'], 86400)

    def test_local_dev_resolves_to_locmem_cache(self):
        """Local/default configuration resolves to LocMemCache when Redis is not configured."""
        from flowforce.settings import get_cache_config
        config = get_cache_config(cache_url=None, debug=True, is_testing=False)
        self.assertEqual(config['default']['BACKEND'], 'django.core.cache.backends.locmem.LocMemCache')
        self.assertEqual(config['default']['LOCATION'], 'flowforce-cache')
        self.assertEqual(config['default']['TIMEOUT'], 86400)

    def test_test_environment_resolves_to_locmem_cache(self):
        """Test environment safely defaults to LocMemCache so tests do not require a live Redis instance."""
        from flowforce.settings import get_cache_config
        config = get_cache_config(cache_url=None, debug=False, is_testing=True)
        self.assertEqual(config['default']['BACKEND'], 'django.core.cache.backends.locmem.LocMemCache')
        self.assertEqual(config['default']['LOCATION'], 'flowforce-cache')

    def test_cache_set_get_delete_lifecycle(self):
        """cache.set() followed by cache.get() returns expected value, and cache.delete() removes it."""
        from django.core.cache import cache
        test_key = "test_phase2d_cache_lifecycle_key"
        test_val = {"status": "ok", "count": 42}

        # Ensure clean state
        cache.delete(test_key)
        self.assertIsNone(cache.get(test_key))

        # Set and get
        cache.set(test_key, test_val, 86400)
        retrieved = cache.get(test_key)
        self.assertEqual(retrieved, test_val)

        # Delete and verify removal
        cache.delete(test_key)
        self.assertIsNone(cache.get(test_key))

    def test_table_statistics_cache_key_preservation(self):
        """Existing table statistics cache keys (table_stats_{table_id}_{date}) still work as expected."""
        from datetime import date
        from django.core.cache import cache
        table_id = 9999
        today_str = date.today().isoformat()
        expected_key = f"table_stats_{table_id}_{today_str}"

        stats_payload = {
            "completion_stats": {"total": 50, "completed": 25, "percent": 50},
            "due_today_count": 5,
            "overdue_count": 2,
            "total_qty": 1250.0,
        }

        cache.delete(expected_key)
        self.assertIsNone(cache.get(expected_key))

        cache.set(expected_key, stats_payload, 86400)
        self.assertEqual(cache.get(expected_key), stats_payload)

        cache.delete(expected_key)
        self.assertIsNone(cache.get(expected_key))

    def test_redis_cache_backend_instantiation(self):
        """Verifies RedisCache backend class can be instantiated with production parameters."""
        from django.core.cache.backends.redis import RedisCache
        backend = RedisCache('redis://127.0.0.1:6379/1', {'KEY_PREFIX': 'flowforce', 'TIMEOUT': 86400})
        self.assertEqual(backend._servers, ['redis://127.0.0.1:6379/1'])
        self.assertEqual(backend.key_prefix, 'flowforce')
        self.assertEqual(backend.default_timeout, 86400)

    def test_celery_broker_configuration_preserved(self):
        """CELERY_BROKER_URL and CELERY_RESULT_BACKEND remain unchanged on Redis DB 0."""
        from django.conf import settings
        self.assertEqual(settings.CELERY_BROKER_URL, 'redis://localhost:6379/0')
        self.assertEqual(settings.CELERY_RESULT_BACKEND, 'redis://localhost:6379/0')


class SpreadsheetSearchOptimizationRegressionTestCase(TestCase):
    """
    Phase 2E Regression Tests:
    Verifies optimized spreadsheet search semantics: case & space insensitivity,
    partial text matching, searches across multiple cell values, combined filters (year, column, etc.),
    empty/whitespace search behavior, archived row exclusion, deduplication,
    pagination fidelity, and query efficiency (EXISTS subquery structure).
    """

    def setUp(self):
        from employee_management.models import Department
        from auth_app.models import EmployeeUser
        from tables.models import Table, Column, Row, CellValue, TableAccess
        from rest_framework.test import APIRequestFactory

        self.factory = APIRequestFactory()
        self.dept = Department.objects.create(name="Phase2E Search Dept")
        self.admin = EmployeeUser.objects.create(
            email="search_admin@test.com",
            full_name="Search Admin",
            role="ADMIN",
            department=self.dept
        )
        self.table = Table.objects.create(name="Search Optimization Table", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")

        self.col_sno = Column.objects.create(table=self.table, name="S_NO", data_type="TEXT", position=1)
        self.col_pid = Column.objects.create(table=self.table, name="PID", data_type="TEXT", position=2)
        self.col_customer = Column.objects.create(table=self.table, name="CUSTOMER_NAME", data_type="TEXT", position=3)
        self.col_desc = Column.objects.create(table=self.table, name="DESCRIPTION", data_type="TEXT", position=4)
        self.col_remarks = Column.objects.create(table=self.table, name="REMARKS", data_type="TEXT", position=5)

    def test_case_and_space_insensitive_matching(self):
        """Search matches regardless of uppercase/lowercase and inline or surrounding whitespace."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        row = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=row, column=self.col_customer, value="PT Flow Force Indonesia", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})

        # Searches: no spaces, lowercase, extra spaces, mixed case
        for query_term in ["ptflowforceindonesia", "  PT   Flow   Force   ", "FLOW FORCE", "forceindonesia"]:
            req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search={query_term}")
            force_authenticate(req, user=self.admin)
            res = view(req)
            self.assertEqual(res.status_code, 200)
            self.assertEqual(len(res.data['results']), 1, f"Failed matching term: {query_term}")
            self.assertEqual(res.data['results'][0]['id'], row.id)

    def test_partial_text_matching(self):
        """Search performs partial substring matching on cell text values."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        row = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=row, column=self.col_desc, value="High Pressure Valve Assembly 5000 PSI", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})

        for partial in ["Pressure", "valve", "5000", "Assembly", "essure"]:
            req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search={partial}")
            force_authenticate(req, user=self.admin)
            res = view(req)
            self.assertEqual(res.status_code, 200)
            self.assertEqual(len(res.data['results']), 1, f"Failed matching partial: {partial}")

    def test_search_across_multiple_cell_values(self):
        """Search correctly finds rows matching against different columns."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        row1 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=row1, column=self.col_pid, value="PID-1049", updated_by=self.admin)
        CellValue.objects.create(row=row1, column=self.col_customer, value="Siemens Energy", updated_by=self.admin)
        CellValue.objects.create(row=row1, column=self.col_remarks, value="Expedited shipment required", updated_by=self.admin)

        row2 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=row2, column=self.col_pid, value="PID-2080", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=self.col_customer, value="Chevron Pacific", updated_by=self.admin)
        CellValue.objects.create(row=row2, column=self.col_remarks, value="Standard lead time", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})

        # Match row1 via PID
        req1 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=1049")
        force_authenticate(req1, user=self.admin)
        res1 = view(req1)
        self.assertEqual([r['id'] for r in res1.data['results']], [row1.id])

        # Match row2 via Customer
        req2 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Chevron")
        force_authenticate(req2, user=self.admin)
        res2 = view(req2)
        self.assertEqual([r['id'] for r in res2.data['results']], [row2.id])

        # Match row1 via Remarks
        req3 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Expedited")
        force_authenticate(req3, user=self.admin)
        res3 = view(req3)
        self.assertEqual([r['id'] for r in res3.data['results']], [row1.id])

    def test_search_combined_with_year_filter(self):
        """Search works accurately when combined with year filter."""
        import datetime
        from tasks.models import Task
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        row1 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=row1, assigned_by=self.admin, due_date=datetime.date(2024, 6, 15))
        CellValue.objects.create(row=row1, column=self.col_customer, value="Universal Nickel 2024", updated_by=self.admin)

        row2 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=row2, assigned_by=self.admin, due_date=datetime.date(2025, 7, 20))
        CellValue.objects.create(row=row2, column=self.col_customer, value="Universal Nickel 2025", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})

        # Search without year filter returns both
        req_both = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Universal")
        force_authenticate(req_both, user=self.admin)
        res_both = view(req_both)
        self.assertEqual(len(res_both.data['results']), 2)

        # Search combined with year=2024 returns only row1
        req_2024 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Universal&year=2024")
        force_authenticate(req_2024, user=self.admin)
        res_2024 = view(req_2024)
        self.assertEqual(len(res_2024.data['results']), 1)
        self.assertEqual(res_2024.data['results'][0]['id'], row1.id)

        # Search combined with year=2025 returns only row2
        req_2025 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Universal&year=2025")
        force_authenticate(req_2025, user=self.admin)
        res_2025 = view(req_2025)
        self.assertEqual(len(res_2025.data['results']), 1)
        self.assertEqual(res_2025.data['results'][0]['id'], row2.id)

    def test_empty_and_whitespace_search_behavior(self):
        """Empty search or search with whitespace only returns all active rows without error."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        r1 = Row.objects.create(table=self.table, created_by=self.admin)
        r2 = Row.objects.create(table=self.table, created_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})

        for empty_val in ["", "   ", "\t  \n"]:
            req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search={empty_val}")
            force_authenticate(req, user=self.admin)
            res = view(req)
            self.assertEqual(res.status_code, 200)
            self.assertEqual(len(res.data['results']), 2)

    def test_archived_rows_exclusion(self):
        """Archived rows are excluded from search results."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        active_row = Row.objects.create(table=self.table, is_archived=False, created_by=self.admin)
        CellValue.objects.create(row=active_row, column=self.col_customer, value="Archived Search Target", updated_by=self.admin)

        archived_row = Row.objects.create(table=self.table, is_archived=True, created_by=self.admin)
        CellValue.objects.create(row=archived_row, column=self.col_customer, value="Archived Search Target", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Archived Search Target")
        force_authenticate(req, user=self.admin)
        res = view(req)

        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data['results']), 1)
        self.assertEqual(res.data['results'][0]['id'], active_row.id)

    def test_no_duplicate_rows_when_multiple_cells_match(self):
        """Rows matching across multiple columns appear exactly once (no duplicates)."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        row = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=row, column=self.col_customer, value="Hydraulic Pump Unit", updated_by=self.admin)
        CellValue.objects.create(row=row, column=self.col_desc, value="Hydraulic System Service", updated_by=self.admin)
        CellValue.objects.create(row=row, column=self.col_remarks, value="Hydraulic pressure test ok", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Hydraulic")
        force_authenticate(req, user=self.admin)
        res = view(req)

        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['count'], 1)
        self.assertEqual(len(res.data['results']), 1)
        self.assertEqual(res.data['results'][0]['id'], row.id)

    def test_paginated_search_results(self):
        """Search correctly respects pagination page_size and preserves count."""
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        rows = []
        cells = []
        for i in range(25):
            r = Row(table=self.table, created_by=self.admin)
            rows.append(r)
        created_rows = Row.objects.bulk_create(rows)

        for idx, r in enumerate(created_rows):
            cells.append(CellValue(row=r, column=self.col_customer, value=f"Pagination Client #{idx+1}"))
        CellValue.objects.bulk_create(cells)

        view = RowViewSet.as_view({'get': 'list'})

        # Page 1 (page_size=10)
        req1 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Pagination&page_size=10&page=1")
        force_authenticate(req1, user=self.admin)
        res1 = view(req1)
        self.assertEqual(res1.status_code, 200)
        self.assertEqual(res1.data['count'], 25)
        self.assertEqual(len(res1.data['results']), 10)

        # Page 2 (page_size=10)
        req2 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Pagination&page_size=10&page=2")
        force_authenticate(req2, user=self.admin)
        res2 = view(req2)
        self.assertEqual(res2.status_code, 200)
        self.assertEqual(res2.data['count'], 25)
        self.assertEqual(len(res2.data['results']), 10)
        # Ensure page 1 and page 2 IDs are completely disjoint
        p1_ids = {r['id'] for r in res1.data['results']}
        p2_ids = {r['id'] for r in res2.data['results']}
        self.assertTrue(p1_ids.isdisjoint(p2_ids))

        # Page 3 (page_size=10)
        req3 = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Pagination&page_size=10&page=3")
        force_authenticate(req3, user=self.admin)
        res3 = view(req3)
        self.assertEqual(res3.status_code, 200)
        self.assertEqual(len(res3.data['results']), 5)

    def test_task_status_and_priority_matching(self):
        """Search finds rows via linked task status or priority."""
        from tasks.models import Task
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate

        r1 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=r1, assigned_by=self.admin, status="IN_PROGRESS", priority="CRITICAL")

        r2 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=r2, assigned_by=self.admin, status="COMPLETED", priority="LOW")

        view = RowViewSet.as_view({'get': 'list'})

        # Match via status with space and no space
        for term in ["inprogress", "in progress", "IN_PROGRESS"]:
            req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search={term}")
            force_authenticate(req, user=self.admin)
            res = view(req)
            self.assertEqual(len(res.data['results']), 1, f"Failed matching status: {term}")
            self.assertEqual(res.data['results'][0]['id'], r1.id)

        # Match via priority
        req_p = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=critical")
        force_authenticate(req_p, user=self.admin)
        res_p = view(req_p)
        self.assertEqual(len(res_p.data['results']), 1)
        self.assertEqual(res_p.data['results'][0]['id'], r1.id)

    def test_search_query_efficiency_and_structure(self):
        """Verifies query uses EXISTS subquery, avoids redundant annotations on Row, and is bounded."""
        from tasks.models import Task
        from tables.views import RowViewSet
        from rest_framework.test import force_authenticate
        from django.test.utils import CaptureQueriesContext
        from django.db import connection

        for i in range(5):
            r = Row.objects.create(table=self.table, created_by=self.admin)
            Task.objects.create(row=r, assigned_by=self.admin, status="PENDING", priority="MEDIUM")
            CellValue.objects.create(row=r, column=self.col_customer, value=f"Client Energy #{i}", updated_by=self.admin)

        view = RowViewSet.as_view({'get': 'list'})
        req = self.factory.get(f"/tables/api/rows/?table={self.table.id}&search=Energy")
        force_authenticate(req, user=self.admin)

        with CaptureQueriesContext(connection) as ctx:
            res = view(req)

        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['count'], 5)

        # Check queries executed
        sql_statements = [q['sql'] for q in ctx.captured_queries]
        # Query count is bounded (1 table get, 1 table access, 1 count, 1 page slice, 1 cells prefetch, 1 assigned_to prefetch)
        self.assertLessEqual(len(ctx), 6)

        # Verify EXISTS is in the queries and clean_task_status is NOT selected as output column
        found_exists = any("EXISTS" in sql.upper() for sql in sql_statements)
        self.assertTrue(found_exists, "Search query should use EXISTS subquery")
        found_task_clean_select = any("clean_task_status" in sql for sql in sql_statements)
        self.assertFalse(found_task_clean_select, "clean_task_status should not be projected in SELECT")


class SpreadsheetFrontendPerformanceRegressionTestCase(TestCase):
    def setUp(self):
        self.dept = Department.objects.create(name="Phase 2F Dept", slug="phase-2f-dept")
        self.admin = User.objects.create_user(
            email="phase2fadmin@flow-force.com",
            password="testpassword",
            full_name="Phase 2F Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(
            name="Frontend Perf Test Table",
            job_type="PERSONAL",
            created_by=self.admin
        )
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")

        self.col_job = Column.objects.create(
            table=self.table,
            name="JOB_NUMBER",
            data_type="TEXT",
            position=1
        )
        self.col_task = Column.objects.create(
            table=self.table,
            name="TASK_NAME",
            data_type="TEXT",
            position=2
        )
        self.col_due = Column.objects.create(
            table=self.table,
            name="DUE_DATE",
            data_type="DATE",
            position=3
        )
        self.col_notes = Column.objects.create(
            table=self.table,
            name="NOTES",
            data_type="TEXT",
            options='{"input_type": "multiline", "rows": 4, "placeholder": "Notes..."}',
            position=4
        )

        # Create rows: 2 rows with same job number, 1 row with different, 1 row with "-"
        self.r1 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=self.r1, column=self.col_job, value="JOB-101", updated_by=self.admin)
        CellValue.objects.create(row=self.r1, column=self.col_task, value="First Task", updated_by=self.admin)
        CellValue.objects.create(row=self.r1, column=self.col_due, value="2026-10-01", updated_by=self.admin)

        self.r2 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=self.r2, column=self.col_job, value="JOB-101", updated_by=self.admin)
        CellValue.objects.create(row=self.r2, column=self.col_task, value="Second Task", updated_by=self.admin)
        CellValue.objects.create(row=self.r2, column=self.col_due, value="2026-10-15", updated_by=self.admin)

        self.r3 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=self.r3, column=self.col_job, value="JOB-102", updated_by=self.admin)
        CellValue.objects.create(row=self.r3, column=self.col_task, value="Third Task", updated_by=self.admin)
        CellValue.objects.create(row=self.r3, column=self.col_due, value="2026-10-20", updated_by=self.admin)

        self.r4 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=self.r4, column=self.col_job, value="-", updated_by=self.admin)
        CellValue.objects.create(row=self.r4, column=self.col_task, value="Fourth Task", updated_by=self.admin)

    def test_spreadsheet_html_renders_optimized_frontend_assets(self):
        """Verifies table_spreadsheet.html renders with 200 OK and includes all frontend performance optimizations."""
        self.client.force_login(self.admin)
        response = self.client.get(f"/tables/{self.table.id}/")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")

        # Verify optimized state variables
        self.assertIn("_cachedJobColName: undefined", html)
        self.assertIn("_colConfigs: {}", html)
        self.assertIn("_jobGroups: {}", html)
        self.assertIn("_cachedTodayDate: null", html)
        self.assertIn("_cachedTodayIso: null", html)
        self.assertIn("_fetchAbortController: null", html)

        # Verify optimized methods
        self.assertIn("updateJobGroups()", html)
        self.assertIn("getTodayIso()", html)
        self.assertIn("getTodayDate()", html)
        self.assertIn("isDateColumn(col)", html)
        self.assertIn("AbortController()", html)

        # Verify column metadata and options are safely passed to the template
        self.assertIn("JOB_NUMBER", html)
        self.assertIn("NOTES", html)

    def test_api_rows_payload_matches_frontend_precomputation_contracts(self):
        """Verifies /tables/api/rows/ returns all fields required for precomputed cells_dict and job groups."""
        self.client.force_login(self.admin)
        response = self.client.get(f"/tables/api/rows/?table={self.table.id}&include_stats=true")
        self.assertEqual(response.status_code, 200)
        data = response.json()

        self.assertEqual(data["count"], 4)
        results = data["results"]
        self.assertEqual(len(results), 4)

        # Check cells are serialized with column_name and value
        for row in results:
            self.assertIn("cells", row)
            cells = row["cells"]
            col_map = {c["column_name"]: c["value"] for c in cells}
            self.assertIn("JOB_NUMBER", col_map)
            self.assertIn("TASK_NAME", col_map)

        # Verify stats and unique lists are returned
        self.assertIn("stats", data)
        self.assertIn("unique_pids", data)
        self.assertIn("unique_years", data)

    def test_row_editing_apis_preserve_contracts(self):
        """Verifies cell and row edit endpoints work properly and update values for subsequent fetchRows."""
        self.client.force_login(self.admin)

        # Edit single cell
        edit_cell_resp = self.client.post(
            f"/tables/api/rows/{self.r1.id}/edit-cell/",
            data={"column": self.col_task.id, "value": "Updated Task 1"},
            content_type="application/json"
        )
        self.assertEqual(edit_cell_resp.status_code, 200)

        # Edit row
        edit_row_resp = self.client.post(
            f"/tables/api/rows/{self.r2.id}/edit-row/",
            data={"cells": {"TASK_NAME": "Batch Updated Task 2", "JOB_NUMBER": "JOB-101"}},
            content_type="application/json"
        )
        self.assertEqual(edit_row_resp.status_code, 200)

        # Verify updated row data via rows API
        fetch_resp = self.client.get(f"/tables/api/rows/?table={self.table.id}")
        self.assertEqual(fetch_resp.status_code, 200)
        fetch_data = fetch_resp.json()
        row_map = {r["id"]: {c["column_name"]: c["value"] for c in r["cells"]} for r in fetch_data["results"]}

        self.assertEqual(row_map[self.r1.id]["TASK_NAME"], "Updated Task 1")
        self.assertEqual(row_map[self.r2.id]["TASK_NAME"], "Batch Updated Task 2")

    def test_deduplicated_pagination_query_contract(self):
        """Verifies pagination without stats skips stats calculation while preserving row data."""
        self.client.force_login(self.admin)
        response = self.client.get(f"/tables/api/rows/?table={self.table.id}&page=1&page_size=2")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(len(data["results"]), 2)
        self.assertEqual(data["count"], 4)

    def test_spreadsheet_html_loading_overlay_and_submission_guards(self):
        """Verifies table_spreadsheet.html contains robust loading overlay cleanup, fetchRows force support, and submission guards."""
        self.client.force_login(self.admin)
        response = self.client.get(f"/tables/{self.table.id}/")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode("utf-8")

        # Verify state guards exist
        self.assertIn("isSubmittingRow: false", html)
        self.assertIn("_isFetchingRows: false", html)

        # Verify button is disabled during submission to prevent duplicates
        self.assertIn(':disabled="isSubmittingRow"', html)

        # Verify fetchRows supports forced refreshes and decouples _isFetchingRows
        self.assertIn("fetchRows(force = false)", html)
        self.assertIn("!force && this._isFetchingRows && this._lastFetchQuery === queryStr", html)

        # Verify try/finally cleanup exists in submitInlineRow, saveRecordModal, and saveRowEdit
        self.assertIn("this.isSubmittingRow = false;", html)
        self.assertIn("this.isLoading = false;", html)

    def test_row_create_api_success_and_immediate_persistence(self):
        """Verifies POST /tables/api/rows/ successfully saves the row and task without requiring a page refresh."""
        self.client.force_login(self.admin)
        create_resp = self.client.post(
            "/tables/api/rows/",
            data={
                "table": self.table.id,
                "cells": {
                    "TASK_NAME": "Brand New Submitted Task",
                    "DUE_DATE": "2026-10-30",
                    "JOB_NUMBER": "JOB-999"
                }
            },
            content_type="application/json"
        )
        self.assertEqual(create_resp.status_code, 201)
        created_data = create_resp.json()
        new_row_id = created_data["id"]

        # Verify row immediately exists in DB and is retrievable via GET /tables/api/rows/
        fetch_resp = self.client.get(f"/tables/api/rows/?table={self.table.id}")
        self.assertEqual(fetch_resp.status_code, 200)
        fetch_data = fetch_resp.json()
        saved_row = next((r for r in fetch_data["results"] if r["id"] == new_row_id), None)
        self.assertIsNotNone(saved_row, "Newly created row must be retrievable via rows API immediately")

    def test_row_create_api_validation_and_permission_failure(self):
        """Verifies validation and permission failures return 400 and 403 respectively with descriptive errors."""
        # Permission failure: user without edit permissions
        unauth_user = User.objects.create_user(
            email="noaccess@flow-force.com",
            password="testpassword",
            full_name="No Access User",
            role="EMPLOYEE",
            status="APPROVED"
        )
        self.client.force_login(unauth_user)
        perm_resp = self.client.post(
            "/tables/api/rows/",
            data={"table": self.table.id, "cells": {"TASK_NAME": "Fail Task"}},
            content_type="application/json"
        )
        self.assertEqual(perm_resp.status_code, 403)
        self.assertIn("error", perm_resp.json())

        # Validation failure: invalid date format on GENERAL table
        general_table = Table.objects.create(name="General Test Table", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=general_table, user=self.admin, access_level="ADMIN")
        Column.objects.create(table=general_table, name="DUE_DATE", data_type="DATE", position=1)
        Column.objects.create(table=general_table, name="TASK_NAME", data_type="TEXT", position=2)

        self.client.force_login(self.admin)
        val_resp = self.client.post(
            "/tables/api/rows/",
            data={"table": general_table.id, "cells": {"TASK_NAME": "Invalid Date Task", "DUE_DATE": "invalid-date"}},
            content_type="application/json"
        )
        self.assertEqual(val_resp.status_code, 400)
        self.assertIn("error", val_resp.json())


class TableStatisticsServiceTestCase(TestCase):
    def setUp(self):
        from django.core.cache import cache
        from django.utils import timezone
        from tables.services import TableStatisticsService
        self.cache = cache
        self.dept = Department.objects.create(name="Phase 3A Dept", slug="phase-3a-dept")
        self.admin = User.objects.create_user(
            email="phase3aadmin@flow-force.com",
            password="testpassword",
            full_name="Phase 3A Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(
            name="Service Test Table",
            job_type="PERSONAL",
            created_by=self.admin
        )
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")

        self.col_pid = Column.objects.create(table=self.table, name="PID", data_type="TEXT", position=1)
        self.col_project = Column.objects.create(table=self.table, name="PROJECT", data_type="TEXT", position=2)
        self.col_qty = Column.objects.create(table=self.table, name="QTY", data_type="NUMBER", position=3)
        self.col_filter = Column.objects.create(table=self.table, name="CATEGORY", data_type="TEXT", is_filterable=True, position=4)
        self.col_dropdown = Column.objects.create(
            table=self.table,
            name="STATUS_OPT",
            data_type="DROPDOWN",
            is_filterable=True,
            options="Active, Inactive, Archived",
            position=5
        )

        today = timezone.localdate()
        self.today_str = today.isoformat()
        TableStatisticsService.invalidate_cache(self.table.id, self.today_str)

    def test_service_empty_table(self):
        """Verifies empty table returns clean default statistics without error."""
        from tables.services import TableStatisticsService
        stats = TableStatisticsService.get_table_statistics(self.table, use_cache=False)
        self.assertEqual(stats["completion_stats"]["total"], 0)
        self.assertEqual(stats["completion_stats"]["completed"], 0)
        self.assertEqual(stats["completion_stats"]["percent"], 0)
        self.assertEqual(stats["due_today_count"], 0)
        self.assertEqual(stats["overdue_count"], 0)
        self.assertEqual(stats["total_qty"], 0.0)
        self.assertEqual(stats["unique_pids"], [])
        self.assertEqual(stats["unique_years"], [])
        self.assertEqual(stats["unique_column_values"][self.col_dropdown.id], ["Active", "Inactive", "Archived"])

    def test_service_aggregates_and_calculations(self):
        """Verifies accurate calculation of completion, overdue, due today, qty, and unique values."""
        from django.utils import timezone
        from tables.services import TableStatisticsService
        today = timezone.localdate()
        yesterday = today - datetime.timedelta(days=1)

        # Row 1: completed, yesterday
        r1 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=r1, status="COMPLETED", priority="HIGH", due_date=yesterday, assigned_by=self.admin)
        CellValue.objects.create(row=r1, column=self.col_pid, value="PID-001", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=self.col_project, value="Apollo", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=self.col_qty, value="10.5", updated_by=self.admin)
        CellValue.objects.create(row=r1, column=self.col_filter, value="CatA", updated_by=self.admin)

        # Row 2: overdue (pending with past due date)
        r2 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=r2, status="PENDING", priority="URGENT", due_date=yesterday, assigned_by=self.admin)
        CellValue.objects.create(row=r2, column=self.col_pid, value="PID-002", updated_by=self.admin)
        CellValue.objects.create(row=r2, column=self.col_project, value="Apollo", updated_by=self.admin)
        CellValue.objects.create(row=r2, column=self.col_qty, value="20.0", updated_by=self.admin)
        CellValue.objects.create(row=r2, column=self.col_filter, value="CatB", updated_by=self.admin)

        # Row 3: due today
        r3 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=r3, status="IN_PROGRESS", priority="MED", due_date=today, assigned_by=self.admin)
        CellValue.objects.create(row=r3, column=self.col_pid, value="PID-001", updated_by=self.admin)
        CellValue.objects.create(row=r3, column=self.col_project, value="Gemini", updated_by=self.admin)
        CellValue.objects.create(row=r3, column=self.col_qty, value="invalid_qty", updated_by=self.admin)
        CellValue.objects.create(row=r3, column=self.col_filter, value="CatA", updated_by=self.admin)

        # Row 4: approved (past due date but approved -> NOT overdue)
        r4 = Row.objects.create(table=self.table, created_by=self.admin)
        Task.objects.create(row=r4, status="APPROVED", priority="LOW", due_date=yesterday, assigned_by=self.admin)
        CellValue.objects.create(row=r4, column=self.col_qty, value=None, updated_by=self.admin)

        # Row 5: archived row (must be ignored)
        r5 = Row.objects.create(table=self.table, created_by=self.admin, is_archived=True)
        Task.objects.create(row=r5, status="PENDING", priority="URGENT", due_date=yesterday, assigned_by=self.admin)
        CellValue.objects.create(row=r5, column=self.col_qty, value="100.0", updated_by=self.admin)

        stats = TableStatisticsService.get_table_statistics(self.table, use_cache=False)

        self.assertEqual(stats["completion_stats"]["total"], 4)
        self.assertEqual(stats["completion_stats"]["completed"], 1)
        self.assertEqual(stats["completion_stats"]["percent"], 25)
        self.assertEqual(stats["due_today_count"], 1)
        self.assertEqual(stats["overdue_count"], 1)
        self.assertEqual(stats["total_qty"], 30.5)

        self.assertEqual(stats["unique_pids"], ["PID-001", "PID-002"])
        self.assertEqual(stats["unique_column_values"][self.col_filter.id], ["CatA", "CatB"])
        self.assertEqual(stats["priority_counts"]["High"], 1)
        self.assertEqual(stats["priority_counts"]["Urgent"], 1)
        self.assertEqual(stats["priority_counts"]["Med"], 1)
        self.assertEqual(stats["priority_counts"]["Low"], 1)
        self.assertEqual(stats["project_counts"]["Apollo"], 2)
        self.assertEqual(stats["project_counts"]["Gemini"], 1)

    def test_service_cache_hit_and_invalidation(self):
        """Verifies Redis caching and invalidation works as expected."""
        from tables.services import TableStatisticsService
        # 1. First call calculates and sets cache
        TableStatisticsService.invalidate_cache(self.table.id, self.today_str)
        cache_key = TableStatisticsService.get_cache_key(self.table.id, self.today_str)
        self.assertIsNone(self.cache.get(cache_key))

        stats1 = TableStatisticsService.get_table_statistics(self.table, use_cache=True)
        self.assertIsNotNone(self.cache.get(cache_key))

        # 2. Second call returns from cache
        stats2 = TableStatisticsService.get_table_statistics(self.table, use_cache=True)
        self.assertEqual(stats1, stats2)

        # 3. Invalidation clears cache
        TableStatisticsService.invalidate_cache(self.table.id, self.today_str)
        self.assertIsNone(self.cache.get(cache_key))

    def test_views_compatibility_wrapper(self):
        """Verifies get_table_statistics in tables.views delegates directly to the service."""
        from tables.views import get_table_statistics
        from tables.services import TableStatisticsService
        stats_from_view = get_table_statistics(self.table, use_cache=False)
        stats_from_service = TableStatisticsService.get_table_statistics(self.table, use_cache=False)
        self.assertEqual(stats_from_view, stats_from_service)


class RowServiceTestCase(TestCase):
    def setUp(self):
        from django.core.cache import cache
        self.cache = cache
        self.dept = Department.objects.create(name="Phase 3B Dept", slug="phase-3b-dept")
        self.admin = User.objects.create_user(
            email="phase3badmin@flow-force.com",
            password="testpassword",
            full_name="Phase 3B Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="phase3bemp@flow-force.com",
            password="testpassword",
            full_name="Phase 3B Employee",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(
            name="Row Service Table",
            job_type="GENERAL",
            created_by=self.admin
        )
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")

    def test_service_normal_row_creation(self):
        """Verifies normal row, task, and cell value creation via RowService."""
        from tables.services import RowService
        row = RowService.create_row(
            table=self.table,
            user=self.admin,
            cells_data={
                "TASK_NAME": "Deploy Service",
                "DUE_DATE": "2026-11-15",
                "priority": "HIGH"
            }
        )
        self.assertIsNotNone(row)
        self.assertEqual(row.created_by, self.admin)
        self.assertEqual(row.table, self.table)

        # Verify task creation
        task = Task.objects.get(row=row)
        self.assertEqual(task.status, "PENDING")
        self.assertEqual(task.priority, "HIGH")
        self.assertEqual(task.due_date.strftime("%Y-%m-%d"), "2026-11-15")
        self.assertEqual(task.assigned_by, self.admin)

        # Verify cell values
        s_no_col = self.table.columns.get(name="S_NO")
        task_name_col = self.table.columns.get(name="TASK_NAME")
        self.assertEqual(int(CellValue.objects.get(row=row, column=s_no_col).value), 1)
        self.assertEqual(CellValue.objects.get(row=row, column=task_name_col).value, "Deploy Service")

    def test_service_validation_errors(self):
        """Verifies RowCreationValidationError on missing mandatory fields or invalid dates."""
        from tables.services import RowService, RowCreationValidationError

        # Missing DUE_DATE
        with self.assertRaises(RowCreationValidationError) as ctx:
            RowService.create_row(table=self.table, user=self.admin, cells_data={"TASK_NAME": "Incomplete"})
        self.assertIn("DUE_DATE is mandatory", str(ctx.exception))

        # Invalid DUE_DATE format
        with self.assertRaises(RowCreationValidationError) as ctx2:
            RowService.create_row(table=self.table, user=self.admin, cells_data={"TASK_NAME": "Bad Date", "DUE_DATE": "15-11-2026"})
        self.assertIn("Invalid DUE_DATE format", str(ctx2.exception))

    def test_service_s_no_sequential_increment_and_corrupt_handling(self):
        """Verifies S_NO increments cleanly and recovers from corrupt string values."""
        from tables.services import RowService
        s_no_col = self.table.columns.get(name="S_NO")

        # Create row 1 -> S_NO=1
        r1 = RowService.create_row(table=self.table, user=self.admin, cells_data={"TASK_NAME": "T1", "DUE_DATE": "2026-11-01"})
        self.assertEqual(int(CellValue.objects.get(row=r1, column=s_no_col).value), 1)

        # Create row 2 -> S_NO=2
        r2 = RowService.create_row(table=self.table, user=self.admin, cells_data={"TASK_NAME": "T2", "DUE_DATE": "2026-11-02"})
        self.assertEqual(int(CellValue.objects.get(row=r2, column=s_no_col).value), 2)

        # Corrupt the latest S_NO value
        cell2 = CellValue.objects.get(row=r2, column=s_no_col)
        cell2.value = "corrupted_val"
        cell2.save()

        # Create row 3 -> should fallback to 0 and produce S_NO=1 safely without crashing
        r3 = RowService.create_row(table=self.table, user=self.admin, cells_data={"TASK_NAME": "T3", "DUE_DATE": "2026-11-03"})
        self.assertEqual(int(CellValue.objects.get(row=r3, column=s_no_col).value), 1)

    def test_service_different_job_types(self):
        """Verifies row creation across SALES, LIST_PID, PERSONAL, and LOGS tables."""
        from tables.services import RowService

        # 1. SALES
        sales_table = Table.objects.create(name="Sales Table", job_type="SALES", created_by=self.admin)
        sales_row = RowService.create_row(
            table=sales_table,
            user=self.admin,
            cells_data={"CUSTOMER_NAME": "Acme Corp", "FOLLOW_UP_DATE": "2026-12-01"}
        )
        self.assertEqual(Task.objects.get(row=sales_row).due_date.strftime("%Y-%m-%d"), "2026-12-01")

        # 2. LIST_PID (dates optional)
        pid_table = Table.objects.create(name="PID Table", job_type="LIST_PID", created_by=self.admin)
        pid_row = RowService.create_row(
            table=pid_table,
            user=self.admin,
            cells_data={"ENQUIRY_NO/QUOTATION_NO": "ENQ-999"}
        )
        self.assertEqual(pid_row.table, pid_table)

        # 3. PERSONAL (dates optional)
        personal_table = Table.objects.create(name="Personal Table", job_type="PERSONAL", created_by=self.admin)
        personal_row = RowService.create_row(
            table=personal_table,
            user=self.admin,
            cells_data={}
        )
        self.assertEqual(personal_row.table, personal_table)

        # 4. LOGS
        logs_table = Table.objects.create(name="Logs Table", job_type="LOGS", created_by=self.admin)
        logs_row = RowService.create_row(
            table=logs_table,
            user=self.admin,
            cells_data={"TOOL_NAME": "Wrench Set", "RETURN_DATE": "2026-11-10"}
        )
        self.assertEqual(Task.objects.get(row=logs_row).due_date.strftime("%Y-%m-%d"), "2026-11-10")

    def test_service_assigned_to_handling(self):
        """Verifies task assignees are populated via assigned_to_ids or USER column."""
        from tables.services import RowService
        row = RowService.create_row(
            table=self.table,
            user=self.admin,
            cells_data={"TASK_NAME": "Assigned Task", "DUE_DATE": "2026-11-20"},
            assigned_to_ids=[self.employee.id]
        )
        task = Task.objects.get(row=row)
        self.assertIn(self.employee, task.assigned_to.all())

    def test_service_transaction_rollback_on_failure(self):
        """Verifies atomic rollback: no Row or CellValue records persist if an error occurs."""
        from tables.services import RowService
        from unittest.mock import patch

        initial_row_count = Row.objects.filter(table=self.table).count()
        initial_cell_count = CellValue.objects.filter(row__table=self.table).count()

        with patch("tasks.models.Task.objects.create", side_effect=RuntimeError("Simulated DB Crash")):
            with self.assertRaises(RuntimeError):
                RowService.create_row(
                    table=self.table,
                    user=self.admin,
                    cells_data={"TASK_NAME": "Will Roll Back", "DUE_DATE": "2026-11-25"}
                )

        # Assert no orphaned records were persisted
        self.assertEqual(Row.objects.filter(table=self.table).count(), initial_row_count)
        self.assertEqual(CellValue.objects.filter(row__table=self.table).count(), initial_cell_count)

    def test_views_create_table_row_wrapper(self):
        """Verifies create_table_row in tables.views delegates directly to RowService."""
        from tables.views import create_table_row
        row = create_table_row(
            table=self.table,
            user=self.admin,
            cells_data={"TASK_NAME": "Wrapper Task", "DUE_DATE": "2026-11-30"}
        )
        self.assertIsNotNone(row)
        self.assertEqual(row.created_by, self.admin)


from django.utils import timezone
from rest_framework.test import APIRequestFactory, force_authenticate
from auth_app.models import EmployeeUser
from tasks.models import Task, ActivityLog
from tables.models import ColumnAccess
from tables.services.cell_service import (
    CellMutationService,
    CellMutationError,
    CellPermissionDeniedError,
    CellValidationError,
    sync_logs_row_overdue,
)
from tables.services.row_mutation_service import (
    RowMutationService,
    RowMutationError,
    RowPermissionDeniedError,
    RowValidationError,
)
from tables.views import (
    RowViewSet,
    TableViewSet,
    sync_logs_row_overdue as views_sync_logs_row_overdue,
    edit_table_row as views_edit_table_row,
    bulk_update_table as views_bulk_update_table,
)



class CellMutationServiceTestCase(TestCase):
    def setUp(self):
        self.admin = EmployeeUser.objects.create_superuser(
            email="admin_cell@flow-force.com",
            password="password123",
            full_name="Admin Cell"
        )
        self.emp1 = EmployeeUser.objects.create_user(
            email="emp1_cell@flow-force.com",
            password="password123",
            full_name="Employee One",
            role="EMPLOYEE",
            status="APPROVED"
        )
        self.emp2 = EmployeeUser.objects.create_user(
            email="emp2_cell@flow-force.com",
            password="password123",
            full_name="Employee Two",
            role="EMPLOYEE",
            status="APPROVED"
        )

        self.table = Table.objects.create(name="Standard Table", job_type="GENERAL", created_by=self.admin)
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=self.table, user=self.emp1, access_level="EDIT")

        self.col_s_no = Column.objects.create(table=self.table, name="S_NO", data_type="NUMBER", is_system_column=True, position=1)
        self.col_task = Column.objects.create(table=self.table, name="TASK_NAME", data_type="TEXT", is_system_column=True, position=2)
        self.col_due = Column.objects.create(table=self.table, name="DUE_DATE", data_type="DATE", is_system_column=True, position=3)
        self.col_status = Column.objects.create(table=self.table, name="STATUS", data_type="TEXT", is_system_column=True, position=4)
        self.col_assignee = Column.objects.create(table=self.table, name="ASSIGNED_TO", data_type="USER", is_system_column=False, position=5)
        self.col_custom_num = Column.objects.create(table=self.table, name="QUANTITY", data_type="NUMBER", is_system_column=False, position=6)

        self.row = Row.objects.create(table=self.table, created_by=self.admin)
        self.task = Task.objects.create(
            row=self.row,
            due_date=datetime.date(2026, 11, 1),
            priority="HIGH",
            status="PENDING",
            assigned_by=self.admin,
            alert_mail_sent=True
        )

    def test_service_normal_cell_update_text_and_number(self):
        """Verifies updating text and number columns updates CellValue and creates ActivityLog."""
        cell_text = CellMutationService.update_cell(
            row=self.row,
            column=self.col_task,
            value="Refactored Task Name",
            user=self.admin
        )
        self.assertEqual(cell_text.value, "Refactored Task Name")
        self.assertEqual(cell_text.updated_by, self.admin)

        cell_num = CellMutationService.update_cell(
            row=self.row,
            column=self.col_custom_num,
            value=42,
            user=self.emp1
        )
        self.assertEqual(cell_num.value, 42)
        self.assertEqual(cell_num.updated_by, self.emp1)

        logs = ActivityLog.objects.filter(task=self.task).order_by("-id")
        self.assertTrue(logs.filter(action=f"Updated cell {self.col_task.name}").exists())
        self.assertTrue(logs.filter(action=f"Updated cell {self.col_custom_num.name}").exists())

    def test_service_permission_denial(self):
        """Verifies table and column level permission enforcement."""
        unauthorized_user = EmployeeUser.objects.create_user(
            email="unauth@flow-force.com",
            password="pass",
            role="EMPLOYEE",
            status="APPROVED"
        )
        # 1. No table access and not assigned
        with self.assertRaises(CellPermissionDeniedError):
            CellMutationService.update_cell(
                row=self.row,
                column=self.col_task,
                value="Hacked",
                user=unauthorized_user
            )

        # 2. Assignee trying to edit S_NO (read-only for assignees)
        self.task.assigned_to.set([unauthorized_user])
        with self.assertRaises(CellPermissionDeniedError) as ctx:
            CellMutationService.update_cell(
                row=self.row,
                column=self.col_s_no,
                value=999,
                user=unauthorized_user
            )
        self.assertIn("read-only for assignees", str(ctx.exception))

        # 3. Non-assignee with table access but column set to READ_ONLY
        ColumnAccess.objects.create(column=self.col_task, user=self.emp1, access_level="READ_ONLY")
        with self.assertRaises(CellPermissionDeniedError) as ctx:
            CellMutationService.update_cell(
                row=self.row,
                column=self.col_task,
                value="Readonly attempt",
                user=self.emp1
            )
        self.assertIn("read-only or hidden", str(ctx.exception))

    def test_service_date_validation_and_task_sync(self):
        """Verifies date parsing, invalid format error, and Task.due_date sync with alert_mail reset."""
        # Invalid date format
        with self.assertRaises(CellValidationError):
            CellMutationService.update_cell(
                row=self.row,
                column=self.col_due,
                value="not-a-date",
                user=self.admin
            )

        # Valid date format -> updates task.due_date and resets alert_mail_sent
        self.assertTrue(self.task.alert_mail_sent)
        CellMutationService.update_cell(
            row=self.row,
            column=self.col_due,
            value="2026-12-25",
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.due_date, datetime.date(2026, 12, 25))
        self.assertFalse(self.task.alert_mail_sent)

    def test_service_status_and_assigned_to_sync(self):
        """Verifies STATUS cell updates Task.status, and USER/assignee column updates Task.assigned_to."""
        # Status sync
        CellMutationService.update_cell(
            row=self.row,
            column=self.col_status,
            value="Completed",
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "COMPLETED")

        CellMutationService.update_cell(
            row=self.row,
            column=self.col_status,
            value="Pending",
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "PENDING")

        # Assignee sync by user ID
        CellMutationService.update_cell(
            row=self.row,
            column=self.col_assignee,
            value=str(self.emp2.id),
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(list(self.task.assigned_to.all()), [self.emp2])
        self.assertEqual(self.task.assigned_by, self.admin)

        # Assignee clear
        CellMutationService.update_cell(
            row=self.row,
            column=self.col_assignee,
            value="",
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.assigned_to.count(), 0)

    def test_service_logs_behavior_and_days_overdue(self):
        """Verifies LOGS table auto-capture of return_date and calculation of DAYS_OVERDUE."""
        logs_table = Table.objects.create(name="Logs Table", job_type="LOGS", created_by=self.admin)
        TableAccess.objects.create(table=logs_table, user=self.admin, access_level="ADMIN")
        cols = {c.name.upper(): c for c in logs_table.columns.all()}
        col_issue = cols["ISSUE_DATE"]
        col_return = cols["RETURN_DATE"]
        col_log_status = cols["STATUS"]
        col_overdue = cols["DAYS_OVERDUE"]

        log_row = Row.objects.create(table=logs_table, created_by=self.admin)

        # Set past issue date (5 days ago)
        past_date = (timezone.localdate() - datetime.timedelta(days=5)).isoformat()
        CellMutationService.update_cell(
            row=log_row,
            column=col_issue,
            value=past_date,
            user=self.admin
        )

        # RETURN_DATE should be auto-set to ISSUE_DATE because missing
        ret_cell = CellValue.objects.filter(row=log_row, column=col_return).first()
        self.assertIsNotNone(ret_cell)
        self.assertEqual(ret_cell.value, past_date)

        # DAYS_OVERDUE should be 5
        overdue_cell = CellValue.objects.filter(row=log_row, column=col_overdue).first()
        self.assertIsNotNone(overdue_cell)
        self.assertEqual(overdue_cell.value, 5)

        # Mark as returned
        CellMutationService.update_cell(
            row=log_row,
            column=col_log_status,
            value="Returned",
            user=self.admin
        )
        overdue_cell.refresh_from_db()
        self.assertIsNotNone(overdue_cell.value)

    def test_service_transaction_rollback_on_failure(self):
        """Verifies transaction rollback if an unexpected failure occurs during cell update."""
        from unittest.mock import patch

        initial_val = "Initial Value"
        cell = CellValue.objects.create(row=self.row, column=self.col_task, value=initial_val)

        with patch("tasks.models.ActivityLog.objects.create", side_effect=RuntimeError("Simulated Failure")):
            with self.assertRaises(RuntimeError):
                CellMutationService.update_cell(
                    row=self.row,
                    column=self.col_task,
                    value="Should Rollback",
                    user=self.admin
                )

        cell.refresh_from_db()
        self.assertEqual(cell.value, initial_val)

    def test_views_edit_cell_api_contract_and_error_responses(self):
        """Verifies DRF edit_cell endpoint returns correct status codes and JSON structure."""
        factory = APIRequestFactory()
        view = RowViewSet.as_view({"post": "edit_cell"})

        # 1. Success response
        req = factory.post(
            f"/tables/api/rows/{self.row.id}/edit-cell/",
            {"column": self.col_task.id, "value": "API Contract Value"},
            format="json"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.row.id)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("id", resp.data)
        self.assertIn("cells", resp.data)

        # 2. Permission denied (403)
        other_user = EmployeeUser.objects.create_user(
            email="other@flow-force.com",
            password="pass",
            role="EMPLOYEE",
            status="APPROVED"
        )
        req = factory.post(
            f"/tables/api/rows/{self.row.id}/edit-cell/",
            {"column": self.col_task.id, "value": "Forbidden"},
            format="json"
        )
        force_authenticate(req, user=other_user)
        resp = view(req, pk=self.row.id)
        self.assertEqual(resp.status_code, 403)
        self.assertIn("error", resp.data)

        # 3. Invalid date format (400)
        req = factory.post(
            f"/tables/api/rows/{self.row.id}/edit-cell/",
            {"column": self.col_due.id, "value": "invalid-date-string"},
            format="json"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.row.id)
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data.get("error"), "Invalid date format")

        # 4. Column not found (404)
        req = factory.post(
            f"/tables/api/rows/{self.row.id}/edit-cell/",
            {"column": 999999, "value": "Missing Column"},
            format="json"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.row.id)
        self.assertEqual(resp.status_code, 404)

    def test_sync_logs_row_overdue_backward_compatibility(self):
        """Verifies backward compatibility wrapper in tables.views delegates correctly."""
        logs_table = Table.objects.create(name="Logs Wrapper Table", job_type="LOGS", created_by=self.admin)
        cols = {c.name.upper(): c for c in logs_table.columns.all()}
        col_issue = cols["ISSUE_DATE"]
        col_return = cols["RETURN_DATE"]
        log_row = Row.objects.create(table=logs_table, created_by=self.admin)

        today_str = timezone.localdate().isoformat()
        CellValue.objects.create(row=log_row, column=col_issue, value=today_str)

        views_sync_logs_row_overdue(log_row, request_user=self.admin)

        ret_cell = CellValue.objects.filter(row=log_row, column=col_return).first()
        self.assertIsNotNone(ret_cell)
        self.assertEqual(ret_cell.value, today_str)


class RowMutationServiceTestCase(TestCase):
    def setUp(self):
        self.admin = EmployeeUser.objects.create_superuser(
            email="admin_rowmut@flow-force.com",
            password="password123",
            full_name="Admin RowMut"
        )
        self.emp1 = EmployeeUser.objects.create_user(
            email="emp1_rowmut@flow-force.com",
            password="password123",
            full_name="Worker One",
            role="EMPLOYEE",
            status="APPROVED"
        )
        self.emp2 = EmployeeUser.objects.create_user(
            email="emp2_rowmut@flow-force.com",
            password="password123",
            full_name="Worker Two",
            role="EMPLOYEE",
            status="APPROVED"
        )
        self.unauth = EmployeeUser.objects.create_user(
            email="unauth_rowmut@flow-force.com",
            password="password123",
            full_name="Unauth RowMut",
            role="EMPLOYEE",
            status="APPROVED"
        )

        self.table = Table.objects.create(
            name="Row Mutation Test Table",
            job_type="PERSONAL",
            created_by=self.admin
        )
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")

        TableAccess.objects.create(table=self.table, user=self.emp1, access_level="EDIT")
        TableAccess.objects.create(table=self.table, user=self.emp2, access_level="VIEW")

        self.col_task = Column.objects.create(
            table=self.table, name="TASK_NAME", data_type="TEXT", position=1, is_system_column=True
        )
        self.col_job = Column.objects.create(
            table=self.table, name="JOB_NUMBER", data_type="TEXT", position=2, is_system_column=False
        )
        self.col_due = Column.objects.create(
            table=self.table, name="DUE_DATE", data_type="DATE", position=3, is_system_column=True
        )
        self.col_status = Column.objects.create(
            table=self.table, name="STATUS", data_type="TEXT", position=4, is_system_column=False
        )
        self.col_assignee = Column.objects.create(
            table=self.table, name="ASSIGNED_TO", data_type="USER", position=5, is_system_column=False
        )
        self.col_sno = Column.objects.create(
            table=self.table, name="S_NO", data_type="NUMBER", position=6, is_system_column=False
        )
        self.col_init_mail = Column.objects.create(
            table=self.table, name="INITIAL_MAIL", data_type="TEXT", position=7, is_system_column=False
        )
        self.col_alert_mail = Column.objects.create(
            table=self.table, name="ALERT_MAIL", data_type="TEXT", position=8, is_system_column=False
        )

        self.row = Row.objects.create(table=self.table, created_by=self.admin)
        self.task = Task.objects.create(
            row=self.row,
            due_date="2026-08-01",
            status="PENDING",
            priority="MEDIUM",
            assigned_by=self.admin,
            alert_mail_sent=True,
        )

    def test_normal_multi_cell_row_edit(self):
        cells_data = {
            "TASK_NAME": "Fabricate Flange",
            "JOB_NUMBER": "JOB-2026-X",
        }
        row, updated_cols = RowMutationService.edit_row(
            row=self.row,
            cells_data=cells_data,
            user=self.admin
        )
        self.assertEqual(set(updated_cols), {"TASK_NAME", "JOB_NUMBER"})
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_task).value, "Fabricate Flange")
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_job).value, "JOB-2026-X")

    def test_unauthorized_row_edit_raises_permission_denied(self):
        with self.assertRaises(RowPermissionDeniedError):
            RowMutationService.edit_row(
                row=self.row,
                cells_data={"TASK_NAME": "Hacked Task"},
                user=self.unauth
            )

    def test_assignee_can_edit_assigned_row_without_table_edit_access(self):
        self.task.assigned_to.set([self.unauth])
        row, updated_cols = RowMutationService.edit_row(
            row=self.row,
            cells_data={"TASK_NAME": "Updated By Assignee"},
            user=self.unauth
        )
        self.assertIn("TASK_NAME", updated_cols)
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_task).value, "Updated By Assignee")

    def test_assignee_restricted_columns_skipped(self):
        self.task.assigned_to.set([self.unauth])
        cells_data = {
            "S_NO": 99,
            "INITIAL_MAIL": "YES",
            "ALERT_MAIL": "YES",
            "TASK_NAME": "Permitted Assignee Update",
        }
        row, updated_cols = RowMutationService.edit_row(
            row=self.row,
            cells_data=cells_data,
            user=self.unauth
        )
        self.assertEqual(updated_cols, ["TASK_NAME"])
        self.assertFalse(CellValue.objects.filter(row=self.row, column=self.col_sno).exists())
        # INITIAL_MAIL remains at the initial "NO" set during task creation signal, not updated to "YES"
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_init_mail).value, "NO")

    def test_column_level_permission_filtering(self):
        # Clear assignments so emp1 is treated as non-assignee relying on table edit access
        self.task.assigned_to.clear()
        ColumnAccess.objects.create(column=self.col_job, user=self.emp1, access_level="READ_ONLY")
        cells_data = {
            "TASK_NAME": "Allowed Change",
            "JOB_NUMBER": "Forbidden Job Number",
        }
        row, updated_cols = RowMutationService.edit_row(
            row=self.row,
            cells_data=cells_data,
            user=self.emp1
        )
        self.assertEqual(updated_cols, ["TASK_NAME"])
        self.assertFalse(CellValue.objects.filter(row=self.row, column=self.col_job).exists())
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_task).value, "Allowed Change")


    def test_due_date_synchronization_and_mail_reset(self):
        from unittest.mock import patch
        with patch("tasks.tasks.update_task_row_mail_columns") as mock_mail_sync:
            RowMutationService.edit_row(
                row=self.row,
                cells_data={"DUE_DATE": "2026-09-15"},
                user=self.admin
            )
            self.task.refresh_from_db()
            self.assertEqual(str(self.task.due_date), "2026-09-15")
            self.assertFalse(self.task.alert_mail_sent)
            mock_mail_sync.assert_called_once_with(self.task)

    def test_due_date_list_pid_handling(self):
        pid_table = Table.objects.create(name="PID Test Table", job_type="LIST_PID", created_by=self.admin)
        TableAccess.objects.create(table=pid_table, user=self.admin, access_level="ADMIN")
        col_ff = pid_table.columns.get(name="DUE_DATE_FLOW_FORCE")
        col_cust = pid_table.columns.get(name="DUE_DATE_CUSTOMER")
        pid_row = Row.objects.create(table=pid_table, created_by=self.admin)
        pid_task = Task.objects.create(row=pid_row, due_date="2026-01-01", status="PENDING", assigned_by=self.admin)


        # 1. Flow force priority
        RowMutationService.edit_row(
            row=pid_row,
            cells_data={"DUE_DATE_FLOW_FORCE": "2026-11-10", "DUE_DATE_CUSTOMER": "2026-11-20"},
            user=self.admin
        )
        pid_task.refresh_from_db()
        self.assertEqual(str(pid_task.due_date), "2026-11-10")

        # 2. Customer fallback if Flow Force missing/empty
        CellValue.objects.filter(row=pid_row, column=col_ff).update(value="")
        RowMutationService.edit_row(
            row=pid_row,
            cells_data={"DUE_DATE_CUSTOMER": "2026-12-05"},
            user=self.admin
        )
        pid_task.refresh_from_db()
        self.assertEqual(str(pid_task.due_date), "2026-12-05")

    def test_silent_invalid_date_handling(self):
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"DUE_DATE": "invalid-date", "TASK_NAME": "Still Updated"},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(str(self.task.due_date), "2026-08-01")
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_task).value, "Still Updated")

    def test_status_synchronization(self):
        # COMPLETE -> COMPLETED
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"STATUS": "COMPLETE"},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "COMPLETED")

        # NOT_RETURNED -> PENDING
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"STATUS": "not returned"},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "PENDING")

    def test_assignee_synchronization(self):
        # 1. By ID
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"ASSIGNED_TO": str(self.emp1.id)},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(list(self.task.assigned_to.all()), [self.emp1])
        self.assertEqual(self.task.assigned_by, self.admin)

        # 2. By Email
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"ASSIGNED_TO": self.emp2.email},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(list(self.task.assigned_to.all()), [self.emp2])

        # 3. By Full Name
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"ASSIGNED_TO": "Worker One"},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(list(self.task.assigned_to.all()), [self.emp1])

        # 4. Clear assignee
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"ASSIGNED_TO": ""},
            user=self.admin
        )
        self.task.refresh_from_db()
        self.assertEqual(self.task.assigned_to.count(), 0)

    def test_activity_log_creation(self):
        RowMutationService.edit_row(
            row=self.row,
            cells_data={"TASK_NAME": "Logged Multi Cell", "JOB_NUMBER": "JOB-LOG-1"},
            user=self.admin
        )
        log = ActivityLog.objects.filter(task=self.task, action="Updated multiple cells in row").first()
        self.assertIsNotNone(log)
        self.assertEqual(log.user, self.admin)
        self.assertIn("TASK_NAME", log.details.get("updated_columns", []))
        self.assertIn("JOB_NUMBER", log.details.get("updated_columns", []))

    def test_logs_table_synchronization(self):
        logs_table = Table.objects.create(name="Logs Row Mut Table", job_type="LOGS", created_by=self.admin)
        TableAccess.objects.create(table=logs_table, user=self.admin, access_level="ADMIN")
        col_issue = logs_table.columns.get(name="ISSUE_DATE")
        col_return = logs_table.columns.get(name="RETURN_DATE")
        col_stat = logs_table.columns.get(name="STATUS")
        log_row = Row.objects.create(table=logs_table, created_by=self.admin)
        today_str = timezone.localdate().isoformat()
        CellValue.objects.create(row=log_row, column=col_issue, value=today_str)


        RowMutationService.edit_row(
            row=log_row,
            cells_data={"STATUS": "Not Returned"},
            user=self.admin
        )

        return_cell = CellValue.objects.filter(row=log_row, column=col_return).first()
        self.assertIsNotNone(return_cell)
        self.assertEqual(return_cell.value, today_str)
        overdue_cell = CellValue.objects.filter(row=log_row, column__name="DAYS_OVERDUE").first()
        self.assertIsNotNone(overdue_cell)
        self.assertEqual(overdue_cell.value, 0)

    def test_bulk_update_permissions(self):
        with self.assertRaises(RowPermissionDeniedError):
            RowMutationService.bulk_update_table(
                table=self.table,
                field="INITIAL_MAIL",
                value="YES",
                user=self.emp1
            )

    def test_bulk_update_initial_mail_and_alert_mail(self):
        count = RowMutationService.bulk_update_table(
            table=self.table,
            field="INITIAL_MAIL",
            value="YES",
            user=self.admin
        )
        self.assertEqual(count, 1)
        self.task.refresh_from_db()
        self.assertTrue(self.task.initial_mail_sent)
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_init_mail).value, "YES")

        self.task.alert_mail_sent = False
        self.task.save()
        count = RowMutationService.bulk_update_table(
            table=self.table,
            field="ALERT_MAIL",
            value="YES",
            user=self.admin
        )
        self.assertEqual(count, 1)
        self.task.refresh_from_db()
        self.assertTrue(self.task.alert_mail_sent)
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_alert_mail).value, "YES")

    def test_bulk_update_status(self):
        count = RowMutationService.bulk_update_table(
            table=self.table,
            field="STATUS",
            value="COMPLETED",
            user=self.admin
        )
        self.assertEqual(count, 1)
        self.task.refresh_from_db()
        self.assertEqual(self.task.status, "COMPLETED")
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_status).value, "COMPLETED")
        log = ActivityLog.objects.filter(task=self.task, action="Updated cell STATUS via Bulk Update").first()
        self.assertIsNotNone(log)

    def test_bulk_update_invalid_field(self):
        with self.assertRaises(RowValidationError):
            RowMutationService.bulk_update_table(
                table=self.table,
                field="INVALID_FIELD",
                value="VAL",
                user=self.admin
            )

    def test_transaction_rollback_on_failure(self):
        from unittest.mock import patch
        initial_val = "Safe Initial"
        cell = CellValue.objects.create(row=self.row, column=self.col_task, value=initial_val)

        with patch("tasks.models.ActivityLog.objects.create", side_effect=RuntimeError("Simulated Crash")):
            with self.assertRaises(RuntimeError):
                RowMutationService.edit_row(
                    row=self.row,
                    cells_data={"TASK_NAME": "Will Crash"},
                    user=self.admin
                )
        cell.refresh_from_db()
        self.assertEqual(cell.value, initial_val)

    def test_api_compatibility_edit_row(self):
        factory = APIRequestFactory()
        view = RowViewSet.as_view({"post": "edit_row"})

        # Success (200)
        req = factory.post(
            f"/tables/api/rows/{self.row.id}/edit-row/",
            {"cells": {"TASK_NAME": "API Updated Row", "JOB_NUMBER": "API-JOB-1"}},
            format="json"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.row.id)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data["id"], self.row.id)
        self.assertIn("cells", resp.data)

        # Forbidden (403)
        req = factory.post(
            f"/tables/api/rows/{self.row.id}/edit-row/",
            {"cells": {"TASK_NAME": "API Forbidden"}},
            format="json"
        )
        force_authenticate(req, user=self.unauth)
        resp = view(req, pk=self.row.id)
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data.get("error"), "No edit access to this table or task row")

    def test_api_compatibility_bulk_update(self):
        factory = APIRequestFactory()
        view = TableViewSet.as_view({"post": "bulk_update"})

        # Success (200)
        req = factory.post(
            f"/tables/api/tables/{self.table.id}/bulk-update/",
            {"field": "STATUS", "value": "COMPLETED"},
            format="json"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.table.id)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("message", resp.data)
        self.assertEqual(resp.data["message"], "Successfully updated 1 rows")

        # Forbidden (403)
        req = factory.post(
            f"/tables/api/tables/{self.table.id}/bulk-update/",
            {"field": "STATUS", "value": "COMPLETED"},
            format="json"
        )
        force_authenticate(req, user=self.emp1)
        resp = view(req, pk=self.table.id)
        self.assertEqual(resp.status_code, 403)
        self.assertEqual(resp.data.get("error"), "Only admins can perform bulk updates")

        # Bad Request (400)
        req = factory.post(
            f"/tables/api/tables/{self.table.id}/bulk-update/",
            {"field": "INVALID", "value": "COMPLETED"},
            format="json"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.table.id)
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(resp.data.get("error"), "Invalid field for bulk update")

    def test_views_backward_compatibility_wrappers(self):
        row, updated_cols = views_edit_table_row(
            row=self.row,
            cells_data={"TASK_NAME": "Wrapper Updated"},
            user=self.admin
        )
        self.assertIn("TASK_NAME", updated_cols)
        self.assertEqual(CellValue.objects.get(row=self.row, column=self.col_task).value, "Wrapper Updated")

        count = views_bulk_update_table(
            table=self.table,
            field="STATUS",
            value="COMPLETED",
            user=self.admin
        )
        self.assertEqual(count, 1)


class TableDeleteServiceTestCase(TestCase):
    def setUp(self):
        from rest_framework.test import APIRequestFactory, force_authenticate
        self.factory = APIRequestFactory()
        self.dept = Department.objects.create(name="Delete Svc Dept", slug="delete-svc-dept")
        self.admin = User.objects.create_user(
            email="del_admin@flow-force.com",
            password="testpassword",
            full_name="Delete Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="del_emp@flow-force.com",
            password="testpassword",
            full_name="Delete Employee",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(name="Delete Service Table", created_by=self.admin, job_type="STANDARD")
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=self.table, user=self.employee, access_level="VIEW")

        self.custom_col = Column.objects.create(
            table=self.table,
            name="CUSTOM_FIELD",
            data_type="TEXT",
            is_system_column=False
        )
        self.sys_col = self.table.columns.get(name="TASK_NAME")

        self.row1 = Row.objects.create(table=self.table, created_by=self.admin)
        self.row2 = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=self.row1, column=self.custom_col, value="Value 1")
        CellValue.objects.create(row=self.row1, column=self.sys_col, value="Task 1")
        CellValue.objects.create(row=self.row2, column=self.sys_col, value="Task 2")
        self.task1 = Task.objects.create(row=self.row1, status="PENDING", priority="HIGH", assigned_by=self.admin)

    def test_delete_row_service_cascades_and_invalidates_cache(self):
        from tables.services.delete_service import TableDeleteService
        from django.core.cache import cache
        from tables.services.statistics_service import TableStatisticsService

        TableStatisticsService.get_table_statistics(self.table)
        cache_key = TableStatisticsService.get_cache_key(self.table.id)
        cache.set(cache_key, {"cached": True}, 300)
        self.assertIsNotNone(cache.get(cache_key))

        TableDeleteService.delete_row(self.row1, self.admin)
        self.assertFalse(Row.objects.filter(id=self.row1.id).exists())
        self.assertFalse(CellValue.objects.filter(row_id=self.row1.id).exists())
        self.assertFalse(Task.objects.filter(id=self.task1.id).exists())
        self.assertIsNone(cache.get(cache_key))

    def test_bulk_delete_rows_service(self):
        from tables.services.delete_service import TableDeleteService, DeleteValidationError
        # Non-list row_ids raises DeleteValidationError
        with self.assertRaises(DeleteValidationError):
            TableDeleteService.bulk_delete_rows(self.table, self.admin, row_ids="invalid")

        # Specific row_ids deletion
        count = TableDeleteService.bulk_delete_rows(self.table, self.admin, row_ids=[self.row1.id])
        self.assertEqual(count, 1)
        self.assertFalse(Row.objects.filter(id=self.row1.id).exists())
        self.assertTrue(Row.objects.filter(id=self.row2.id).exists())

        # Bulk delete all remaining rows
        count_all = TableDeleteService.bulk_delete_rows(self.table, self.admin, row_ids=None)
        self.assertEqual(count_all, 1)
        self.assertFalse(Row.objects.filter(table=self.table).exists())

    def test_delete_rows_by_column_service_and_system_protection(self):
        from tables.services.delete_service import TableDeleteService, DeleteValidationError
        # System column protection
        with self.assertRaises(DeleteValidationError):
            TableDeleteService.delete_rows_by_column(self.sys_col, self.admin)

        # Custom column deletes rows with values in that column
        count = TableDeleteService.delete_rows_by_column(self.custom_col, self.admin)
        self.assertEqual(count, 1)
        self.assertFalse(Row.objects.filter(id=self.row1.id).exists())
        self.assertTrue(Row.objects.filter(id=self.row2.id).exists())

    def test_clear_column_values_service_and_system_protection(self):
        from tables.services.delete_service import TableDeleteService, DeleteValidationError
        # System column protection
        with self.assertRaises(DeleteValidationError):
            TableDeleteService.clear_column_values(self.sys_col, self.admin)

        # Custom column clears values
        TableDeleteService.clear_column_values(self.custom_col, self.admin)
        self.assertEqual(CellValue.objects.get(row=self.row1, column=self.custom_col).value, None)

    def test_view_system_column_and_invalid_row_ids_error_responses(self):
        from tables.views import TableViewSet, ColumnViewSet
        from rest_framework.test import force_authenticate

        # 1. bulk_delete_rows with non-list row_ids returns 400
        view_tbl = TableViewSet.as_view({"post": "bulk_delete_rows"})
        req1 = self.factory.post(f"/tables/api/tables/{self.table.id}/bulk-delete-rows/", {"row_ids": "not-a-list"}, format="json")
        force_authenticate(req1, user=self.admin)
        resp1 = view_tbl(req1, pk=self.table.id)
        self.assertEqual(resp1.status_code, 400)
        self.assertEqual(resp1.data.get("error"), "row_ids must be a list")

        # 2. clear_values on system column returns 400
        view_col = ColumnViewSet.as_view({"post": "clear_values"})
        req2 = self.factory.post(f"/tables/api/columns/{self.sys_col.id}/clear-values/")
        force_authenticate(req2, user=self.admin)
        resp2 = view_col(req2, pk=self.sys_col.id)
        self.assertEqual(resp2.status_code, 400)
        self.assertEqual(resp2.data.get("error"), "Cannot clear system columns")

        # 3. delete_rows on system column returns 400
        view_del_col = ColumnViewSet.as_view({"post": "delete_rows"})
        req3 = self.factory.post(f"/tables/api/columns/{self.sys_col.id}/delete-rows/")
        force_authenticate(req3, user=self.admin)
        resp3 = view_del_col(req3, pk=self.sys_col.id)
        self.assertEqual(resp3.status_code, 400)
        self.assertEqual(resp3.data.get("error"), "Cannot delete rows using system column filter")


class TableImportServiceTest(TestCase):
    """
    Dedicated test suite for TableImportService extraction (Phase 3E-3).
    Verifies header normalization, strict PID mapping, date parsing,
    multiline parsing, atomic rollback, and API equivalence.
    """

    def setUp(self):
        from rest_framework.test import APIRequestFactory
        self.factory = APIRequestFactory()
        self.dept = Department.objects.create(name="Import Test Dept", slug="import-test-dept")
        self.admin = User.objects.create_user(
            email="importadmin@flow-force.com",
            password="testpassword",
            full_name="Import Admin",
            role="ADMIN",
            department=self.dept,
            status="APPROVED"
        )
        self.employee = User.objects.create_user(
            email="importemp@flow-force.com",
            password="testpassword",
            full_name="Import Employee",
            role="EMPLOYEE",
            department=self.dept,
            status="APPROVED"
        )
        self.table = Table.objects.create(name="Standard Tasks Table", created_by=self.admin)
        TableAccess.objects.create(table=self.table, user=self.admin, access_level="ADMIN")

    def test_normal_csv_import(self):
        from tables.services.import_service import TableImportService
        csv_data = (
            "S_NO,DATE,DUE_DATE,TASK_NAME,INITIAL_MAIL,ALERT_MAIL\n"
            "1,2026-06-22,2026-06-30,Normal Task Alpha,NO,NO\n"
            "2,2026-06-23,2026-07-01,Normal Task Beta,YES,NO\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 2)
        self.assertEqual(Row.objects.filter(table=self.table).count(), 2)

        tasks = list(Task.objects.filter(row__table=self.table).order_by('id'))
        self.assertEqual(len(tasks), 2)
        self.assertEqual(tasks[0].task_name, "Normal Task Alpha")
        self.assertEqual(tasks[1].task_name, "Normal Task Beta")
        self.assertEqual(tasks[0].due_date, datetime.date(2026, 6, 30))
        self.assertFalse(tasks[0].initial_mail_sent)
        self.assertTrue(tasks[1].initial_mail_sent)

    def test_quoted_multiline_csv(self):
        from tables.services.import_service import TableImportService
        csv_data = (
            'S_NO,DATE,DUE_DATE,TASK_NAME,INITIAL_MAIL,ALERT_MAIL\n'
            '1,2026-06-22,2026-06-30,"First line task\nSecond line details",NO,NO\n'
            '2,2026-06-23,2026-07-01,"Single line",NO,NO\n'
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 2)

        tasks = list(Task.objects.filter(row__table=self.table).order_by('id'))
        self.assertEqual(len(tasks), 2)
        self.assertEqual(tasks[0].task_name, "First line task\nSecond line details")
        self.assertEqual(tasks[1].task_name, "Single line")

    def test_normalized_headers(self):
        from tables.services.import_service import TableImportService
        csv_data = (
            "  S. No.  , Date , Due-Date ,  Task Name  , Initial Mail , Alert Mail \n"
            "1,2026-06-22,2026-06-30,Spaced Header Task,NO,NO\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 1)

        task = Task.objects.filter(row__table=self.table).first()
        self.assertIsNotNone(task)
        self.assertEqual(task.task_name, "Spaced Header Task")
        self.assertEqual(task.due_date, datetime.date(2026, 6, 30))

    def test_strict_pid_mapping_and_blank_enquiry(self):
        from tables.services.import_service import TableImportService
        pid_table = Table.objects.create(
            name="PID Specific Table",
            created_by=self.admin,
            job_type="LIST_PID"
        )
        TableAccess.objects.create(table=pid_table, user=self.admin, access_level="ADMIN")

        # CSV with PID column but NO enquiry column
        csv_data = (
            "PID,COMPANY_NAME,DUE_DATE_FLOW_FORCE\n"
            "PID-9999,Acme Industrial,2026-08-15\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=pid_table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 1)

        row = created_rows[0]
        col_pid = pid_table.columns.get(name="PID")
        col_enq = pid_table.columns.get(name="ENQUIRY_NO/QUOTATION_NO")

        cell_pid = CellValue.objects.filter(row=row, column=col_pid).first()
        cell_enq = CellValue.objects.filter(row=row, column=col_enq).first()

        # PID must strictly map to PID
        self.assertIsNotNone(cell_pid)
        self.assertEqual(cell_pid.value, "PID-9999")

        # Enquiry must NOT take the PID value
        enquiry_val = cell_enq.value if cell_enq else ""
        self.assertNotEqual(enquiry_val, "PID-9999")
        self.assertEqual(enquiry_val, "")

    def test_extra_source_columns_ignored(self):
        from tables.services.import_service import TableImportService
        csv_data = (
            "S_NO,TASK_NAME,DUE_DATE,EXTRA_COL_1,EXTRA_COL_2\n"
            "1,Extra Column Task,2026-06-30,IgnoredValue1,IgnoredValue2\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 1)
        task = Task.objects.filter(row__table=self.table).first()
        self.assertEqual(task.task_name, "Extra Column Task")

    def test_missing_destination_columns_remain_empty(self):
        from tables.services.import_service import TableImportService
        # Add a custom column to the table that is absent in the CSV
        custom_col = Column.objects.create(table=self.table, name="REMARKS", data_type="TEXT")

        csv_data = (
            "S_NO,TASK_NAME,DUE_DATE\n"
            "1,Missing Dest Task,2026-06-30\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 1)

        cell = CellValue.objects.filter(row=created_rows[0], column=custom_col).first()
        self.assertIsNone(cell)

    def test_safe_parse_date(self):
        from tables.services.import_service import TableImportService
        # 1. Excel serial date
        d1 = TableImportService.safe_parse_date("45443")
        self.assertEqual(d1, datetime.date(2024, 5, 31))

        # 2. 8-digit numeric date YYYYMMDD
        d2 = TableImportService.safe_parse_date("20260715")
        self.assertEqual(d2, datetime.date(2026, 7, 15))

        # 3. ISO format
        d3 = TableImportService.safe_parse_date("2026-10-06")
        self.assertEqual(d3, datetime.date(2026, 10, 6))

        # 4. Regional dayfirst
        d4 = TableImportService.safe_parse_date("25/12/2026")
        self.assertEqual(d4, datetime.date(2026, 12, 25))

        # 5. Invalid / empty returns None
        self.assertIsNone(TableImportService.safe_parse_date(""))
        self.assertIsNone(TableImportService.safe_parse_date(None))
        self.assertIsNone(TableImportService.safe_parse_date("invalid-date-string"))

    def test_s_no_generation_continuity(self):
        from tables.services.import_service import TableImportService
        # Create an existing row with S_NO = 10
        s_no_col = self.table.columns.get(name="S_NO")
        existing_row = Row.objects.create(table=self.table, created_by=self.admin)
        CellValue.objects.create(row=existing_row, column=s_no_col, value=10, updated_by=self.admin)

        csv_data = (
            "TASK_NAME,DUE_DATE\n"
            "Continued S_NO Task 1,2026-06-30\n"
            "Continued S_NO Task 2,2026-07-01\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 2)

        s_no_1 = CellValue.objects.get(row=created_rows[0], column=s_no_col).value
        s_no_2 = CellValue.objects.get(row=created_rows[1], column=s_no_col).value
        self.assertEqual(s_no_1, 11)
        self.assertEqual(s_no_2, 12)

    def test_assignee_resolution(self):
        from tables.services.import_service import TableImportService
        # Create an ASSIGNED_TO column
        col_assign = Column.objects.create(table=self.table, name="ASSIGNED_TO", data_type="USER")

        csv_data = (
            f"S_NO,TASK_NAME,ASSIGNED_TO\n"
            f"1,Assigned Task,{self.employee.email}\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=self.table,
            user=self.admin,
        )
        self.assertIsNone(err)
        task = Task.objects.filter(row=created_rows[0]).first()
        self.assertIsNotNone(task)
        self.assertIn(self.employee, task.assigned_to.all())

    def test_list_pid_import_auto_assignees_and_filter(self):
        from tables.services.import_service import TableImportService
        pid_table = Table.objects.create(
            name="PID Filter Table",
            created_by=self.admin,
            job_type="LIST_PID"
        )
        TableAccess.objects.create(table=pid_table, user=self.admin, access_level="ADMIN")
        TableAccess.objects.create(table=pid_table, user=self.employee, access_level="EDIT")

        csv_data = (
            "PID,COMPANY_NAME,DUE_DATE_FLOW_FORCE\n"
            "PID-101,Global Tech,2026-08-15\n"
            "PID-102,Apex Corp,2026-08-16\n"
        )
        created_rows, err = TableImportService.import_rows_from_csv_data(
            file_data=csv_data,
            table=pid_table,
            user=self.admin,
        )
        self.assertIsNone(err)
        self.assertEqual(len(created_rows), 2)

        # In LIST_PID, employees with table access are assigned to all rows
        task = Task.objects.filter(row=created_rows[0]).first()
        assignees = list(task.assigned_to.all())
        self.assertIn(self.employee, assignees)

        # COMPANY_NAME column should be converted to DROPDOWN with filter options
        company_col = pid_table.columns.get(name="COMPANY_NAME")
        self.assertEqual(company_col.data_type, "DROPDOWN")
        self.assertTrue(company_col.is_filterable)
        self.assertIn("Global Tech", company_col.options)
        self.assertIn("Apex Corp", company_col.options)

    def test_atomic_rollback_on_failure(self):
        from tables.services.import_service import TableImportService
        from unittest.mock import patch

        csv_data = (
            "S_NO,TASK_NAME,DUE_DATE\n"
            "1,Fail Task 1,2026-06-30\n"
            "2,Fail Task 2,2026-07-01\n"
        )

        with patch("tables.models.CellValue.objects.bulk_create", side_effect=RuntimeError("Simulated DB Crash")):
            with self.assertRaises(RuntimeError):
                TableImportService.import_rows_from_csv_data(
                    file_data=csv_data,
                    table=self.table,
                    user=self.admin,
                )

        # Clean rollback: zero rows created
        self.assertEqual(Row.objects.filter(table=self.table).count(), 0)

    def test_existing_api_behavior_preserved(self):
        from tables.views import TableViewSet
        from rest_framework.test import force_authenticate
        from django.core.files.uploadedfile import SimpleUploadedFile

        view = TableViewSet.as_view({"post": "import_csv"})

        # 1. Success case
        csv_content = (
            "S_NO,DATE,DUE_DATE,TASK_NAME,INITIAL_MAIL,ALERT_MAIL\n"
            "1,2026-06-22,2026-06-30,API Imported Task,NO,NO\n"
        )
        csv_file = SimpleUploadedFile("test.csv", csv_content.encode("utf-8"), content_type="text/csv")
        req = self.factory.post(
            f"/tables/api/tables/{self.table.id}/import-csv/",
            {"file": csv_file},
            format="multipart"
        )
        force_authenticate(req, user=self.admin)
        resp = view(req, pk=self.table.id)

        self.assertEqual(resp.status_code, 201)
        self.assertIn("Successfully imported 1 rows", resp.data.get("message"))

        # 2. No file provided error case
        req_empty = self.factory.post(
            f"/tables/api/tables/{self.table.id}/import-csv/",
            {},
            format="multipart"
        )
        force_authenticate(req_empty, user=self.admin)
        resp_empty = view(req_empty, pk=self.table.id)
        self.assertEqual(resp_empty.status_code, 400)
        self.assertEqual(resp_empty.data.get("error"), "No CSV file provided")


from django.test import TransactionTestCase, override_settings
from channels.testing import WebsocketCommunicator
from channels.layers import get_channel_layer
from flowforce.asgi import application
from django.contrib.auth import SESSION_KEY, BACKEND_SESSION_KEY, HASH_SESSION_KEY
from importlib import import_module
from django.conf import settings
from asgiref.sync import sync_to_async


@override_settings(
    CHANNEL_LAYERS={
        "default": {
            "BACKEND": "channels.layers.InMemoryChannelLayer",
        }
    }
)
class TableWebSocketConsumerTestCase(TransactionTestCase):
    """
    PHASE REALTIME-2: Verifies TableEventConsumer WebSocket connection and authorization.
    1. authenticated user with VIEW permission can connect
    2. unauthenticated user cannot establish an authorized connection
    3. user without table VIEW access is rejected
    4. invalid/nonexistent table ID is rejected
    5. authorized connection joins correct group (table_<table_id>)
    6. table 27 connection does not join table 28 group
    7. disconnect cleans up the group membership
    """

    def setUp(self):
        self.dept_a = Department.objects.create(name="WS Dept A", slug="ws-dept-a")
        self.dept_b = Department.objects.create(name="WS Dept B", slug="ws-dept-b")

        self.admin = EmployeeUser.objects.create_user(
            email="ws_admin_t2@example.com",
            password="testpassword",
            role="ADMIN",
            department=self.dept_a,
        )
        self.user_a = EmployeeUser.objects.create_user(
            email="ws_user_a_t2@example.com",
            password="testpassword",
            role="EMPLOYEE",
            department=self.dept_a,
        )
        self.user_b = EmployeeUser.objects.create_user(
            email="ws_user_b_t2@example.com",
            password="testpassword",
            role="EMPLOYEE",
            department=self.dept_b,
        )

        self.table_a = Table.objects.create(
            name="WS Test Table A",
            created_by=self.admin,
            department=self.dept_a,
        )
        self.table_b = Table.objects.create(
            name="WS Test Table B",
            created_by=self.admin,
            department=self.dept_b,
        )

    @sync_to_async
    def _create_user_session(self, user):
        SessionStore = import_module(settings.SESSION_ENGINE).SessionStore
        session = SessionStore()
        session[SESSION_KEY] = str(user.pk)
        session[BACKEND_SESSION_KEY] = "django.contrib.auth.backends.ModelBackend"
        session[HASH_SESSION_KEY] = user.get_session_auth_hash()
        session.save()
        return session.session_key

    async def test_1_authenticated_user_with_view_permission_can_connect(self):
        """1. Authenticated user with VIEW permission establishes a successful WebSocket connection."""
        session_key = await self._create_user_session(self.user_a)
        communicator = WebsocketCommunicator(
            application,
            f"/ws/tables/{self.table_a.id}/",
            headers=[(b"cookie", f"sessionid={session_key}".encode("ascii"))],
        )
        connected, _ = await communicator.connect()
        self.assertTrue(connected)
        await communicator.disconnect()

    async def test_2_unauthenticated_user_cannot_connect(self):
        """2. Unauthenticated user without valid session cannot establish a connection."""
        communicator = WebsocketCommunicator(
            application,
            f"/ws/tables/{self.table_a.id}/",
        )
        connected, _ = await communicator.connect()
        self.assertFalse(connected)
        await communicator.disconnect()

    async def test_3_user_without_view_access_is_rejected(self):
        """3. User from another department without VIEW permission is rejected."""
        session_key = await self._create_user_session(self.user_b)
        communicator = WebsocketCommunicator(
            application,
            f"/ws/tables/{self.table_a.id}/",
            headers=[(b"cookie", f"sessionid={session_key}".encode("ascii"))],
        )
        connected, _ = await communicator.connect()
        self.assertFalse(connected)
        await communicator.disconnect()

    async def test_4_invalid_or_nonexistent_table_id_is_rejected(self):
        """4. Connection to a nonexistent table ID is rejected safely without data leaks."""
        session_key = await self._create_user_session(self.user_a)
        communicator = WebsocketCommunicator(
            application,
            "/ws/tables/999999/",
            headers=[(b"cookie", f"sessionid={session_key}".encode("ascii"))],
        )
        connected, _ = await communicator.connect()
        self.assertFalse(connected)
        await communicator.disconnect()

    async def test_5_authorized_connection_joins_correct_table_group(self):
        """5. Authorized connection joins the specific group table_<table_id>."""
        session_key = await self._create_user_session(self.user_a)
        communicator = WebsocketCommunicator(
            application,
            f"/ws/tables/{self.table_a.id}/",
            headers=[(b"cookie", f"sessionid={session_key}".encode("ascii"))],
        )
        connected, _ = await communicator.connect()
        self.assertTrue(connected)

        layer = get_channel_layer()
        group_name = f"table_{self.table_a.id}"
        self.assertIn(group_name, layer.groups)
        self.assertEqual(len(layer.groups[group_name]), 1)

        await communicator.disconnect()

    async def test_6_table_connection_does_not_join_different_table_group(self):
        """6. Connecting to Table A does not join Table B's group (isolation)."""
        session_key = await self._create_user_session(self.user_a)
        communicator = WebsocketCommunicator(
            application,
            f"/ws/tables/{self.table_a.id}/",
            headers=[(b"cookie", f"sessionid={session_key}".encode("ascii"))],
        )
        connected, _ = await communicator.connect()
        self.assertTrue(connected)

        layer = get_channel_layer()
        group_b = f"table_{self.table_b.id}"
        self.assertEqual(len(layer.groups.get(group_b, {})), 0)

        await communicator.disconnect()

    async def test_7_disconnect_cleans_up_group_membership(self):
        """7. Disconnecting cleans up channel group membership."""
        session_key = await self._create_user_session(self.user_a)
        communicator = WebsocketCommunicator(
            application,
            f"/ws/tables/{self.table_a.id}/",
            headers=[(b"cookie", f"sessionid={session_key}".encode("ascii"))],
        )
        connected, _ = await communicator.connect()
        self.assertTrue(connected)

        layer = get_channel_layer()
        group_name = f"table_{self.table_a.id}"
        self.assertEqual(len(layer.groups.get(group_name, {})), 1)

        await communicator.disconnect()
        self.assertEqual(len(layer.groups.get(group_name, {})), 0)

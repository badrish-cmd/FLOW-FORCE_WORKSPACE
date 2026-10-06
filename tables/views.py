from rest_framework import viewsets, permissions, status
from rest_framework.response import Response
from rest_framework.decorators import action
from django.shortcuts import get_object_or_404
from django.db import transaction
from django.utils import timezone
import logging
from datetime import datetime

logger = logging.getLogger(__name__)

from .models import Table, Column, Row, CellValue, TableAccess, ColumnAccess
from .serializers import (
    TableSerializer, ColumnSerializer, RowSerializer,
    CellValueSerializer, TableAccessSerializer, ColumnAccessSerializer
)
from .permissions import get_accessible_tables, has_table_access, get_column_access_level
from tasks.models import Task, ActivityLog
from auth_app.models import EmployeeUser
from .services.statistics_service import TableStatisticsService
from .services.row_service import RowService, RowCreationValidationError
from .services.cell_service import (
    CellMutationService,
    CellMutationError,
    CellPermissionDeniedError,
    CellValidationError,
    sync_logs_row_overdue as _sync_logs_row_overdue,
)
from .services.row_mutation_service import (
    RowMutationService,
    RowMutationError,
    RowPermissionDeniedError,
    RowValidationError,
)
from .services.duplicate_service import TableDuplicateService
from .services.delete_service import TableDeleteService, DeleteValidationError
from .services.import_service import TableImportService

def import_rows_from_csv_data(file_data, table, request_user, header_row=None, data_row=None):
    """
    Backward-compatibility wrapper for importing rows from CSV data.
    Delegates directly to TableImportService.import_rows_from_csv_data.
    """
    return TableImportService.import_rows_from_csv_data(
        file_data=file_data,
        table=table,
        user=request_user,
        header_row=header_row,
        data_row=data_row,
    )

def get_table_statistics(table, today_date=None, use_cache=True):
    """
    Backward-compatibility wrapper for table statistics calculation.
    Delegates directly to TableStatisticsService.get_table_statistics.
    """
    return TableStatisticsService.get_table_statistics(table, today_date=today_date, use_cache=use_cache)

def create_table_row(table, user, cells_data=None, assigned_to_ids=None):
    """
    Backward-compatibility wrapper for row creation.
    Delegates directly to RowService.create_row.
    """
    return RowService.create_row(table=table, user=user, cells_data=cells_data, assigned_to_ids=assigned_to_ids)


def sync_logs_row_overdue(row, request_user=None):
    """
    Backward-compatibility wrapper for LOGS overdue synchronization.
    Delegates directly to CellMutationService / cell_service.sync_logs_row_overdue.
    """
    return _sync_logs_row_overdue(row, request_user=request_user)


def edit_table_row(row, cells_data, user, check_permissions=True):
    """
    Backward-compatibility wrapper for row mutation.
    Delegates directly to RowMutationService.edit_row.
    """
    return RowMutationService.edit_row(row=row, cells_data=cells_data, user=user, check_permissions=check_permissions)


def bulk_update_table(table, field, value, user, check_permissions=True):
    """
    Backward-compatibility wrapper for table bulk update.
    Delegates directly to RowMutationService.bulk_update_table.
    """
    return RowMutationService.bulk_update_table(table=table, field=field, value=value, user=user, check_permissions=check_permissions)


def get_filtered_table_rows(table, query_params, user=None):
    """
    Returns a filtered, ordered queryset of rows for a given table based on query parameters.
    Shared across RowViewSet.get_queryset and TableViewSet.export_excel.
    """
    if user and not has_table_access(user, table, "VIEW"):
        return Row.objects.none()

    from django.db.models import Prefetch, Q, Value, TextField, Exists, OuterRef
    from django.db.models.functions import Cast, Lower, Replace
    import datetime

    queryset = Row.objects.filter(
        table=table, is_archived=False
    ).select_related(
        'created_by', 'task', 'task__assigned_by'
    ).prefetch_related(
        Prefetch('cells', queryset=CellValue.objects.select_related('column', 'updated_by')),
        'task__assigned_to'
    )

    # 1. Filter by task_id
    task_id = query_params.get("task_id")
    if task_id:
        queryset = queryset.filter(task__id=task_id)

    # 2. General Row Search Filter (irrespective of case-sensitivity and inline spaces)
    search = query_params.get("search")
    if search and search.strip():
        search_stripped = search.strip()
        clean_search = search_stripped.lower().replace(" ", "")

        if clean_search:
            # Correlated subquery on CellValue for this row only:
            # Avoids loading entire table cells, eliminates redundant inner joins to Row,
            # and short-circuits via LIMIT 1 (EXISTS) as soon as the first matching cell is found.
            cell_match = CellValue.objects.filter(
                row=OuterRef('pk')
            ).annotate(
                clean_val=Lower(Replace(Cast('value', TextField()), Value(' '), Value(''), output_field=TextField()))
            ).filter(
                clean_val__contains=clean_search
            )

            # Pre-match enum choices for task status and priority in Python to utilize database B-tree indexes
            from tasks.models import Task
            search_norm = clean_search.replace("_", "")
            matching_statuses = [code for code, _ in Task.STATUS_CHOICES if search_norm in code.lower().replace("_", "")]
            matching_priorities = [code for code, _ in Task.PRIORITY_CHOICES if search_norm in code.lower().replace("_", "")]

            task_cond = Q()
            if matching_statuses:
                task_cond |= Q(task__status__in=matching_statuses)
            if matching_priorities:
                task_cond |= Q(task__priority__in=matching_priorities)

            search_filter = Q(Exists(cell_match)) | task_cond
            queryset = queryset.filter(search_filter)

    # 3. Custom dynamic column filters (col_<id>)
    for key, val in query_params.items():
        if key.startswith("col_") and val:
            try:
                col_id = int(key.replace("col_", ""))
                queryset = queryset.filter(cells__column_id=col_id, cells__value=val)
            except ValueError:
                pass

    # 4. Filter by PID
    pid = query_params.get("pid")
    if pid:
        queryset = queryset.filter(
            Q(cells__column__name__iexact='PID', cells__value=pid) |
            Q(cells__column__name__iexact='PID', cells__value__icontains=pid)
        ).distinct()

    # 5. Filter by Year
    year = query_params.get("year")
    if year:
        try:
            year_int = int(str(year).strip())
            queryset = queryset.filter(task__due_date__year=year_int)
        except (ValueError, TypeError):
            pass

    # 6. Filter by Month
    month = query_params.get("month")
    if month:
        try:
            month_int = int(str(month).strip())
            queryset = queryset.filter(task__due_date__month=month_int)
        except (ValueError, TypeError):
            pass

    # 7. Filter by Due Status
    due = query_params.get("due")
    if due:
        today = timezone.localdate()
        if due == "today":
            queryset = queryset.filter(task__due_date=today)
        elif due == "this_week":
            monday = today - datetime.timedelta(days=today.weekday())
            sunday = monday + datetime.timedelta(days=6)
            queryset = queryset.filter(task__due_date__range=[monday, sunday])

    # 8. Apply Sorting (PostgreSQL-safe without unsafe DateField casts of JSONB)
    sort_by = query_params.get("sort_by")
    sort_dir = query_params.get("sort_dir", "asc")
    if sort_by:
        if sort_by.lower() == 'date':
            sort_by = 'date_assigned'
        elif sort_by.lower() in ['enquiry_no', 'enquiry_number', 'enquiry_no/quotation_no']:
            sort_by = 'enquiry_no'
        
        if sort_by == 'enquiry_no':
            from django.db.models import Subquery, OuterRef
            enquiry_col = Column.objects.filter(table=table, name__iexact="ENQUIRY_NO/QUOTATION_NO").first()
            if not enquiry_col:
                enquiry_col = Column.objects.filter(table=table, name__icontains="ENQUIRY").first()
            if enquiry_col:
                cell_subquery = Subquery(
                    CellValue.objects.filter(row=OuterRef('pk'), column=enquiry_col).values('value')[:1]
                )
                queryset = queryset.annotate(
                    enquiry_val=Cast(
                        Replace(
                            Cast(cell_subquery, output_field=TextField()),
                            Value('"'),
                            Value(''),
                            output_field=TextField()
                        ),
                        output_field=TextField()
                    )
                )
                if sort_dir == 'desc':
                    queryset = queryset.order_by('-enquiry_val', '-id')
                else:
                    queryset = queryset.order_by('enquiry_val', 'id')
            else:
                queryset = queryset.order_by('id')
        elif sort_by in ['due_date', 'return_date'] and table.job_type in ['GENERAL', 'ENGINEER', 'LIST_PID', 'LOGS']:
            if table.job_type == 'LIST_PID':
                from django.db.models import Subquery, OuterRef
                due_col = Column.objects.filter(table=table, name__iexact="DUE_DATE_FLOW_FORCE").first()
                if due_col:
                    cell_subquery = Subquery(
                        CellValue.objects.filter(row=OuterRef('pk'), column=due_col).values('value')[:1]
                    )
                    queryset = queryset.annotate(
                        due_date_str=Cast(
                            Replace(
                                Cast(cell_subquery, output_field=TextField()),
                                Value('"'),
                                Value(''),
                                output_field=TextField()
                            ),
                            output_field=TextField()
                        )
                    )
                    if sort_dir == 'desc':
                        queryset = queryset.order_by('-due_date_str', '-id')
                    else:
                        queryset = queryset.order_by('due_date_str', 'id')
                else:
                    if sort_dir == 'desc':
                        queryset = queryset.order_by('-task__due_date', '-id')
                    else:
                        queryset = queryset.order_by('task__due_date', 'id')
            else:
                if sort_dir == 'desc':
                    queryset = queryset.order_by('-task__due_date', '-id')
                else:
                    queryset = queryset.order_by('task__due_date', 'id')
        elif sort_by == 'follow_up_date' and table.job_type == 'SALES':
            from django.db.models import Max, DateField
            from django.db.models.functions import Coalesce
            queryset = queryset.annotate(
                latest_follow_up=Max('task__follow_ups__follow_up_date')
            ).annotate(
                sorted_follow_up=Coalesce('latest_follow_up', 'task__due_date', output_field=DateField())
            )
            if sort_dir == 'desc':
                queryset = queryset.order_by('-sorted_follow_up', '-id')
            else:
                queryset = queryset.order_by('sorted_follow_up', 'id')
        elif sort_by == 'date_assigned':
            from django.db.models import Subquery, OuterRef
            from django.db.models.functions import Coalesce
            date_col = Column.objects.filter(table=table, name__iexact="DATE").first()
            if date_col:
                cell_subquery = Subquery(
                    CellValue.objects.filter(row=OuterRef('pk'), column=date_col).values('value')[:1]
                )
                queryset = queryset.annotate(
                    date_assigned_val=Cast(
                        Replace(
                            Cast(cell_subquery, output_field=TextField()),
                            Value('"'),
                            Value(''),
                            output_field=TextField()
                        ),
                        output_field=TextField()
                    )
                ).annotate(
                    sorted_date_assigned=Coalesce('date_assigned_val', Cast('created_at', output_field=TextField()), output_field=TextField())
                )
            else:
                queryset = queryset.annotate(
                    sorted_date_assigned=Cast('created_at', output_field=TextField())
                )
            if sort_dir == 'desc':
                queryset = queryset.order_by('-sorted_date_assigned', '-id')
            else:
                queryset = queryset.order_by('sorted_date_assigned', 'id')

    return queryset

class TableViewSet(viewsets.ModelViewSet):
    serializer_class = TableSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return get_accessible_tables(self.request.user).select_related('department').prefetch_related('columns')

    def perform_create(self, serializer):
        # Automatically assign creator and department if admin/department admin
        dept = self.request.user.department if self.request.user.role in ["ADMIN", "DEPARTMENT_ADMIN"] else None
        serializer.save(created_by=self.request.user, department=dept)

    @action(detail=True, methods=["post"], url_path="share")
    def share_table(self, request, pk=None):
        table = self.get_object_or_404(pk)
        if not has_table_access(request.user, table, "ADMIN"):
            return Response({"error": "Only admins can share this table"}, status=status.HTTP_403_FORBIDDEN)

        user_id = request.data.get("user")
        dept_id = request.data.get("department")
        access_level = request.data.get("access_level", "VIEW")

        if not user_id and not dept_id:
            return Response({"error": "Must provide user or department"}, status=status.HTTP_400_BAD_REQUEST)

        with transaction.atomic():
            if user_id:
                user = get_object_or_404(EmployeeUser, id=user_id)
                access, created = TableAccess.objects.update_or_create(
                    table=table, user=user,
                    defaults={"access_level": access_level}
                )
            else:
                from employee_management.models import Department
                dept = get_object_or_404(Department, id=dept_id)
                access, created = TableAccess.objects.update_or_create(
                    table=table, department=dept,
                    defaults={"access_level": access_level}
                )
        return Response(TableAccessSerializer(access).data, status=status.HTTP_200_OK)

    def update(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance, "ADMIN"):
            return Response({"error": "Only admins can edit this table"}, status=status.HTTP_403_FORBIDDEN)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance, "ADMIN"):
            return Response({"error": "Only admins can edit this table"}, status=status.HTTP_403_FORBIDDEN)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance, "ADMIN"):
            return Response({"error": "Only admins can delete this table"}, status=status.HTTP_403_FORBIDDEN)
        return super().destroy(request, *args, **kwargs)

    @action(detail=True, methods=["post"], url_path="duplicate")
    def duplicate_table(self, request, pk=None):
        table = self.get_object_or_404(pk)
        if not has_table_access(request.user, table, "ADMIN"):
            return Response({"error": "Only admins can duplicate this table"}, status=status.HTTP_403_FORBIDDEN)

        try:
            new_table = TableDuplicateService.duplicate_table(table, request.user)
            return Response(TableSerializer(new_table).data, status=status.HTTP_201_CREATED)
        except Exception as e:
            logger.exception("Internal Server Error during table duplication: %s", str(e))
            return Response(
                {"error": "Failed to duplicate table. An internal server error occurred."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )


    @action(detail=True, methods=["get"], url_path="stats")
    def stats(self, request, pk=None):
        table = self.get_object_or_404(pk)
        if not has_table_access(request.user, table, "VIEW"):
            return Response({"error": "No view access to this table"}, status=status.HTTP_403_FORBIDDEN)
        stats_data = TableStatisticsService.get_table_statistics(table)
        return Response({
            'unique_pids': stats_data['unique_pids'],
            'unique_years': stats_data['unique_years'],
            'unique_column_values': stats_data['unique_column_values'],
            'stats': {
                'status_counts': stats_data['status_counts'],
                'priority_counts': stats_data['priority_counts'],
                'project_counts': stats_data['project_counts'],
                'due_today_count': stats_data['due_today_count'],
                'overdue_count': stats_data['overdue_count'],
                'total_qty': stats_data['total_qty'],
                'completion_stats': stats_data['completion_stats'],
                'week_actuals': stats_data['week_actuals'],
            }
        }, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="bulk-delete-rows")
    @transaction.atomic
    def bulk_delete_rows(self, request, pk=None):
        table = self.get_object_or_404(pk)
        if not has_table_access(request.user, table, "EDIT"):
            return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)
        
        row_ids = request.data.get("row_ids")
        try:
            count = TableDeleteService.bulk_delete_rows(table, request.user, row_ids=row_ids)
            return Response({"message": f"Successfully deleted {count} rows"}, status=status.HTTP_200_OK)
        except DeleteValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=["post"], url_path="send-escalation")
    @transaction.atomic
    def send_manual_escalation(self, request, pk=None):
        table = self.get_object_or_404(pk)
        # Verify access: Admin or Super Admin role globally, or table access ADMIN
        if not (request.user.role in ["SUPER_ADMIN", "ADMIN"] or has_table_access(request.user, table, "ADMIN")):
            return Response({"error": "Only admins can trigger escalation emails"}, status=status.HTTP_403_FORBIDDEN)

        row_ids = request.data.get("row_ids")
        today = timezone.localdate()
        if row_ids is None:
            rows = Row.objects.filter(table=table)
            # Find all associated tasks that are overdue (due_date < today) and not completed/approved
            tasks = Task.objects.filter(row__in=rows, due_date__lt=today).exclude(status__in=['COMPLETED', 'APPROVED'])
        else:
            if not isinstance(row_ids, list):
                return Response({"error": "row_ids must be a list"}, status=status.HTTP_400_BAD_REQUEST)
            rows = Row.objects.filter(table=table, id__in=row_ids)
            # When rows are explicitly selected, we do not require the tasks to be overdue. We send for all selected tasks that are not completed/approved.
            tasks = Task.objects.filter(row__in=rows).exclude(status__in=['COMPLETED', 'APPROVED'])

        from django.template.loader import render_to_string
        from tasks.models import EmailLog
        from tasks.tasks import send_email_log_task
        from django.conf import settings

        sent_count = 0
        for task in tasks:
            if task.due_date:
                days_overdue = (today - task.due_date).days
                if days_overdue < 0:
                    days_overdue = 0
            else:
                days_overdue = 0

            recipients = list(task.assigned_to.all())
            unique_recipients = []
            seen_emails = set()
            for r in recipients:
                if r.email and r.email not in seen_emails:
                    seen_emails.add(r.email)
                    unique_recipients.append(r)

            if unique_recipients:
                task.last_escalation_level = days_overdue
                task.last_escalation_at = timezone.now()
                task.save()

                for recipient in unique_recipients:
                    if days_overdue > 0:
                        subject = f"ESCALATION: Overdue Task - {task.task_name} ({days_overdue} days overdue)"
                        intro_html = f"A task is <strong>{days_overdue}</strong> days overdue and requires immediate attention."
                    else:
                        subject = f"ESCALATION: Task Escalated - {task.task_name}"
                        intro_html = "A task has been escalated and requires immediate attention."

                    site_url = getattr(settings, 'SITE_URL', 'https://flowforceworkspace.cloud')
                    task_link = f"{site_url}/tables/{task.row.table_id}/?open_task_id={task.id}"

                    context = {
                        'recipient_name': recipient.full_name,
                        'days': days_overdue,
                        'intro_html': intro_html,
                        'task_name': task.task_name,
                        'due_date': str(task.due_date) if task.due_date else "Not Set",
                        'employee_name': ", ".join([u.full_name for u in task.assigned_to.all()]),
                        'department_name': task.row.table.department.name if task.row.table.department else "Global",
                        'status': task.status,
                        'priority': task.priority,
                        'task_link': task_link,
                        'pid_data': task.pid_data,
                        'customer_name_data': task.customer_name_data,
                        'task_name_data': task.task_name_data,
                    }

                    html_message = render_to_string('emails/overdue_escalation_mail.html', context)

                    email_log = EmailLog.objects.create(
                        recipient_email=recipient.email,
                        subject=subject,
                        body=html_message,
                        task=task,
                        email_type='OVERDUE_ESCALATION_MAIL',
                        status='PENDING',
                        max_retries=3,
                    )
                    send_email_log_task.delay(email_log.id)
                    sent_count += 1

        return Response({"message": f"Successfully sent escalation emails to {sent_count} recipients"}, status=status.HTTP_200_OK)

    def safe_parse_date(self, val):
        """Backward-compatibility delegate to TableImportService.safe_parse_date."""
        return TableImportService.safe_parse_date(val)

    def _import_rows_from_csv_data(self, file_data, table, request_user, header_row=None, data_row=None):
        """Backward-compatibility delegate to TableImportService.import_rows_from_csv_data."""
        return TableImportService.import_rows_from_csv_data(
            file_data=file_data,
            table=table,
            user=request_user,
            header_row=header_row,
            data_row=data_row,
        )

    @action(detail=True, methods=["post"], url_path="import-csv")
    @transaction.atomic
    def import_csv(self, request, pk=None):
        try:
            table = self.get_object_or_404(pk)
            if not has_table_access(request.user, table, "EDIT"):
                return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)

            csv_file = request.FILES.get("file")
            if not csv_file:
                return Response({"error": "No CSV file provided"}, status=status.HTTP_400_BAD_REQUEST)

            try:
                file_data = csv_file.read().decode("utf-8")
            except Exception:
                return Response({"error": "Failed to decode CSV file. Make sure it is encoded in UTF-8."}, status=status.HTTP_400_BAD_REQUEST)

            header_row = request.data.get("header_row") or request.POST.get("header_row")
            data_row = request.data.get("data_row") or request.POST.get("data_row")
            try:
                header_row = int(header_row) if header_row else None
                data_row = int(data_row) if data_row else None
            except ValueError:
                return Response({"error": "header_row and data_row must be integers"}, status=status.HTTP_400_BAD_REQUEST)

            created_rows, err = TableImportService.import_rows_from_csv_data(
                file_data=file_data,
                table=table,
                user=request.user,
                header_row=header_row,
                data_row=data_row,
            )
            if err:
                return Response({"error": err}, status=status.HTTP_400_BAD_REQUEST)

            return Response({"message": f"Successfully imported {len(created_rows)} rows"}, status=status.HTTP_201_CREATED)
        except Exception as e:
            logger.exception("Internal Server Error during CSV import: %s", str(e))
            return Response({
                "error": "Failed to import CSV. An internal server error occurred."
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=["post"], url_path="import-google-sheet")
    @transaction.atomic
    def import_google_sheet(self, request, pk=None):
        try:
            table = self.get_object_or_404(pk)
            if not has_table_access(request.user, table, "EDIT"):
                return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)

            sheet_url = request.data.get("url")
            if not sheet_url:
                return Response({"error": "No Google Sheet URL provided"}, status=status.HTTP_400_BAD_REQUEST)

            import re
            match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", sheet_url)
            if not match:
                return Response({"error": "Invalid Google Sheets URL format. Make sure it contains '/spreadsheets/d/[ID]'"}, status=status.HTTP_400_BAD_REQUEST)

            header_row = request.data.get("header_row")
            data_row = request.data.get("data_row")
            try:
                header_row = int(header_row) if header_row else None
                data_row = int(data_row) if data_row else None
            except ValueError:
                return Response({"error": "header_row and data_row must be integers"}, status=status.HTTP_400_BAD_REQUEST)

            created_rows, err = TableImportService.import_from_google_sheet(
                sheet_url=sheet_url,
                table=table,
                user=request.user,
                header_row=header_row,
                data_row=data_row,
            )
            if err:
                return Response({"error": err}, status=status.HTTP_400_BAD_REQUEST)

            return Response({"message": f"Successfully imported {len(created_rows)} rows from Google Sheets"}, status=status.HTTP_201_CREATED)
        except Exception as e:
            logger.exception("Internal Server Error during Google Sheet import: %s", str(e))
            return Response({
                "error": "Failed to import Google Sheet. An internal server error occurred."
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=["post"], url_path="bulk-update")
    @transaction.atomic
    def bulk_update(self, request, pk=None):
        table = self.get_object_or_404(pk)
        field = request.data.get("field")
        value = request.data.get("value")

        try:
            updated_count = RowMutationService.bulk_update_table(
                table=table,
                field=field,
                value=value,
                user=request.user,
            )
        except RowPermissionDeniedError as e:
            return Response({"error": str(e)}, status=status.HTTP_403_FORBIDDEN)
        except RowValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response({"message": f"Successfully updated {updated_count} rows"}, status=status.HTTP_200_OK)


    @action(detail=True, methods=["get"], url_path="export-excel")
    def export_excel(self, request, pk=None):
        import io
        import re
        import logging
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from django.http import HttpResponse
        from datetime import datetime, date

        logger = logging.getLogger(__name__)

        try:
            table = get_object_or_404(Table, pk=pk)
            if not has_table_access(request.user, table, "VIEW"):
                return Response({"error": "No view access to this table"}, status=status.HTTP_403_FORBIDDEN)

            rows_qs = get_filtered_table_rows(table, request.query_params, user=request.user)
            columns = list(table.columns.all().order_by("position", "id"))

            wb = openpyxl.Workbook()
            ws = wb.active
            
            # Clean title (max 31 chars, no invalid chars : \ / ? * [ ])
            clean_title = re.sub(r'[:\\/?*\[\]]', '_', (table.name or "Export"))[:31]
            ws.title = clean_title or "Export"

            # Define Styles
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            data_font = Font(name="Calibri", size=10)
            data_alignment_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
            data_alignment_center = Alignment(horizontal="center", vertical="top")
            data_alignment_right = Alignment(horizontal="right", vertical="top")

            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            # Write Header Row
            headers = [col.name for col in columns]
            ws.append(headers)
            ws.row_dimensions[1].height = 28

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Check multiline columns
            def is_col_multiline(c):
                if c.name and c.name.upper() == "DESCRIPTION":
                    return True
                if c.data_type == "TEXT" and c.options and str(c.options).strip().startswith("{"):
                    try:
                        import json
                        c_opts = json.loads(c.options)
                        if c_opts.get("input_type") == "multiline" or c_opts.get("multiline") is True:
                            return True
                    except Exception:
                        pass
                return False

            # Write data rows
            row_idx = 1
            for row in rows_qs.iterator(chunk_size=1000):
                row_idx += 1
                cell_map = {c.column_id: c.value for c in row.cells.all()}
                row_data = []
                
                for col in columns:
                    raw_val = cell_map.get(col.id)
                    formatted_val = raw_val
                    
                    if raw_val is None or str(raw_val).strip() == "":
                        formatted_val = ""
                    elif col.data_type == "NUMBER":
                        try:
                            val_str = str(raw_val).strip()
                            if "." in val_str:
                                formatted_val = float(val_str)
                            else:
                                formatted_val = int(val_str)
                        except (ValueError, TypeError):
                            formatted_val = str(raw_val)
                    elif col.data_type in ["DATE", "DATETIME"]:
                        try:
                            date_str = str(raw_val).split("T")[0].strip()
                            formatted_val = datetime.strptime(date_str, "%Y-%m-%d").date()
                        except (ValueError, TypeError):
                            formatted_val = str(raw_val)
                    elif col.data_type == "CHECKBOX" or col.name.upper() in ["INITIAL_MAIL", "ALERT_MAIL"]:
                        is_true = str(raw_val).lower().strip() in ["true", "1", "yes"]
                        formatted_val = "YES" if is_true else "NO"
                    else:
                        formatted_val = str(raw_val)
                    
                    row_data.append(formatted_val)

                ws.append(row_data)
                ws.row_dimensions[row_idx].height = 20

                for col_idx, (col, val) in enumerate(zip(columns, row_data), 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = data_font
                    cell.border = thin_border
                    
                    if col.data_type == "NUMBER":
                        cell.alignment = data_alignment_right
                    elif col.data_type in ["DATE", "DATETIME"]:
                        cell.alignment = data_alignment_center
                        if isinstance(val, (datetime, date)):
                            cell.number_format = 'yyyy-mm-dd'
                    elif col.data_type == "CHECKBOX":
                        cell.alignment = data_alignment_center
                    elif is_col_multiline(col):
                        cell.alignment = data_alignment_left
                    else:
                        cell.alignment = data_alignment_left

            # Auto-adjust column widths
            for col_idx, col in enumerate(columns, 1):
                col_letter = get_column_letter(col_idx)
                header_len = len(col.name or "")
                max_len = header_len
                
                col_cells = ws[col_letter]
                sample_cells = col_cells[1:101] if len(col_cells) > 1 else ()
                for cell in sample_cells:
                    if cell.value is not None:
                        lines = str(cell.value).split("\n")
                        longest_line = max((len(l) for l in lines), default=0)
                        max_len = max(max_len, longest_line)
                
                if is_col_multiline(col):
                    ws.column_dimensions[col_letter].width = min(max(max_len + 4, 30), 60)
                elif col.data_type in ["DATE", "DATETIME"]:
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
                elif col.data_type == "NUMBER":
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
                else:
                    ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 45)

            # Generate Sensible Filename
            today_str = timezone.localdate().strftime("%Y-%m-%d")
            safe_name = re.sub(r'[^a-zA-Z0-9_-]', '_', (table.name or "table").strip().lower())
            safe_name = re.sub(r'_+', '_', safe_name).strip('_') or "export"
            filename = f"{safe_name}_export_{today_str}.xlsx"

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)

            response = HttpResponse(
                output.getvalue(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            response["Access-Control-Expose-Headers"] = "Content-Disposition"
            return response

        except Exception as e:
            logger.exception("Error exporting Excel for table %s: %s", pk, str(e))
            return Response(
                {"error": "Failed to generate Excel export. An error occurred on the server."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    def get_object_or_404(self, pk):
        obj = get_object_or_404(Table, pk=pk)
        self.check_object_permissions(self.request, obj)
        return obj

class ColumnViewSet(viewsets.ModelViewSet):
    serializer_class = ColumnSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        table_id = self.request.query_params.get("table")
        if not table_id:
            if self.action in ["retrieve", "update", "partial_update", "destroy", "clear_values", "delete_rows"]:
                from .permissions import get_accessible_tables
                accessible_tables = get_accessible_tables(self.request.user)
                return Column.objects.filter(table__in=accessible_tables)
            return Column.objects.none()
        table = get_object_or_404(Table, id=table_id)
        if not has_table_access(self.request.user, table, "VIEW"):
            return Column.objects.none()
        return Column.objects.filter(table=table)

    def create(self, request, *args, **kwargs):
        table_id = request.data.get("table")
        table = get_object_or_404(Table, id=table_id)
        if not has_table_access(request.user, table, "ADMIN"):
            return Response({"error": "Only admins can add columns"}, status=status.HTTP_403_FORBIDDEN)
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance.table, "ADMIN"):
            return Response({"error": "Only admins can update columns"}, status=status.HTTP_403_FORBIDDEN)
        if instance.is_system_column:
            if request.data.get("name") and request.data.get("name") != instance.name:
                return Response({"error": "Cannot rename system columns"}, status=status.HTTP_400_BAD_REQUEST)
            if request.data.get("data_type") and request.data.get("data_type") != instance.data_type:
                if instance.name not in ("TASK_NAME", "CUSTOMER_NAME"):
                    return Response({"error": "Cannot change data type of system columns"}, status=status.HTTP_400_BAD_REQUEST)
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance.table, "ADMIN"):
            return Response({"error": "Only admins can update columns"}, status=status.HTTP_403_FORBIDDEN)
        if instance.is_system_column:
            if request.data.get("name") and request.data.get("name") != instance.name:
                return Response({"error": "Cannot rename system columns"}, status=status.HTTP_400_BAD_REQUEST)
            if request.data.get("data_type") and request.data.get("data_type") != instance.data_type:
                if instance.name not in ("TASK_NAME", "CUSTOMER_NAME"):
                    return Response({"error": "Cannot change data type of system columns"}, status=status.HTTP_400_BAD_REQUEST)
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance.table, "ADMIN"):
            return Response({"error": "Only admins can delete columns"}, status=status.HTTP_403_FORBIDDEN)
        if instance.is_system_column:
            return Response({"error": "Cannot delete system columns"}, status=status.HTTP_400_BAD_REQUEST)
        return super().destroy(request, *args, **kwargs)

    @action(detail=True, methods=["post"], url_path="clear-values")
    @transaction.atomic
    def clear_values(self, request, pk=None):
        column = self.get_object()
        if not has_table_access(request.user, column.table, "EDIT"):
            return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)
        
        try:
            TableDeleteService.clear_column_values(column, request.user)
            return Response({"message": f"Successfully cleared all values in column {column.name}"}, status=status.HTTP_200_OK)
        except DeleteValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

    @action(detail=True, methods=["post"], url_path="delete-rows")
    @transaction.atomic
    def delete_rows(self, request, pk=None):
        column = self.get_object()
        if not has_table_access(request.user, column.table, "EDIT"):
            return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)
        
        try:
            count = TableDeleteService.delete_rows_by_column(column, request.user)
            return Response({"message": f"Successfully deleted {count} rows containing values in column {column.name}"}, status=status.HTTP_200_OK)
        except DeleteValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)


    def perform_create(self, serializer):
        table = serializer.validated_data["table"]
        from django.db.models import Max
        max_pos = Column.objects.filter(table=table).aggregate(Max("position"))["position__max"] or 0
        serializer.save(position=max_pos + 1)

    @action(detail=False, methods=["post"], url_path="reorder")
    @transaction.atomic
    def reorder_columns(self, request):
        column_ids = request.data.get("columns", [])
        if not column_ids:
            return Response({"error": "No columns list provided"}, status=status.HTTP_400_BAD_REQUEST)
        
        columns = Column.objects.filter(id__in=column_ids)
        if not columns.exists():
            return Response({"error": "No columns found for provided IDs"}, status=status.HTTP_404_NOT_FOUND)
        
        table = columns.first().table
        if columns.filter(table=table).count() != len(column_ids):
            return Response({"error": "All columns must belong to the same table"}, status=status.HTTP_400_BAD_REQUEST)
            
        if not has_table_access(request.user, table, "ADMIN"):
            return Response({"error": "Only admins can reorder columns"}, status=status.HTTP_403_FORBIDDEN)
            
        for index, col_id in enumerate(column_ids):
            Column.objects.filter(id=col_id, table=table).update(position=index + 1)
            
        return Response({"status": "reordered"}, status=status.HTTP_200_OK)

from rest_framework.pagination import PageNumberPagination
from django.db.models import Count

class RowPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = 'page_size'
    max_page_size = 500

    def get_paginated_response(self, data):
        table_id = self.request.query_params.get("table")
        if not table_id:
            return super().get_paginated_response(data)
            
        table = getattr(self.request, "_cached_table", None)
        if not table or str(table.id) != str(table_id):
            table = get_object_or_404(Table, id=table_id)
        include_stats = self.request.query_params.get("include_stats") == "true"

        if include_stats:
            stats_data = TableStatisticsService.get_table_statistics(table)
        else:
            stats_data = TableStatisticsService.get_default_statistics()

        return Response({
            'count': self.page.paginator.count,
            'next': self.get_next_link(),
            'previous': self.get_previous_link(),
            'results': data,
            'unique_pids': stats_data['unique_pids'],
            'unique_years': stats_data['unique_years'],
            'unique_column_values': stats_data['unique_column_values'],
            'stats': {
                'status_counts': stats_data['status_counts'],
                'priority_counts': stats_data['priority_counts'],
                'project_counts': stats_data['project_counts'],
                'due_today_count': stats_data['due_today_count'],
                'overdue_count': stats_data['overdue_count'],
                'total_qty': stats_data['total_qty'],
                'completion_stats': stats_data['completion_stats'],
                'week_actuals': stats_data['week_actuals'],
            }
        })


class RowViewSet(viewsets.ModelViewSet):
    serializer_class = RowSerializer
    permission_classes = [permissions.IsAuthenticated]
    pagination_class = RowPagination

    def get_queryset(self):
        table_id = self.request.query_params.get("table")
        if not table_id:
            if self.action in ["retrieve", "update", "partial_update", "destroy"]:
                from .permissions import get_accessible_tables
                from django.db.models import Prefetch
                accessible_tables = get_accessible_tables(self.request.user)
                return Row.objects.filter(
                    table__in=accessible_tables, is_archived=False
                ).select_related(
                    'created_by', 'task', 'task__assigned_by'
                ).prefetch_related(
                    Prefetch('cells', queryset=CellValue.objects.select_related('column', 'updated_by')),
                    'task__assigned_to'
                )
            return Row.objects.none()
            
        table = get_object_or_404(Table, id=table_id)
        if not has_table_access(self.request.user, table, "VIEW"):
            return Row.objects.none()
        self.request._cached_table = table
        return get_filtered_table_rows(table, self.request.query_params, user=self.request.user)

    def destroy(self, request, *args, **kwargs):
        instance = self.get_object()
        if not has_table_access(request.user, instance.table, "EDIT"):
            return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)
        TableDeleteService.delete_row(instance, request.user)
        return Response(status=status.HTTP_204_NO_CONTENT)

    @transaction.atomic
    def create(self, request, *args, **kwargs):
        table_id = request.data.get("table")
        table = get_object_or_404(Table.objects.select_for_update(), id=table_id)
        
        if not has_table_access(request.user, table, "EDIT"):
            return Response({"error": "No edit access to this table"}, status=status.HTTP_403_FORBIDDEN)

        cells_data = request.data.get("cells", {})
        assigned_to_ids = request.data.get("assigned_to", [])

        try:
            row = RowService.create_row(
                table=table,
                user=request.user,
                cells_data=cells_data,
                assigned_to_ids=assigned_to_ids,
            )
        except RowCreationValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(RowSerializer(row).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"], url_path="edit-cell")
    @transaction.atomic
    def edit_cell(self, request, pk=None):
        row = get_object_or_404(Row, pk=pk)
        column_id = request.data.get("column")
        value = request.data.get("value")
        column = get_object_or_404(Column, id=column_id, table=row.table)

        try:
            CellMutationService.update_cell(
                row=row,
                column=column,
                value=value,
                user=request.user,
                check_permissions=True,
            )
        except CellPermissionDeniedError as e:
            return Response({"error": str(e)}, status=status.HTTP_403_FORBIDDEN)
        except CellValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(RowSerializer(row).data, status=status.HTTP_200_OK)


    @action(detail=True, methods=["post"], url_path="toggle-status")
    @transaction.atomic
    def toggle_status(self, request, pk=None):
        row = get_object_or_404(Row, pk=pk)
        table = row.table

        # Check permissions: must have VIEW access to table
        if not has_table_access(request.user, table, "VIEW"):
            return Response({"error": "No access to this table"}, status=status.HTTP_403_FORBIDDEN)

        from tasks.models import Task, ActivityLog
        from django.utils import timezone
        task = getattr(row, "task", None)
        if not task:
            from datetime import datetime
            due_date = None
            date_cell = row.cells.filter(column__name__in=["DUE_DATE", "FOLLOW_UP_DATE", "RETURN_DATE", "DUE_DATE_FLOW_FORCE"]).first()
            if date_cell and date_cell.value:
                try:
                    due_date = datetime.strptime(str(date_cell.value).split("T")[0], "%Y-%m-%d").date()
                except Exception:
                    due_date = None
            task = Task.objects.create(
                row=row,
                due_date=due_date,
                priority="MEDIUM",
                status="PENDING",
                assigned_by=request.user
            )

        requested_status = request.data.get("status")
        if requested_status in dict(Task.STATUS_CHOICES):
            new_status = requested_status
        else:
            new_status = "PENDING" if (task.status in ["COMPLETED", "APPROVED"]) else "COMPLETED"

        old_status = task.status
        task.status = new_status
        task.save(update_fields=["status"])

        # Sync back to STATUS cell if such a column exists
        status_col = Column.objects.filter(table=table, name__iexact="STATUS").first()
        if status_col:
            CellValue.objects.update_or_create(
                row=row, column=status_col,
                defaults={"value": new_status, "updated_by": request.user}
            )

        from tasks.models import ActivityLog
        ActivityLog.objects.create(
            task=task,
            action=f"Changed status from {old_status} to {new_status}",
            user=request.user,
            details={"old_status": old_status, "new_status": new_status}
        )

        return Response(RowSerializer(row).data, status=status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="edit-row")
    @transaction.atomic
    def edit_row(self, request, pk=None):
        row = get_object_or_404(Row, pk=pk)
        try:
            row, _ = RowMutationService.edit_row(
                row=row,
                cells_data=request.data.get("cells", {}),
                user=request.user,
            )
        except RowPermissionDeniedError as e:
            return Response({"error": str(e)}, status=status.HTTP_403_FORBIDDEN)
        except RowValidationError as e:
            return Response({"error": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(RowSerializer(row).data, status=status.HTTP_200_OK)


class TableAccessViewSet(viewsets.ModelViewSet):
    serializer_class = TableAccessSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return TableAccess.objects.all().select_related('user', 'department')

class ColumnAccessViewSet(viewsets.ModelViewSet):
    serializer_class = ColumnAccessSerializer
    permission_classes = [permissions.IsAuthenticated]

    def get_queryset(self):
        return ColumnAccess.objects.all().select_related('user')

from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required

@login_required
def table_spreadsheet_view(request, table_id):
    table = get_object_or_404(
        Table.objects.select_related('department').prefetch_related('columns'),
        id=table_id
    )
    if not has_table_access(request.user, table, "VIEW"):
        return redirect("/")
    has_edit = has_table_access(request.user, table, "EDIT")
    has_admin = has_table_access(request.user, table, "ADMIN")
    has_follow_up = (
        table.job_type in ["SALES", "LIST_PID"] or
        (table.department and table.department.name.lower() == "sales") or
        any(c.name.upper() in ["FOLLOW_UP_DATE", "FOLLOW-UP DATE", "FOLLOW - UP DATE"] for c in table.columns.all())
    )
    return render(request, "tables/table_spreadsheet.html", {
        "table": table,
        "has_edit_access": has_edit,
        "has_admin_access": has_admin,
        "has_follow_up": has_follow_up
    })

@login_required
def table_create_view(request):
    if request.user.role not in ["SUPER_ADMIN", "ADMIN"]:
        return redirect("/")
    if request.method == "POST":
        name = request.POST.get("name")
        description = request.POST.get("description")
        job_type = request.POST.get("job_type", "GENERAL")
        table = Table.objects.create(name=name, description=description, job_type=job_type, created_by=request.user)
        # Create TableAccess for the creator as ADMIN
        TableAccess.objects.create(table=table, user=request.user, access_level="ADMIN")
        return redirect(f"/tables/{table.id}/")
    return render(request, "tables/table_create.html")

@login_required
def table_list_view(request):
    from django.contrib import messages
    from auth_app.models import EmployeeUser
    from .models import Table, TableAccess
    from .permissions import get_accessible_tables

    is_admin = request.user.role in ["SUPER_ADMIN", "ADMIN"]

    if request.method == "POST" and is_admin:
        action = request.POST.get("action")

        if action == "create":
            name = request.POST.get("name")
            description = request.POST.get("description")
            job_type = request.POST.get("job_type", "GENERAL")
            if name:
                table = Table.objects.create(name=name, description=description, job_type=job_type, created_by=request.user)
                TableAccess.objects.create(
                    table=table, user=request.user, access_level="ADMIN"
                )
                messages.success(request, f"Table '{name}' created successfully.")
            return redirect("tables:table_list")

        elif action == "delete":
            table_id = request.POST.get("table_id")
            if not table_id or not str(table_id).isdigit():
                messages.error(request, "Invalid table selected.")
                return redirect("tables:table_list")
            table = get_object_or_404(Table, id=table_id)
            if table.created_by == request.user or request.user.role == "SUPER_ADMIN":
                table.delete()
                messages.success(request, "Table deleted successfully.")
            else:
                messages.error(request, "You do not have permission to delete this table.")
            return redirect("tables:table_list")

        elif action == "grant":
            table_id = request.POST.get("table_id")
            user_id = request.POST.get("user_id")
            access_level = request.POST.get("access_level", "EDIT")
            if not table_id or not str(table_id).isdigit() or not user_id or not str(user_id).isdigit():
                messages.error(request, "Please select a valid table and employee.")
                return redirect("tables:table_list")
            table = get_object_or_404(Table, id=table_id)
            user = get_object_or_404(EmployeeUser, id=user_id)

            TableAccess.objects.update_or_create(
                table=table, user=user,
                defaults={"access_level": access_level}
            )
            messages.success(request, f"Access granted to {user.full_name or user.email}.")
            return redirect("tables:table_list")

        elif action == "revoke":
            table_id = request.POST.get("table_id")
            user_id = request.POST.get("user_id")
            if not table_id or not str(table_id).isdigit() or not user_id or not str(user_id).isdigit():
                messages.error(request, "Please select a valid table and employee.")
                return redirect("tables:table_list")
            table = get_object_or_404(Table, id=table_id)
            user = get_object_or_404(EmployeeUser, id=user_id)

            TableAccess.objects.filter(table=table, user=user).delete()
            messages.success(request, f"Access revoked for {user.full_name or user.email}.")
            return redirect("tables:table_list")

        elif action == "change_access":
            table_id = request.POST.get("table_id")
            user_id = request.POST.get("user_id")
            access_level = request.POST.get("access_level")
            if not table_id or not str(table_id).isdigit() or not user_id or not str(user_id).isdigit():
                messages.error(request, "Please select a valid table and employee.")
                return redirect("tables:table_list")
            table = get_object_or_404(Table, id=table_id)
            user = get_object_or_404(EmployeeUser, id=user_id)

            TableAccess.objects.filter(table=table, user=user).update(access_level=access_level)
            messages.success(request, f"Access level updated to {access_level}.")
            return redirect("tables:table_list")

    # GET handling
    if is_admin:
        tables = Table.objects.filter(is_active=True).prefetch_related('access_rules__user')
        employees = EmployeeUser.objects.filter(is_active=True).exclude(role="SUPER_ADMIN")
    else:
        tables = get_accessible_tables(request.user)
        employees = None

    return render(
        request,
        "tables/table_list.html",
        {
            "tables": tables,
            "employees": employees,
            "is_admin": is_admin,
        }
    )

@login_required
def tables_analytics_dashboard(request):
    from tables.models import Table
    from tables.permissions import get_accessible_tables
    from tasks.models import Task
    from django.utils import timezone

    if request.user.role in ["SUPER_ADMIN", "ADMIN"]:
        tables = Table.objects.filter(is_active=True).select_related('department')
    else:
        tables = get_accessible_tables(request.user).select_related('department')

    from django.db.models import Count, Q
    today = timezone.localdate()

    annotated_tables = tables.annotate(
        total_tasks=Count('rows__task', filter=Q(rows__is_archived=False)),
        pending_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__status="PENDING")),
        in_progress_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__status="IN_PROGRESS")),
        ready_for_review_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__status="READY_FOR_REVIEW")),
        completed_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__status__in=["COMPLETED", "APPROVED"])),
        overdue_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__due_date__lt=today) & ~Q(rows__task__status__in=["COMPLETED", "APPROVED"])),
        due_today_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__due_date=today)),
        low_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__priority="LOW")),
        medium_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__priority="MEDIUM")),
        high_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__priority="HIGH")),
        critical_tasks=Count('rows__task', filter=Q(rows__is_archived=False, rows__task__priority="CRITICAL"))
    )

    tables_data = []
    for table in annotated_tables:
        total = table.total_tasks
        completed = table.completed_tasks
        completion_rate = int(completed * 100 / total) if total > 0 else 0

        tables_data.append({
            "table": table,
            "total": total,
            "pending": table.pending_tasks,
            "in_progress": table.in_progress_tasks,
            "ready_for_review": table.ready_for_review_tasks,
            "completed": completed,
            "overdue": table.overdue_tasks,
            "due_today": table.due_today_tasks,
            "low": table.low_tasks,
            "medium": table.medium_tasks,
            "high": table.high_tasks,
            "critical": table.critical_tasks,
            "completion_rate": completion_rate,
        })

    return render(
        request,
        "tables/analytics_dashboard.html",
        {
            "tables_data": tables_data,
        },
    )


@login_required
def pid_dashboard_view(request):
    from tables.models import Table, Row, CellValue
    from tables.permissions import get_accessible_tables
    from django.utils import timezone
    from datetime import datetime
    import re

    today = timezone.localdate()

    if request.user.role in ["SUPER_ADMIN", "ADMIN"]:
        pid_tables = Table.objects.filter(is_active=True, job_type="LIST_PID").select_related('department').prefetch_related('columns')
    else:
        pid_tables = get_accessible_tables(request.user).filter(job_type="LIST_PID").select_related('department').prefetch_related('columns')

    pid_tables = list(pid_tables)

    from collections import defaultdict

    # 1. Batch fetch all rows across all PID tables in a single query
    all_rows = Row.objects.filter(
        table__in=pid_tables,
        is_archived=False
    ).select_related('created_by', 'task')

    rows_by_table = defaultdict(list)
    all_row_ids = []
    for r in all_rows:
        rows_by_table[r.table_id].append(r)
        all_row_ids.append(r.id)

    # 2. Identify relevant column IDs needed for the PID dashboard
    PID_RELEVANT_COLUMNS = {
        "PID", "NEW_PID_NO", "NEW PID NO",
        "ENQUIRY_NO/QUOTATION_NO", "QUOTATION_NO", "ENQUIRY_NO", "QUOTATION NO",
        "PO", "PO_NO", "PURCHASE ORDER",
        "SALES_ORDER", "SO", "SALES ORDER",
        "COMPANY_NAME", "CUSTOMER_NAME", "CUSTOMER",
        "DUE_DATE_CUSTOMER", "DUE_DATE_CUST",
        "DUE_DATE_FLOW_FORCE", "DUE_DATE_FF", "DUE_DATE",
        "DATE",
        "STATUS", "CURRENT_STATUS",
        "DESCRIPTION",
        "QTY",
        "PROJECT"
    }

    col_name_map = {}
    relevant_col_ids = []
    for table in pid_tables:
        for col in table.columns.all():
            upper_name = col.name.upper()
            if upper_name in PID_RELEVANT_COLUMNS:
                col_name_map[col.id] = upper_name
                relevant_col_ids.append(col.id)

    # 3. Batch fetch only relevant cell values in a single lightweight query
    cells_by_row = defaultdict(dict)
    if all_row_ids and relevant_col_ids:
        cell_values = CellValue.objects.filter(
            row_id__in=all_row_ids,
            column_id__in=relevant_col_ids
        ).values('row_id', 'column_id', 'value')

        for cv in cell_values:
            val = cv['value']
            if val is not None:
                col_name = col_name_map.get(cv['column_id'])
                if col_name:
                    cells_by_row[cv['row_id']][col_name] = val

    tables_pid_data = []
    quick_pid_summary = []
    available_years = set()
    overall_total = 0
    overall_in_progress = 0
    overall_completed = 0
    overall_overdue = 0
    overall_due_today = 0

    def extract_year(date_str, fallback_year):
        if not date_str or str(date_str).strip() in ["-", "", "None"]:
            return str(fallback_year)
        match = re.search(r'\b(20\d\d|19\d\d)\b', str(date_str))
        if match:
            return match.group(1)
        return str(fallback_year)

    for table in pid_tables:
        rows = rows_by_table.get(table.id, [])

        pids = []
        tbl_total = 0
        tbl_in_progress = 0
        tbl_completed = 0
        tbl_overdue = 0
        tbl_due_today = 0

        for r in rows:
            cells = cells_by_row.get(r.id, {})

            pid_val = str(cells.get("PID") or cells.get("NEW_PID_NO") or cells.get("NEW PID NO") or f"PID-{r.id}").strip()
            new_pid_val = str(cells.get("NEW_PID_NO") or cells.get("NEW PID NO") or "").strip()
            quotation_val = str(cells.get("ENQUIRY_NO/QUOTATION_NO") or cells.get("QUOTATION_NO") or cells.get("ENQUIRY_NO") or cells.get("QUOTATION NO") or "-").strip()
            po_val = str(cells.get("PO") or cells.get("PO_NO") or cells.get("PURCHASE ORDER") or "-").strip()
            so_val = str(cells.get("SALES_ORDER") or cells.get("SO") or cells.get("SALES ORDER") or "-").strip()
            customer_val = str(cells.get("COMPANY_NAME") or cells.get("CUSTOMER_NAME") or cells.get("CUSTOMER") or "-").strip()
            due_date_cust = str(cells.get("DUE_DATE_CUSTOMER") or cells.get("DUE_DATE_CUST") or "-").strip()
            due_date_ff = str(cells.get("DUE_DATE_FLOW_FORCE") or cells.get("DUE_DATE_FF") or cells.get("DUE_DATE") or "-").strip()
            date_cell = str(cells.get("DATE") or "").strip()
            status_cell = str(cells.get("STATUS") or cells.get("CURRENT_STATUS") or "").strip()
            desc_val = str(cells.get("DESCRIPTION") or "").strip()
            qty_val = str(cells.get("QTY") or "").strip()
            project_val = str(cells.get("PROJECT") or "").strip()

            fallback_y = r.created_at.year if r.created_at else today.year
            record_year = extract_year(date_cell, fallback_y)
            if record_year == str(fallback_y) and (due_date_ff or due_date_cust):
                record_year = extract_year(due_date_ff or due_date_cust, fallback_y)

            available_years.add(str(record_year))

            task_obj = getattr(r, 'task', None)
            exact_status = status_cell if status_cell else (task_obj.status if task_obj else "PENDING")
            current_status_upper = exact_status.upper()
            last_status = task_obj.status if task_obj else exact_status

            is_overdue = False
            is_due_today = False

            target_date = None
            if task_obj and task_obj.due_date:
                target_date = task_obj.due_date
            elif due_date_ff and due_date_ff != "-":
                try:
                    target_date = datetime.strptime(due_date_ff, "%Y-%m-%d").date()
                except Exception:
                    pass

            if target_date:
                if target_date < today and current_status_upper not in ["COMPLETED", "APPROVED"]:
                    is_overdue = True
                elif target_date == today:
                    is_due_today = True

            tbl_total += 1
            if current_status_upper in ["COMPLETED", "APPROVED"]:
                tbl_completed += 1
            else:
                tbl_in_progress += 1
                if is_overdue:
                    tbl_overdue += 1

            if is_due_today:
                tbl_due_today += 1

            pid_entry = {
                "row_id": r.id,
                "pid": pid_val,
                "new_pid": new_pid_val,
                "quotation_no": quotation_val,
                "po_number": po_val,
                "so_number": so_val,
                "customer_name": customer_val,
                "due_date_customer": due_date_cust,
                "due_date_flow_force": due_date_ff,
                "current_status": exact_status,
                "last_status": last_status,
                "description": desc_val,
                "qty": qty_val,
                "project": project_val,
                "year": str(record_year),
                "is_overdue": is_overdue,
                "is_due_today": is_due_today,
                "created_at": r.created_at,
            }
            pids.append(pid_entry)

            quick_pid_summary.append({
                "row_id": r.id,
                "pid": pid_val,
                "status": exact_status,
                "year": str(record_year),
                "table_id": table.id,
                "table_name": table.name,
                "customer": customer_val,
            })

        overall_total += tbl_total
        overall_in_progress += tbl_in_progress
        overall_completed += tbl_completed
        overall_overdue += tbl_overdue
        overall_due_today += tbl_due_today

        tables_pid_data.append({
            "table": table,
            "pids": pids,
            "total": tbl_total,
            "in_progress": tbl_in_progress,
            "completed": tbl_completed,
            "overdue": tbl_overdue,
            "due_today": tbl_due_today,
        })

    sorted_years = sorted(list(available_years), reverse=True)
    if "2026" not in sorted_years:
        sorted_years.insert(0, "2026")

    context = {
        "tables_pid_data": tables_pid_data,
        "quick_pid_summary": quick_pid_summary,
        "available_years": sorted_years,
        "selected_year_default": "2026",
        "overall_total": overall_total,
        "overall_in_progress": overall_in_progress,
        "overall_completed": overall_completed,
        "overall_overdue": overall_overdue,
        "overall_due_today": overall_due_today,
    }

    return render(request, "tables/pid_dashboard.html", context)




